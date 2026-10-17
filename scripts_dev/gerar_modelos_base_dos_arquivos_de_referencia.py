# scripts_dev/gerar_modelos_base_dos_arquivos_de_referencia.py

# Função Objetivo: Gera os MODELOS-BASE (arquivos .xlsx) que o botão "Modelo" da tela "Arquivos de referência" entrega, a partir dos arquivos
# REAIS da pasta "Arquivos usados para Popular Banco". Rode de novo sempre que o layout de um arquivo real mudar (cabeçalho, cores, colunas).
#
# * [DECISÃO, 17/10/2026] → O modelo tem de ser IGUAL ao arquivo real no visual e na estrutura (fonte, cores, bordas, larguras, alturas, formato
#   dos números, grupos mesclados, notas, fórmulas). Por isso nasce do próprio arquivo real.
# * [DECISÃO, 17/10/2026, Matheus] → O modelo só GUIA, não é uma planilha cheia: tem o cabeçalho, NO MÁXIMO 1 LINHA DE EXEMPLO e, logo depois,
#   uma LINHA CINZA ESCURA inteira que marca o fim do modelo. Exceção: o Mercado Livre tem 2 linhas de exemplo, porque cada faixa de peso é um
#   conjunto de 2 linhas (Sem e Com Frete Grátis Rápido). Nenhum valor de frete vem preenchido (o frete muda; um valor esquecido seria preço errado).
#   - O exemplo de cada tabela de frete é a PRIMEIRA faixa (fechada, com o valor em branco): se alguém enviar o modelo sem mexer, a auditoria barra
#     (última faixa de peso fechada e/ou valor vazio), em vez de gravar uma tabela de 1 faixa.
#   - A linha cinza não tem nenhum valor: os importadores e as auditorias ignoram linha sem dado. A nota (balãozinho) da célula A dela explica
#     o modelo e as regras do arquivo.
#   - Busca Legal: cada empresa tem a SUA ordem de colunas (a Samvale tem o CST antes do PIS); o exemplo é a 1ª linha real com CST 00.
# * [IMPORTANTE] → O script NUNCA altera os arquivos reais: abre, ajusta numa cópia em memória e grava em arquivos_referencia/modelos_base/.

from copy import copy
from math import ceil
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.formatting.formatting import ConditionalFormattingList
from openpyxl.styles import PatternFill
from openpyxl.styles.cell_style import StyleArray

RAIZ = Path(__file__).resolve().parents[1]
PASTA_DOS_ARQUIVOS_REAIS = RAIZ / 'Arquivos usados para Popular Banco'
PASTA_DOS_MODELOS = RAIZ / 'arquivos_referencia' / 'modelos_base'

COR_DA_LINHA_DE_FIM = 'FF404040'   # cinza escuro
ALTURA_DA_LINHA_DE_FIM = 12

NOTAS = {
    'frete-mercado-livre': (
        'FIM DO MODELO. As 2 linhas acima são o exemplo de UMA faixa de peso: cada faixa tem sempre 2 linhas, “Sem Frete Grátis Rápido” e “Com Frete Grátis Rápido” '
        '(com este texto exato). Apague-as junto com esta linha cinza e monte a tabela inteira a partir da linha 2. Preencha o frete (R$) das 8 faixas de preço. '
        'Na última faixa de peso, o Peso máx. é o texto “sem limite”.'
    ),
    'frete-magalu': (
        'FIM DO MODELO. A linha acima é o exemplo de UMA faixa de peso. Apague-a junto com esta linha cinza e monte a tabela inteira a partir da linha 2. '
        'Cada linha: o nome da faixa em A, os 3 fretes por reputação em B, C e D e, em E e F, o peso mínimo e o máximo (fórmulas que leem o nome da faixa: '
        'copie-as para baixo). Os descontos do cabeçalho são conferidos: C deve ser B menos 25% e D, B menos 50%. Na última faixa, o peso máximo fica vazio.'
    ),
    'frete-tiktok-shop': (
        'FIM DO MODELO. A linha acima é o exemplo de UMA faixa de peso. Apague-a junto com esta linha cinza e monte a tabela inteira a partir da linha 2. '
        'Cada linha: o nome da faixa em A, o valor médio do frete em B e o peso mínimo e o máximo em C e D. Na última faixa, o peso máximo fica vazio.'
    ),
    'frete-amazon': (
        'FIM DO MODELO. A linha acima é o exemplo de UMA combinação de faixa de peso e faixa de preço. Apague-a junto com esta linha cinza e monte a tabela '
        'inteira a partir da linha 2 (toda faixa de peso precisa de todas as faixas de preço). Nas faixas de preço mais baixas, que não dependem de peso '
        '(até R$ 78,99), o Peso mínimo e o Peso máximo ficam vazios. Preço máximo vazio = sem teto (última faixa de preço). O arquivo precisa das 2 abas, '
        'DBA e FBA, com os nomes exatos.'
    ),
    'busca-legal': (
        'FIM DO MODELO. A linha acima é só um exemplo de como preencher. Apague-a junto com esta linha cinza e coloque os seus produtos a partir da linha 3: '
        'uma linha por produto (o EAN é a chave). EAN, NCM e CST como TEXTO. PIS, COFINS e as UFs em percentual (18% = 0,18); vazio = sem imposto. '
        'Produtos com o mesmo NCM e CST precisam ter o mesmo percentual: se divergirem, o sistema rejeita o grupo inteiro.'
    ),
}


# Função Objetivo: Apaga por completo (valor, estilo, nota) as linhas da primeira até a última, desfazendo as mesclagens que tocam nelas.
def _remover_linhas(aba, primeira):
    ultima = aba.max_row
    for faixa in list(aba.merged_cells.ranges):
        if faixa.max_row >= primeira and faixa.min_row <= ultima:
            aba.unmerge_cells(str(faixa))
    for chave in [k for k in aba._cells if k[0] >= primeira]:
        del aba._cells[chave]
    for linha in [n for n in aba.row_dimensions if n >= primeira]:
        del aba.row_dimensions[linha]


# Função Objetivo: Esvazia o VALOR das células (o estilo, o formato e as notas ficam como no arquivo real).
def _esvaziar(aba, primeira_linha, ultima_linha, colunas):
    for linha in range(primeira_linha, ultima_linha + 1):
        for coluna in colunas:
            aba.cell(row=linha, column=coluna).value = None


# Função Objetivo: A linha cinza escura inteira (da coluna A à última), sem nenhum valor, com a nota que explica o modelo na célula A.
def _linha_de_fim(aba, linha, ultima_coluna, texto_da_nota):
    for coluna in range(1, ultima_coluna + 1):
        celula = aba.cell(row=linha, column=coluna)
        celula._style = StyleArray()
        celula.fill = PatternFill('solid', fgColor=COR_DA_LINHA_DE_FIM)
    aba.row_dimensions[linha].height = ALTURA_DA_LINHA_DE_FIM
    nota = Comment(texto_da_nota, 'Sistema Interno')
    nota.width = 380
    nota.height = 40 + 18 * ceil(len(texto_da_nota) / 52)
    aba.cell(row=linha, column=1).comment = nota


# Função Objetivo: Copia valor e estilo das células de uma linha para outra (A até a última coluna).
def _copiar_linha(aba, origem, destino, ultima_coluna):
    for coluna in range(1, ultima_coluna + 1):
        de = aba.cell(row=origem, column=coluna)
        para = aba.cell(row=destino, column=coluna)
        para.value = de.value
        para._style = copy(de._style)


def _modelo_ml(livro):
    aba = livro.active
    # As notas do rodapé (fonte, peso, reputação) ficam como no arquivo real: guardadas aqui e recolocadas depois da linha cinza.
    notas = []
    for linha in range(63, 67):
        celulas = {c: (aba.cell(row=linha, column=c).value if c == 1 else None, copy(aba.cell(row=linha, column=c)._style)) for c in range(1, 14)}
        notas.append((aba.row_dimensions[linha].height, celulas))
    assert all(isinstance(n[1][1][0], str) and n[1][1][0].strip() for n in notas), 'as 4 notas do rodapé do ML mudaram de lugar'

    _remover_linhas(aba, 4)   # fica o cabeçalho (1) e a 1ª faixa de peso (linhas 2 e 3: Sem e Com Frete Grátis Rápido)
    _esvaziar(aba, 2, 3, range(6, 14))   # F..M: os 8 valores de frete
    _linha_de_fim(aba, 4, 13, NOTAS['frete-mercado-livre'])
    for posicao, (altura, celulas) in enumerate(notas):
        linha = 6 + posicao
        for coluna, (valor, estilo) in celulas.items():
            celula = aba.cell(row=linha, column=coluna)
            celula._style = estilo
            if coluna == 1:
                celula.value = valor
        aba.merge_cells(start_row=linha, start_column=1, end_row=linha, end_column=13)
        aba.row_dimensions[linha].height = altura

    # A cor do "false" da coluna E (formatação condicional) valia de E2 a E61: agora só as 2 linhas do exemplo.
    antigas = list(aba.conditional_formatting)
    aba.conditional_formatting = ConditionalFormattingList()
    for formato in antigas:
        for regra in formato.rules:
            aba.conditional_formatting.add('E2:E3', regra)


def _modelo_magalu(livro):
    aba = livro.active
    _remover_linhas(aba, 3)   # fica o cabeçalho e a 1ª faixa ("0-0,499")
    _esvaziar(aba, 2, 2, range(2, 5))   # B..D: os 3 fretes por reputação (as colunas E e F são FÓRMULAS do rótulo da coluna A: ficam)
    aba.freeze_panes = 'A2'   # o arquivo real tem o congelamento solto em A20; no modelo o cabeçalho (linha 1) é que fica fixo
    _linha_de_fim(aba, 3, 6, NOTAS['frete-magalu'])


def _modelo_tiktok(livro):
    aba = livro.active
    _remover_linhas(aba, 3)
    _esvaziar(aba, 2, 2, [2])   # B: valor médio
    _linha_de_fim(aba, 3, 4, NOTAS['frete-tiktok-shop'])


def _modelo_amazon(livro):
    for aba in livro.worksheets:
        assert aba.cell(row=5, column=3).value == 79, 'a linha 5 da Amazon deixou de ser a 1ª faixa que depende de peso'
        _copiar_linha(aba, 5, 2, 5)   # exemplo: a 1ª faixa que depende de peso (0 a 0,249 kg, R$ 79 a R$ 99,99)
        _esvaziar(aba, 2, 2, [5])   # E: Valor
        _remover_linhas(aba, 3)
        if aba.auto_filter.ref:
            aba.auto_filter.ref = 'A1:E2'
        _linha_de_fim(aba, 3, 5, NOTAS['frete-amazon'])


def _modelo_busca_legal(livro):
    aba = livro.active
    ultima_coluna = aba.max_column
    colunas = {aba.cell(row=2, column=c).value: c for c in range(1, ultima_coluna + 1)}
    # Exemplo: a 1ª linha real com CST 00 (o caso comum), copiada para a linha 3 (o estilo da linha 3 fica).
    origem = next(n for n in range(3, aba.max_row + 1) if str(aba.cell(row=n, column=colunas['CST']).value) == '00')
    for coluna in range(1, ultima_coluna + 1):
        aba.cell(row=3, column=coluna).value = aba.cell(row=origem, column=coluna).value
    for nome in ('Cód Barras', 'NCM', 'CST'):
        assert isinstance(aba.cell(row=3, column=colunas[nome]).value, str), 'EAN, NCM e CST do exemplo têm de ser texto, como no arquivo real'
    _remover_linhas(aba, 4)
    _linha_de_fim(aba, 4, ultima_coluna, NOTAS['busca-legal'])
    aba.auto_filter.ref = f'A2:{aba.cell(row=2, column=ultima_coluna).column_letter}3'
    aba.sheet_view.selection[0].activeCell = 'A3'
    aba.sheet_view.selection[0].sqref = 'A3'


# (id do arquivo no catálogo) -> (arquivo real, o que esvaziar/ajustar)
MODELOS = {
    'frete-mercado-livre': ('Tabelas de Frete/Tabela_Frete_Mercado_Livre.xlsx', _modelo_ml),
    'frete-magalu': ('Tabelas de Frete/Tabela_Frete_Magalu.xlsx', _modelo_magalu),
    'frete-tiktok-shop': ('Tabelas de Frete/Tabela_Frete_Tiktok_Shop.xlsx', _modelo_tiktok),
    'frete-amazon': ('Tabelas de Frete/Tabela_Frete_Amazon.xlsx', _modelo_amazon),
    'busca-legal-magazine': ('Impostos Saida/TABELA DE IMPOSTOS DE SAIDA MAGAZINE.xlsx', _modelo_busca_legal),
    'busca-legal-samvale': ('Impostos Saida/TABELA DE IMPOSTOS DE SAIDA SAMVALE.xlsx', _modelo_busca_legal),
}


def main():
    PASTA_DOS_MODELOS.mkdir(parents=True, exist_ok=True)
    for id_do_arquivo, (caminho_real, ajustar) in MODELOS.items():
        origem = PASTA_DOS_ARQUIVOS_REAIS / caminho_real
        livro = load_workbook(origem)     # cópia em memória: o arquivo real não é tocado
        ajustar(livro)
        destino = PASTA_DOS_MODELOS / f'{id_do_arquivo}.xlsx'
        livro.save(destino)
        print(f'{id_do_arquivo}: {origem.name} -> {destino.relative_to(RAIZ)} ({destino.stat().st_size / 1024:.0f} KB)')


if __name__ == '__main__':
    main()
