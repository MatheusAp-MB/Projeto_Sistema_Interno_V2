# scripts_exploracao_ERP/duble_precificacao.py

# Função Objetivo: Duble de Precificação — passo a passo didático, só leitura, cobrindo
# os 6 marketplaces (Mercado Livre, Raia, Magalu, Shopee, TikTok, Amazon) e todas as
# variações internas reais de cada um (ML: Clássico/Premium; TikTok: Sem Afiliado/Com
# Afiliado; Amazon: DBA/FBA; Raia/Magalu/Shopee: nenhuma), nas 4 margens (Mínima/Padrão/
# Máxima/Competição), pros produtos de EANS_TESTE. Substitui o antigo
# duble_precificacao_ml.py (quebrado desde 16/08 — produto.pis_cofins removido — e cobria
# só ML/Clássico/Padrão). Não reimplementa nenhuma conta: instancia as 6 classes de
# fórmula de PRODUÇÃO (as mesmas usadas pelos calcular_grade_precificacao_*) e usa os
# métodos de auditoria delas (formula_abstrata/formula_preenchida/passos/entrada/
# intermediarios/saida) — o que este script mostra é EXATAMENTE o que o sistema real
# calculou, nunca uma conta paralela. Erro de assert numa combinação não derruba o
# script inteiro — é capturado, registrado no resumo, e o Duble segue pras próximas.
# Só leitura no banco, nenhuma escrita.

import os
import sys
from dataclasses import fields as dataclass_fields
from datetime import datetime


def _adicionar_raiz_do_projeto_ao_path():
    caminho_atual = os.path.dirname(os.path.abspath(__file__))
    while caminho_atual != os.path.dirname(caminho_atual):
        if os.path.exists(os.path.join(caminho_atual, 'manage.py')):
            sys.path.insert(0, caminho_atual)
            return
        caminho_atual = os.path.dirname(caminho_atual)
    raise RuntimeError('Não foi possível encontrar manage.py subindo a partir deste script.')


_adicionar_raiz_do_projeto_ao_path()

import django

os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'projeto_sistema_interno_mb_sv.settings')
django.setup()

from django.core.exceptions import ObjectDoesNotExist

from rich.console import Console
from rich.table import Table
from rich.panel import Panel

from produtos.models import Produto
from precificacao.models import (
    ConfiguracaoOperacional, FaixaArmazenagem, TabelaComissaoShopee, TabelaComissaoTiktok,
    FreteAmazon, TaxaKgAdicionalAmazon,
)

from mercado_livre.models import (
    FreteML, TipoDeAnuncioMercadoLivre, ConfiguracaoTipoAnuncioMercadoLivre,
)
from mercado_livre.funcoes_auxiliares.dimensoes_efetivas import resolver_dimensoes_efetivas
from precificacao.funcoes_auxiliares.mercado_livre.formula_precificacao import FormulaPrecificacao

from raia.models import ConfiguracaoRaia
from precificacao.funcoes_auxiliares.raia.formula_precificacao_raia import FormulaPrecificacaoRaia

from magalu.models import ConfiguracaoMagalu, FreteMagalu
from precificacao.funcoes_auxiliares.magalu.formula_precificacao_magalu import FormulaPrecificacaoMagalu

from shopee.models import ConfiguracaoShopee
from precificacao.funcoes_auxiliares.shopee.formula_precificacao_shopee import FormulaPrecificacaoShopee

from tiktok.models import ConfiguracaoTiktok, FreteTiktok
from precificacao.funcoes_auxiliares.tiktok.formula_precificacao_tiktok import FormulaPrecificacaoTiktok

from amazon.models import ConfiguracaoAmazon
from precificacao.funcoes_auxiliares.amazon.formula_precificacao_amazon import FormulaPrecificacaoAmazon

from produtos.funcoes_auxiliares.dimensoes_fisicas import resolver_dimensao_produto

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import CellIsRule
from openpyxl.worksheet.table import Table as TabelaExcel, TableStyleInfo


# ========== Configuração do Duble ==========

EANS_TESTE = [
    '7908050719121',  # pulverizador — produto de referência, usado desde o Duble antigo
    '7909436926904',  # CONJUNTO REP. MOTOR 1.0 CV 127V — 1 dos 2 SKUs com erro de assert em Raia/Magalu
    '7891988014072',  # PARTE APARELHO MECANICO/KIT DA BARRA... — o outro SKU problemático
]

_PASTA_ATUAL = os.path.dirname(os.path.abspath(__file__))
PASTA_SAIDAS = os.path.join(_PASTA_ATUAL, 'saidas')

console = Console(record=True)

resumo_geral = []  # 1 dict por combinação (produto, marketplace, variação, margem) — vira a tabela final
detalhamento_geral = []  # 1 dict por combinação, campo a campo (custo/créditos/FIXO/taxa/preço) — vira a aba "Detalhamento" do Excel
passos_geral = []  # 1 dict por PASSO de cada combinação resolvida — vira a aba "Passos" do Excel
entrada_xml_por_produto = {}  # 1 dict por produto — dado cru da nota fiscal (ICMS/ICMS ST/IPI/PIS/COFINS) — vira a aba "Impostos Entrada (XML)"
produtos_processados = {}  # 1 dict por produto — identificação básica, pra aba "Produtos" do Excel


# ========== Coleta do dado cru da nota fiscal (impostos_entrada) — pra aba "Impostos Entrada (XML)" ==========

# Função Objetivo: Lê produto.impostos_entrada (XML da nota, via Sysemp) e devolve TODOS os campos
# crus — ICMS, ICMS ST, IPI, PIS, COFINS, quantidade/custo da nota — sem processar nada. É o mesmo
# dado que hoje só dava pra conferir abrindo o Django Admin; aqui vira parte da auditoria, pronto
# pra cruzar com o crédito já calculado (aba "Créditos Fiscais") e confirmar se bate com a nota real.
# Nunca recalcula nada — só lê e organiza. None quando o produto não tem impostos_entrada sincronizado.
def _coletar_dados_entrada_xml(produto):
    try:
        ie = produto.impostos_entrada
    except ObjectDoesNotExist:
        return None

    return {
        'nr_nf': ie.nr_nf,
        'data_entrada_nota': ie.data_entrada_nota,
        'fornecedor': ie.fornecedor,
        'quantidade_nota': ie.quantidade_nota,
        'custo_total': ie.custo_total,
        'custo_unitario': ie.custo_unitario,
        # * [EXPLICAÇÃO] → mesmo critério de "é ST?" usado em creditos_fiscais_para_precificacao.py —
        #                  se a nota trouxe base ou valor de ICMS ST maior que zero, é regime ST.
        'regime_st': ie.icms_st.valor > 0 or ie.icms_st.base_calculo > 0,
        'icms_cst': ie.icms.cst_xml,
        'icms_base_calculo': ie.icms.base_calculo,
        'icms_aliquota': ie.icms.aliquota,
        'icms_reducao': ie.icms.reducao,
        'icms_valor': ie.icms.valor,
        'icms_st_base_calculo': ie.icms_st.base_calculo,
        'icms_st_aliquota': ie.icms_st.aliquota,
        'icms_st_reducao': ie.icms_st.reducao,
        'icms_st_valor': ie.icms_st.valor,
        'icms_st_aliquota_fcp': ie.icms_st.aliquota_fcp,
        'icms_st_valor_fcp': ie.icms_st.valor_fcp,
        'ipi_cst': ie.ipi.cst_xml,
        'ipi_base_calculo': ie.ipi.base_calculo,
        'ipi_aliquota': ie.ipi.aliquota,
        'ipi_valor': ie.ipi.valor,
        'pis_cst': ie.pis.cst_xml,
        'pis_base_calculo': ie.pis.base_calculo,
        'pis_aliquota': ie.pis.aliquota,
        'pis_reducao': ie.pis.reducao,
        'pis_valor': ie.pis.valor,
        'cofins_cst': ie.cofins.cst_xml,
        'cofins_base_calculo': ie.cofins.base_calculo,
        'cofins_aliquota': ie.cofins.aliquota,
        'cofins_reducao': ie.cofins.reducao,
        'cofins_valor': ie.cofins.valor,
    }


# ========== Helpers de impressão (Rich) — reaproveitados por todos os marketplaces ==========

# Função Objetivo: Imprime qualquer dataclass de auditoria (Entrada/Intermediarios/Saida)
# como tabela — 1 linha por campo, igual foi de fato usado/calculado naquela margem.
def _imprimir_dataclass(titulo, dado):
    tabela = Table(title=titulo)
    tabela.add_column('Campo', style='cyan', no_wrap=True)
    tabela.add_column('Valor', style='green')
    for campo in dataclass_fields(dado):
        tabela.add_row(campo.name, str(getattr(dado, campo.name)))
    console.print(tabela)


# Função Objetivo: Imprime passos() — o passo a passo real da fórmula, na ordem certa.
def _imprimir_passos(passos):
    tabela = Table(title='Passo a Passo (auditoria real da fórmula)', show_lines=True)
    tabela.add_column('#', justify='right', width=3)
    tabela.add_column('Rótulo', style='cyan')
    tabela.add_column('Fórmula (valores reais)')
    tabela.add_column('Resultado', style='bold green', justify='right')
    for passo in passos:
        tabela.add_row(str(passo['ordem']), passo['rotulo'], passo['formula'], str(passo['resultado']))
    console.print(tabela)


# Função Objetivo: Painel de destaque com o resultado final — preço, frete, margem obtida
# vs meta (com alerta visual se ficar abaixo), e margem exata antes do RoundUp90.
def _imprimir_resultado_final(formula, margem_alvo_percentual):
    margem_obtida = formula.saida.margem_percentual_obtida
    margem_ok = margem_obtida >= margem_alvo_percentual
    cor = 'green' if margem_ok else 'red'
    alerta = '' if margem_ok else '  ⚠ MARGEM ABAIXO DA META'

    linhas = [
        f'[bold {cor}]Preço Final: R$ {formula.saida.preco_final:.2f}[/bold {cor}]',
        f'Frete usado: R$ {formula.saida.frete_usado:.2f}',
        f'Margem obtida: {margem_obtida:.2f}%  (meta: {margem_alvo_percentual:.2f}%){alerta}',
        f'Margem exata (antes do RoundUp90): {formula.saida.margem_exata_percentual:.2f}%',
    ]
    preco_de_exibicao = getattr(formula.saida, 'preco_de_exibicao', None)
    if preco_de_exibicao is not None:
        linhas.append(f'Preço de exibição (decorativo): R$ {preco_de_exibicao:.2f}')

    console.print(Panel('\n'.join(linhas), border_style=cor))


# Função Objetivo: Guarda 1 linha da combinação (produto, marketplace, variação, margem)
# no resumo geral — sempre, mesmo em erro/sem-cálculo, pra aparecer na tabela final.
def _registrar_resumo(produto, marketplace, variacao, margem_chave, status, formula=None, margem_alvo=None, detalhe=''):
    resumo_geral.append({
        'produto': f'{produto.ean} — {produto.titulo[:40]}',
        'marketplace': marketplace,
        'variacao': variacao or '—',
        'margem': margem_chave,
        'status': status,
        'preco_final': f'R$ {formula.saida.preco_final:.2f}' if formula and formula.resolvida else '—',
        'margem_obtida': f'{formula.saida.margem_percentual_obtida:.2f}%' if formula and formula.resolvida else '—',
        'margem_meta': f'{margem_alvo:.2f}%' if margem_alvo is not None else '—',
        'detalhe': detalhe,
    })


# Função Objetivo: Guarda 1 linha campo-a-campo da combinação — mesmos nomes que
# CAMPOS_DIAGNOSTICO_FALHA já usa, então funciona igual em sucesso e em falha (o que não
# foi calculado fica None). Vira a aba "Detalhamento" do Excel — a auditoria campo a campo
# que antes só dava pra montar lendo o log inteiro na mão.
def _montar_linha_detalhamento(produto, marketplace, variacao_rotulo, margem_chave, margem_alvo_percentual, formula, status):
    preco_exato = formula.intermediarios.preco_exato_antes_arredondar if formula.intermediarios else None
    preco_final = formula.saida.preco_final if formula.saida else None
    frete_usado = formula.saida.frete_usado if formula.saida else None
    margem_obtida = formula.saida.margem_percentual_obtida if formula.saida else None

    detalhamento_geral.append({
        'ean': produto.ean, 'titulo': produto.titulo, 'marketplace': marketplace,
        'variacao': variacao_rotulo or '—', 'margem': margem_chave, 'margem_meta': margem_alvo_percentual,
        'status': status,
        'custo_final': getattr(formula, '_custo_final', None),
        'coleta': getattr(formula, '_coleta', None),
        'armazenagem': getattr(formula, '_armazenagem', None),
        'credito_icms_entrada': getattr(formula, '_credito_icms_entrada', None),
        'credito_pis': getattr(formula, '_credito_pis', None),
        'credito_cofins': getattr(formula, '_credito_cofins', None),
        'fixo': getattr(formula, '_fixo', None),
        'comissao_percentual': getattr(formula, '_comissao_percentual', None),
        'icms_saida_percentual': getattr(formula, '_icms_saida_percentual', None),
        'pis_saida_percentual': getattr(formula, '_pis_saida_percentual', None),
        'cofins_saida_percentual': getattr(formula, '_cofins_saida_percentual', None),
        'taxa_percentual': getattr(formula, '_taxa_percentual', None),
        'denominador': getattr(formula, '_denominador', None),
        'preco_exato': preco_exato, 'preco_final': preco_final,
        'frete_usado': frete_usado, 'margem_obtida': margem_obtida,
    })


# Função Objetivo: Guarda 1 linha por PASSO de formula.passos() — só chamada quando a
# fórmula resolveu (passos() precisa de entrada/intermediarios/saida já montados).
# Vira a aba "Passos" do Excel.
def _registrar_passos(produto, marketplace, variacao_rotulo, margem_chave, formula):
    for passo in formula.passos():
        passos_geral.append({
            'ean': produto.ean, 'titulo': produto.titulo, 'marketplace': marketplace,
            'variacao': variacao_rotulo or '—', 'margem': margem_chave,
            'ordem': passo['ordem'], 'rotulo': passo['rotulo'],
            'formula': passo['formula'], 'resultado': passo['resultado'],
        })


# * [EXPLICAÇÃO] → nomes internos (com "_" na frente) são IDÊNTICOS nas 6 classes de
#                  fórmula — conferido direto no código de calcular_fixo/montar_taxa_e_
#                  denominador de todas elas. Por isso 1 lista só, genérica, cobre todo
#                  mundo — o que existir na instância no momento da falha é mostrado.
CAMPOS_DIAGNOSTICO_FALHA = [
    '_custo_final', '_coleta', '_armazenagem',
    '_credito_icms_entrada', '_credito_pis', '_credito_cofins', '_fixo',
    '_comissao_percentual', '_icms_saida_percentual', '_pis_saida_percentual',
    '_cofins_saida_percentual', '_taxa_percentual', '_denominador',
]


# Função Objetivo: Mostra o que já tinha sido calculado até o ponto da falha — essencial
# pra diagnosticar ERRO DE ASSERT/SEM CÁLCULO, já que formula.entrada/intermediarios/
# saida continuam None quando a fórmula não resolve (só são montados em caso de sucesso).
def _imprimir_estado_interno_na_falha(formula):
    tabela = Table(title='Estado interno no momento da falha (o que já tinha sido calculado)')
    tabela.add_column('Campo', style='cyan', no_wrap=True)
    tabela.add_column('Valor', style='yellow')
    algum_campo = False
    for nome in CAMPOS_DIAGNOSTICO_FALHA:
        if hasattr(formula, nome):
            tabela.add_row(nome.lstrip('_'), str(getattr(formula, nome)))
            algum_campo = True
    if algum_campo:
        console.print(tabela)
    else:
        console.print('[dim]Nenhum campo intermediário chegou a ser calculado (falhou logo na busca de crédito fiscal).[/dim]')


# Função Objetivo: Roda 1 combinação (produto, marketplace, variação, margem) por
# completo — instancia a fórmula real (via construir_formula, sem calcular ainda),
# chama .calcular(), imprime Entrada/Intermediários/Passos/Saída/Resultado, e sempre
# registra no resumo geral. Erro de assert é capturado aqui — não derruba o script.
def _rodar_e_exibir(produto, marketplace, variacao_rotulo, margem_chave, margem_alvo_percentual, construir_formula):
    rotulo_variacao = f' — {variacao_rotulo}' if variacao_rotulo else ''
    console.rule(f'[bold]{marketplace}{rotulo_variacao} | margem {margem_chave} ({margem_alvo_percentual}%) | {produto.ean}[/bold]')

    # * [EXPLICAÇÃO] → formula é instanciada e SÓ DEPOIS .calcular() é chamado, separado —
    #                  se estivesse tudo numa linha só (formula = construir_formula().calcular()),
    #                  um AssertionError no meio do .calcular() faria a atribuição nunca
    #                  completar, e o except não teria acesso a NADA do estado já calculado.
    formula = construir_formula()
    try:
        formula.calcular()
    except AssertionError as e:
        console.print(Panel(f'[bold red]ERRO DE ASSERT[/bold red]\n{e}', border_style='red'))
        console.print()
        _imprimir_estado_interno_na_falha(formula)
        console.print()
        _registrar_resumo(produto, marketplace, variacao_rotulo, margem_chave, 'ERRO DE ASSERT', margem_alvo=margem_alvo_percentual, detalhe=str(e))
        _montar_linha_detalhamento(produto, marketplace, variacao_rotulo, margem_chave, margem_alvo_percentual, formula, 'ERRO DE ASSERT')
        return

    if not formula.resolvida:
        motivo = (
            'Sem crédito fiscal de entrada (impostos_entrada ausente/incompleto)'
            if getattr(formula, '_creditos', None) is None
            else 'Nenhuma faixa gerou solução consistente (meta de margem inatingível)'
        )
        console.print(Panel(f'[bold yellow]SEM CÁLCULO POSSÍVEL[/bold yellow]\n{motivo}', border_style='yellow'))
        console.print()
        _imprimir_estado_interno_na_falha(formula)
        console.print()
        _registrar_resumo(produto, marketplace, variacao_rotulo, margem_chave, 'SEM CÁLCULO', margem_alvo=margem_alvo_percentual, detalhe=motivo)
        _montar_linha_detalhamento(produto, marketplace, variacao_rotulo, margem_chave, margem_alvo_percentual, formula, 'SEM CÁLCULO')
        return

    _imprimir_dataclass('Entrada (dado cru — banco/planilha/config)', formula.entrada)
    console.print()
    _imprimir_dataclass('Intermediários (cada pedaço calculado)', formula.intermediarios)
    console.print()
    console.print(f'[dim]{formula.formula_abstrata()}[/dim]')
    console.print(f'[dim]{formula.formula_preenchida()}[/dim]')
    console.print()
    _imprimir_passos(formula.passos())
    console.print()
    _imprimir_dataclass('Saída (resultado)', formula.saida)
    _imprimir_resultado_final(formula, margem_alvo_percentual)
    console.print()

    status = 'OK' if formula.saida.margem_percentual_obtida >= margem_alvo_percentual else 'MARGEM ABAIXO DA META'
    _registrar_resumo(produto, marketplace, variacao_rotulo, margem_chave, status, formula=formula, margem_alvo=margem_alvo_percentual)
    _montar_linha_detalhamento(produto, marketplace, variacao_rotulo, margem_chave, margem_alvo_percentual, formula, status)
    _registrar_passos(produto, marketplace, variacao_rotulo, margem_chave, formula)


# ========== Carga de configuração (1x, fora do loop de produtos) ==========

console.print(Panel('[bold]Duble de Precificação — carregando configuração real do banco...[/bold]', border_style='blue'))

config_geral = ConfiguracaoOperacional.obter()
faixas_armazenagem = list(FaixaArmazenagem.objects.filter(ativo=True).order_by('ordem'))

TipoAnuncio = TipoDeAnuncioMercadoLivre.TipoAnuncio
configs_ml = {c.tipo_anuncio: c for c in ConfiguracaoTipoAnuncioMercadoLivre.objects.all()}
frete_todas_ml = list(FreteML.objects.all())

config_raia = ConfiguracaoRaia.obter()

config_magalu = ConfiguracaoMagalu.obter()
frete_todas_magalu = list(FreteMagalu.objects.all())

config_shopee = ConfiguracaoShopee.obter()
faixas_comissao_shopee = list(TabelaComissaoShopee.objects.all().order_by('preco_min'))

config_tiktok = ConfiguracaoTiktok.obter()
faixas_comissao_tiktok = list(TabelaComissaoTiktok.objects.all().order_by('preco_min'))
frete_todas_tiktok = list(FreteTiktok.objects.all())

config_amazon = ConfiguracaoAmazon.obter()
fretes_amazon = list(FreteAmazon.objects.all())
taxas_kg_adicional = list(TaxaKgAdicionalAmazon.objects.all())

# * [EXPLICAÇÃO] → Raia/Magalu/Shopee/TikTok/Amazon usam margem FIXA, hardcoded no
#                  próprio calcular_grade_precificacao_* (não vem de config) — copiado
#                  literal daqui pra manter o Duble fiel ao que roda de verdade.
MARGENS_PADRAO = [('minima', 10), ('padrao', 15), ('maxima', 20), ('competicao', 5)]


# Função Objetivo: ML é o único marketplace que lê as 4 margens de uma Configuração
# real por tipo de anúncio (Clássico/Premium têm margens próprias, não fixas).
def _margens_do_tipo_ml(config):
    return [
        ('minima', config.margem_minima), ('padrao', config.margem_padrao),
        ('maxima', config.margem_maxima), ('competicao', config.margem_competicao),
    ]


# ========== 1 função por marketplace — mesmo padrão dos calcular_grade_precificacao_* ==========

def _processar_mercado_livre(produto):
    dim = resolver_dimensoes_efetivas(produto, variacao=None)
    for tipo, tipo_rotulo in [(TipoAnuncio.CLASSICO, 'Clássico'), (TipoAnuncio.PREMIUM, 'Premium')]:
        config_tipo = configs_ml.get(tipo)
        if not config_tipo:
            console.print(f'[yellow]Sem ConfiguracaoTipoAnuncioMercadoLivre para {tipo_rotulo} — pulado.[/yellow]')
            continue
        for margem_chave, margem_valor in _margens_do_tipo_ml(config_tipo):
            _rodar_e_exibir(
                produto, 'Mercado Livre', tipo_rotulo, margem_chave, margem_valor,
                lambda config_tipo=config_tipo, margem_valor=margem_valor: FormulaPrecificacao(
                    produto=produto, dimensoes_efetivas=dim, config_tipo=config_tipo,
                    config_geral=config_geral, margem_alvo_percentual=margem_valor,
                    frete_todas=frete_todas_ml, faixas_armazenagem=faixas_armazenagem,
                )
            )


def _processar_raia(produto):
    for margem_chave, margem_valor in MARGENS_PADRAO:
        _rodar_e_exibir(
            produto, 'Raia', None, margem_chave, margem_valor,
            lambda margem_valor=margem_valor: FormulaPrecificacaoRaia(
                produto=produto, config_raia=config_raia, config_geral=config_geral,
                margem_alvo_percentual=margem_valor, faixas_armazenagem=faixas_armazenagem,
            )
        )


def _processar_magalu(produto):
    for margem_chave, margem_valor in MARGENS_PADRAO:
        _rodar_e_exibir(
            produto, 'Magalu', None, margem_chave, margem_valor,
            lambda margem_valor=margem_valor: FormulaPrecificacaoMagalu(
                produto=produto, config_magalu=config_magalu, config_geral=config_geral,
                margem_alvo_percentual=margem_valor, frete_todas=frete_todas_magalu,
                faixas_armazenagem=faixas_armazenagem,
            )
        )


def _processar_shopee(produto):
    for margem_chave, margem_valor in MARGENS_PADRAO:
        _rodar_e_exibir(
            produto, 'Shopee', None, margem_chave, margem_valor,
            lambda margem_valor=margem_valor: FormulaPrecificacaoShopee(
                produto=produto, config_shopee=config_shopee, config_geral=config_geral,
                margem_alvo_percentual=margem_valor, faixas_comissao=faixas_comissao_shopee,
                faixas_armazenagem=faixas_armazenagem,
            )
        )


def _processar_tiktok(produto):
    for tipo, tipo_rotulo in [('sem_afiliado', 'Sem Afiliado'), ('com_afiliado', 'Com Afiliado')]:
        for margem_chave, margem_valor in MARGENS_PADRAO:
            _rodar_e_exibir(
                produto, 'TikTok', tipo_rotulo, margem_chave, margem_valor,
                lambda tipo=tipo, margem_valor=margem_valor: FormulaPrecificacaoTiktok(
                    produto=produto, config_tiktok=config_tiktok, config_geral=config_geral,
                    margem_alvo_percentual=margem_valor, tipo=tipo,
                    faixas_comissao=faixas_comissao_tiktok, frete_todas=frete_todas_tiktok,
                    faixas_armazenagem=faixas_armazenagem,
                )
            )


def _processar_amazon(produto):
    for tipo, tipo_rotulo in [('dba', 'DBA'), ('fba', 'FBA')]:
        for margem_chave, margem_valor in MARGENS_PADRAO:
            _rodar_e_exibir(
                produto, 'Amazon', tipo_rotulo, margem_chave, margem_valor,
                lambda tipo=tipo, margem_valor=margem_valor: FormulaPrecificacaoAmazon(
                    produto=produto, config_amazon=config_amazon, config_geral=config_geral,
                    margem_alvo_percentual=margem_valor, tipo=tipo,
                    fretes_amazon=fretes_amazon, taxas_kg_adicional=taxas_kg_adicional,
                    faixas_armazenagem=faixas_armazenagem,
                )
            )


# ========== Geração do Excel — 8 abas, direto do que já foi coletado acima ==========

_FONTE = 'Arial'
_FUNDO_CABECALHO = PatternFill('solid', fgColor='1F3864')
_FONTE_CABECALHO = Font(name=_FONTE, bold=True, color='FFFFFF', size=10)
_FONTE_TITULO = Font(name=_FONTE, bold=True, size=14, color='1F3864')
_FONTE_SUBTITULO = Font(name=_FONTE, italic=True, size=10, color='595959')
_FONTE_CORPO = Font(name=_FONTE, size=10)
_FONTE_CORPO_NEGRITO = Font(name=_FONTE, size=10, bold=True)
_FUNDO_OK = PatternFill('solid', fgColor='C6EFCE')
_FONTE_OK = Font(name=_FONTE, size=10, color='006100')
_FUNDO_ERRO = PatternFill('solid', fgColor='FFC7CE')
_FONTE_ERRO = Font(name=_FONTE, size=10, color='9C0006', bold=True)
_FUNDO_SEM_CALCULO = PatternFill('solid', fgColor='FFEB9C')
_FONTE_SEM_CALCULO = Font(name=_FONTE, size=10, color='9C6500')
_FUNDO_SEM_DADO = PatternFill('solid', fgColor='F2F2F2')
_FONTE_SEM_DADO = Font(name=_FONTE, size=10, italic=True, color='7F7F7F')
_BORDA_FINA = Side(style='thin', color='D9D9D9')
_BORDA = Border(left=_BORDA_FINA, right=_BORDA_FINA, top=_BORDA_FINA, bottom=_BORDA_FINA)

MARGEM_LABEL = {'minima': 'Mínima', 'padrao': 'Padrão', 'maxima': 'Máxima', 'competicao': 'Competição'}


# Função Objetivo: Conta pura — Decimal/None pro tipo que o openpyxl grava melhor (float/None).
def _num(valor):
    return float(valor) if valor is not None else None


def _estilo_cabecalho(ws, linha=1, n_colunas=None):
    n_colunas = n_colunas or ws.max_column
    for c in range(1, n_colunas + 1):
        celula = ws.cell(row=linha, column=c)
        celula.font = _FONTE_CABECALHO
        celula.fill = _FUNDO_CABECALHO
        celula.alignment = Alignment(vertical='center', wrap_text=True, horizontal='center')
    ws.row_dimensions[linha].height = 30


def _autosize(ws, larguras):
    for i, largura in enumerate(larguras, start=1):
        ws.column_dimensions[get_column_letter(i)].width = largura


# Função Objetivo: Gera o log de saída "de verdade" — um .xlsx com 1 aba por tipo de dado,
# a partir só do que já foi coletado durante a execução (produtos_processados,
# entrada_xml_por_produto, detalhamento_geral, passos_geral). Segue a estrutura do mockup
# aprovado (10/09) — a diferença é que agora TODOS os dados são reais, inclusive a aba
# "Impostos Entrada (XML)" (antes só existia como EXEMPLO no mockup, porque o Duble ainda
# não buscava o dado — agora busca, via _coletar_dados_entrada_xml). Nunca recalcula nada,
# só organiza o que já foi apurado durante a execução acima.
def _gerar_excel(caminho):
    wb = Workbook()
    wb.remove(wb.active)

    # ---------- Capa ----------
    ws = wb.create_sheet('Capa')
    ws.sheet_view.showGridLines = False
    ws['B2'] = 'Duble de Precificação — Auditoria Completa'
    ws['B2'].font = _FONTE_TITULO
    ws['B3'] = (
        f'Gerado em {datetime.now().strftime("%d/%m/%Y %H:%M:%S")} — todos os dados vêm '
        f'direto do banco, sem precisar abrir o Django Admin.'
    )
    ws['B3'].font = _FONTE_SUBTITULO

    linhas = [
        '',
        'O que é este arquivo:',
        'Auditoria completa do Duble de Precificação — os 6 marketplaces (Mercado Livre, Raia, Magalu,',
        'Shopee, TikTok, Amazon), todas as variações internas e as 4 margens, pros produtos de EANS_TESTE.',
        'Cada aba cobre uma camada da conta: identificação do produto, dado cru da nota fiscal de entrada,',
        'créditos fiscais já calculados, impostos de saída, o cálculo campo a campo (inclusive combinações',
        'que falharam) e o passo a passo completo de auditoria de cada combinação que resolveu.',
        '',
        'Legenda de cores:',
    ]
    r = 5
    for linha in linhas:
        celula = ws.cell(row=r, column=2, value=linha)
        celula.font = _FONTE_CORPO_NEGRITO if linha.endswith(':') else _FONTE_CORPO
        r += 1

    legenda = [
        (_FUNDO_OK, _FONTE_OK, 'OK — margem obtida atingiu a meta'),
        (_FUNDO_ERRO, _FONTE_ERRO, 'ERRO DE ASSERT — RoundUp90 devolveu margem abaixo da meta'),
        (_FUNDO_SEM_CALCULO, _FONTE_SEM_CALCULO, 'SEM CÁLCULO — nenhuma faixa gerou solução consistente'),
        (_FUNDO_SEM_DADO, _FONTE_SEM_DADO, 'SEM DADO — produto sem nota de entrada sincronizada (Sysemp)'),
    ]
    r += 1
    for fundo, fonte, texto in legenda:
        ws.cell(row=r, column=2, value='   ').fill = fundo
        celula_texto = ws.cell(row=r, column=3, value=texto)
        celula_texto.font = fonte
        r += 1

    r += 2
    ws.cell(row=r, column=2, value='Abas deste arquivo:').font = _FONTE_CORPO_NEGRITO
    r += 1
    abas_desc = [
        ('Resumo', 'Visão geral — 1 linha por combinação (produto × marketplace × variação × margem), status colorido.'),
        ('Produtos', 'Identificação, dimensões e ICMS Saída SP dos produtos testados.'),
        ('Impostos Entrada (XML)', 'Dado cru da nota fiscal (ICMS, ICMS ST, IPI, PIS, COFINS, nº NF, quantidade) — direto do Sysemp.'),
        ('Créditos Fiscais', 'Créditos de entrada já calculados (ICMS/PIS/COFINS por unidade) + proporção crédito ICMS ÷ custo final.'),
        ('Impostos Saída', 'Os campos da Camada 1 (ICMS Saída SP, ICMS Saída Média, PIS, COFINS) por produto.'),
        ('Detalhamento', 'Todas as combinações, campo a campo (Custo Final, Coleta, Armazenagem, FIXO, Taxa, Denominador, Preço, Margem) — inclusive as que falharam.'),
        ('Passos', 'O passo a passo completo (fórmula + valor real) de cada combinação que resolveu com sucesso.'),
    ]
    for nome, desc in abas_desc:
        ws.cell(row=r, column=2, value=nome).font = _FONTE_CORPO_NEGRITO
        ws.cell(row=r, column=3, value=desc).font = _FONTE_CORPO
        r += 1

    _autosize(ws, [3, 24, 95])

    # ---------- Resumo ----------
    ws = wb.create_sheet('Resumo')
    colunas = ['EAN', 'Produto', 'Marketplace', 'Variação', 'Margem', 'Status',
               'Preço Final (R$)', 'Frete Usado (R$)', 'Margem Obtida (%)', 'Margem Meta (%)', 'Motivo (se falhou)']
    for c, cabecalho in enumerate(colunas, start=1):
        ws.cell(row=1, column=c, value=cabecalho)
    _estilo_cabecalho(ws)
    ws.freeze_panes = 'A2'

    linha_atual = 2
    for d in detalhamento_geral:
        motivo = ''
        if d['status'] == 'ERRO DE ASSERT':
            motivo = 'RoundUp90 devolveu margem abaixo da meta (ver aba Detalhamento)'
        elif d['status'] == 'SEM CÁLCULO':
            motivo = 'Nenhuma faixa gerou solução consistente'
        valores = [
            d['ean'], d['titulo'], d['marketplace'], d['variacao'],
            MARGEM_LABEL.get(d['margem'], d['margem']), d['status'],
            _num(d['preco_final']), _num(d['frete_usado']), _num(d['margem_obtida']), _num(d['margem_meta']), motivo,
        ]
        for c, valor in enumerate(valores, start=1):
            celula = ws.cell(row=linha_atual, column=c, value=valor)
            celula.font = _FONTE_CORPO
            celula.border = _BORDA
            if c in (7, 8):
                celula.number_format = 'R$ #,##0.00'
            if c in (9, 10):
                celula.number_format = '0.00"%"'
        linha_atual += 1

    ultima_linha = linha_atual - 1
    if ultima_linha >= 2:
        tabela = TabelaExcel(displayName='TabResumo', ref=f'A1:K{ultima_linha}')
        tabela.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
        ws.add_table(tabela)
        ws.conditional_formatting.add(f'F2:F{ultima_linha}', CellIsRule(operator='equal', formula=['"OK"'], fill=_FUNDO_OK, font=_FONTE_OK))
        ws.conditional_formatting.add(f'F2:F{ultima_linha}', CellIsRule(operator='equal', formula=['"ERRO DE ASSERT"'], fill=_FUNDO_ERRO, font=_FONTE_ERRO))
        ws.conditional_formatting.add(f'F2:F{ultima_linha}', CellIsRule(operator='equal', formula=['"SEM CÁLCULO"'], fill=_FUNDO_SEM_CALCULO, font=_FONTE_SEM_CALCULO))

    _autosize(ws, [14, 34, 16, 14, 12, 16, 14, 14, 15, 12, 48])

    # ---------- Produtos ----------
    ws = wb.create_sheet('Produtos')
    colunas = ['EAN', 'SKU', 'Cód. Fabricante', 'Título', 'Marca', 'Categoria', 'NCM',
               'Custo (R$)', 'ICMS Saída SP (%)', 'Altura (cm)', 'Largura (cm)', 'Comprimento (cm)', 'Peso (kg)']
    for c, cabecalho in enumerate(colunas, start=1):
        ws.cell(row=1, column=c, value=cabecalho)
    _estilo_cabecalho(ws)
    ws.freeze_panes = 'A2'

    linha_atual = 2
    for ean, p in produtos_processados.items():
        valores = [
            p['ean'], p['sku'], p['cod_fabricante'], p['titulo'], p['marca'], p['categoria'], p['ncm'],
            _num(p['custo']), _num(p['icms_saida_sp']), _num(p['altura']), _num(p['largura']),
            _num(p['comprimento']), _num(p['peso']),
        ]
        for c, valor in enumerate(valores, start=1):
            celula = ws.cell(row=linha_atual, column=c, value=valor)
            celula.font = _FONTE_CORPO
            celula.border = _BORDA
            if c == 8:
                celula.number_format = 'R$ #,##0.00'
            if c == 9:
                celula.number_format = '0.00"%"'
            if c in (10, 11, 12, 13):
                celula.number_format = '#,##0.000'
        linha_atual += 1

    _autosize(ws, [14, 20, 16, 46, 14, 18, 12, 14, 15, 12, 12, 15, 10])

    # ---------- Impostos Entrada (XML) ----------
    ws = wb.create_sheet('Impostos Entrada (XML)')
    colunas = [
        'EAN', 'Produto', 'Nº NF', 'Data Entrada', 'Fornecedor', 'Qtde Nota',
        'Custo Total Nota (R$)', 'Custo Unitário Nota (R$)', 'Regime ST?',
        'ICMS CST', 'ICMS Base Cálc. (R$)', 'ICMS Alíquota (%)', 'ICMS Redução (%)', 'ICMS Valor (R$)',
        'ICMS ST Base Cálc. (R$)', 'ICMS ST Alíquota (%)', 'ICMS ST Redução (%)', 'ICMS ST Valor (R$)',
        'ICMS ST Alíq. FCP (%)', 'ICMS ST Valor FCP (R$)',
        'IPI CST', 'IPI Base Cálc. (R$)', 'IPI Alíquota (%)', 'IPI Valor (R$)',
        'PIS CST', 'PIS Base Cálc. (R$)', 'PIS Alíquota (%)', 'PIS Redução (%)', 'PIS Valor (R$)',
        'COFINS CST', 'COFINS Base Cálc. (R$)', 'COFINS Alíquota (%)', 'COFINS Redução (%)', 'COFINS Valor (R$)',
    ]
    for c, cabecalho in enumerate(colunas, start=1):
        ws.cell(row=1, column=c, value=cabecalho)
    _estilo_cabecalho(ws)
    ws.freeze_panes = 'A2'

    colunas_moeda = {7, 8, 11, 14, 15, 18, 20, 22, 24, 26, 29, 31, 34}
    colunas_percentual = {12, 13, 16, 17, 19, 23, 27, 28, 32, 33}

    linha_atual = 2
    for ean, p in produtos_processados.items():
        entrada = entrada_xml_por_produto.get(ean)
        if entrada is None:
            for c in (1, 2, 3):
                celula = ws.cell(
                    row=linha_atual, column=c,
                    value=[ean, p['titulo'], 'SEM DADO — produto sem nota de entrada sincronizada (Sysemp)'][c - 1],
                )
                celula.font = _FONTE_SEM_DADO
                celula.fill = _FUNDO_SEM_DADO
            ws.merge_cells(start_row=linha_atual, start_column=3, end_row=linha_atual, end_column=34)
            linha_atual += 1
            continue

        valores = [
            ean, p['titulo'], entrada['nr_nf'], entrada['data_entrada_nota'], entrada['fornecedor'],
            entrada['quantidade_nota'], _num(entrada['custo_total']), _num(entrada['custo_unitario']),
            'Sim' if entrada['regime_st'] else 'Não',
            entrada['icms_cst'], _num(entrada['icms_base_calculo']), _num(entrada['icms_aliquota']),
            _num(entrada['icms_reducao']), _num(entrada['icms_valor']),
            _num(entrada['icms_st_base_calculo']), _num(entrada['icms_st_aliquota']),
            _num(entrada['icms_st_reducao']), _num(entrada['icms_st_valor']),
            _num(entrada['icms_st_aliquota_fcp']), _num(entrada['icms_st_valor_fcp']),
            entrada['ipi_cst'], _num(entrada['ipi_base_calculo']), _num(entrada['ipi_aliquota']), _num(entrada['ipi_valor']),
            entrada['pis_cst'], _num(entrada['pis_base_calculo']), _num(entrada['pis_aliquota']),
            _num(entrada['pis_reducao']), _num(entrada['pis_valor']),
            entrada['cofins_cst'], _num(entrada['cofins_base_calculo']), _num(entrada['cofins_aliquota']),
            _num(entrada['cofins_reducao']), _num(entrada['cofins_valor']),
        ]
        for c, valor in enumerate(valores, start=1):
            celula = ws.cell(row=linha_atual, column=c, value=valor)
            celula.font = _FONTE_CORPO
            celula.border = _BORDA
            if c in colunas_moeda:
                celula.number_format = 'R$ #,##0.00'
            if c in colunas_percentual:
                celula.number_format = '0.00"%"'
            if c == 4 and valor is not None:
                celula.number_format = 'DD/MM/YYYY'
        linha_atual += 1

    nota_linha = linha_atual + 1
    ws.cell(
        row=nota_linha, column=1,
        value='"Regime ST?" usa o mesmo critério de creditos_fiscais_para_precificacao.py: ICMS ST > 0 (valor ou base) na nota de entrada.',
    ).font = _FONTE_SUBTITULO
    ws.merge_cells(start_row=nota_linha, start_column=1, end_row=nota_linha, end_column=34)

    _autosize(ws, [
        14, 30, 14, 12, 26, 9, 15, 16, 10,
        9, 14, 12, 12, 12,
        14, 12, 12, 12, 12, 13,
        9, 14, 12, 12,
        9, 14, 12, 12, 12,
        10, 14, 12, 12, 12,
    ])

    # ---------- Créditos Fiscais ----------
    ws = wb.create_sheet('Créditos Fiscais')
    colunas = ['EAN', 'Produto', 'Marketplace/Margem de referência', 'Custo Final (R$)', 'Coleta (R$)',
               'Armazenagem (R$)', 'Crédito ICMS/unid (R$)', 'Crédito PIS/unid (R$)', 'Crédito COFINS/unid (R$)',
               'FIXO (R$)', 'Créd. ICMS ÷ Custo Final']
    for c, cabecalho in enumerate(colunas, start=1):
        ws.cell(row=1, column=c, value=cabecalho)
    _estilo_cabecalho(ws)
    ws.freeze_panes = 'A2'

    linha_atual = 2
    primeira_linha_dados = linha_atual
    for ean, p in produtos_processados.items():
        referencia = next((d for d in detalhamento_geral if d['ean'] == ean), None)
        rotulo_referencia = (
            f"{referencia['marketplace']} — {MARGEM_LABEL.get(referencia['margem'], referencia['margem'])}"
            if referencia else '—'
        )
        custo_final = _num(referencia['custo_final']) if referencia else None
        coleta = _num(referencia['coleta']) if referencia else None
        armazenagem = _num(referencia['armazenagem']) if referencia else None
        credito_icms = _num(referencia['credito_icms_entrada']) if referencia else None
        credito_pis = _num(referencia['credito_pis']) if referencia else None
        credito_cofins = _num(referencia['credito_cofins']) if referencia else None
        fixo = _num(referencia['fixo']) if referencia else None

        valores = [ean, p['titulo'], rotulo_referencia, custo_final, coleta, armazenagem,
                   credito_icms, credito_pis, credito_cofins, fixo]
        for c, valor in enumerate(valores, start=1):
            celula = ws.cell(row=linha_atual, column=c, value=valor)
            celula.font = _FONTE_CORPO
            celula.border = _BORDA
            if c in (4, 5, 6, 7, 8, 9, 10):
                celula.number_format = 'R$ #,##0.00'
        if custo_final:
            razao = ws.cell(row=linha_atual, column=11, value=f'=G{linha_atual}/D{linha_atual}')
            razao.number_format = '0.0"x"'
            razao.font = _FONTE_CORPO
            razao.border = _BORDA
        if fixo is not None and fixo < 0:
            ws.cell(row=linha_atual, column=10).font = Font(name=_FONTE, size=10, bold=True, color='C00000')
        linha_atual += 1

    ultima_linha = linha_atual - 1
    if ultima_linha >= primeira_linha_dados:
        ws.conditional_formatting.add(
            f'K{primeira_linha_dados}:K{ultima_linha}',
            CellIsRule(operator='greaterThan', formula=['1'], fill=_FUNDO_ERRO, font=_FONTE_ERRO),
        )
    nota_linha = ultima_linha + 2
    ws.cell(
        row=nota_linha, column=1,
        value=(
            'Coluna K em vermelho quando o crédito de ICMS de entrada sozinho já é maior que o custo final — sinal de '
            'alerta. "Marketplace/Margem de referência" é a 1ª combinação processada pra esse produto — Coleta/Armazenagem '
            'variam por marketplace, mas os créditos fiscais vêm da nota de entrada e não mudam (ver aba Detalhamento pra '
            'todas as combinações).'
        ),
    ).font = _FONTE_SUBTITULO
    ws.merge_cells(start_row=nota_linha, start_column=1, end_row=nota_linha, end_column=11)

    _autosize(ws, [14, 40, 26, 15, 12, 14, 16, 15, 17, 13, 16])

    # ---------- Impostos Saída ----------
    ws = wb.create_sheet('Impostos Saída')
    colunas = ['EAN', 'Produto', 'ICMS Saída SP (%)', 'ICMS Saída Média (%)', 'PIS (%)', 'COFINS (%)']
    for c, cabecalho in enumerate(colunas, start=1):
        ws.cell(row=1, column=c, value=cabecalho)
    _estilo_cabecalho(ws)
    ws.freeze_panes = 'A2'

    linha_atual = 2
    for ean, p in produtos_processados.items():
        referencia = next((d for d in detalhamento_geral if d['ean'] == ean), None)
        icms_saida_media = _num(referencia['icms_saida_percentual']) if referencia else None
        pis_saida = _num(referencia['pis_saida_percentual']) if referencia else None
        cofins_saida = _num(referencia['cofins_saida_percentual']) if referencia else None
        icms_saida_sp = _num(p['icms_saida_sp'])

        valores = [ean, p['titulo'], icms_saida_sp, icms_saida_media, pis_saida, cofins_saida]
        for c, valor in enumerate(valores, start=1):
            celula = ws.cell(row=linha_atual, column=c, value=valor)
            celula.font = _FONTE_CORPO
            celula.border = _BORDA
            if c in (3, 4, 5, 6):
                celula.number_format = '0.00"%"'
        if icms_saida_media == 0 and pis_saida == 0 and cofins_saida == 0:
            for c in (3, 4, 5, 6):
                ws.cell(row=linha_atual, column=c).font = Font(name=_FONTE, size=10, italic=True, color='9C6500')
        linha_atual += 1

    nota_linha = linha_atual + 1
    ws.cell(
        row=nota_linha, column=1,
        value='Linhas com ICMS Saída Média/PIS/COFINS zerados: produto não foi alcançado pela Camada 1 (Busca Legal).',
    ).font = _FONTE_SUBTITULO
    ws.merge_cells(start_row=nota_linha, start_column=1, end_row=nota_linha, end_column=6)

    _autosize(ws, [14, 40, 16, 18, 10, 12])

    # ---------- Detalhamento ----------
    ws = wb.create_sheet('Detalhamento')
    colunas = ['EAN', 'Produto', 'Marketplace', 'Variação', 'Margem', 'Status',
               'Custo Final (R$)', 'Coleta (R$)', 'Armazenagem (R$)',
               'Créd. ICMS (R$)', 'Créd. PIS (R$)', 'Créd. COFINS (R$)', 'FIXO (R$)',
               'Comissão (%)', 'ICMS Saída (%)', 'PIS Saída (%)', 'COFINS Saída (%)',
               'Taxa Total (fração)', 'Denominador', 'Preço Exato (R$)', 'Preço Final (R$)',
               'Frete Usado (R$)', 'Margem Obtida (%)', 'Margem Meta (%)']
    for c, cabecalho in enumerate(colunas, start=1):
        ws.cell(row=1, column=c, value=cabecalho)
    _estilo_cabecalho(ws)
    ws.freeze_panes = 'A2'

    linha_atual = 2
    primeira_linha_dados = linha_atual
    for d in detalhamento_geral:
        valores = [
            d['ean'], d['titulo'], d['marketplace'], d['variacao'],
            MARGEM_LABEL.get(d['margem'], d['margem']), d['status'],
            _num(d['custo_final']), _num(d['coleta']), _num(d['armazenagem']),
            _num(d['credito_icms_entrada']), _num(d['credito_pis']), _num(d['credito_cofins']), _num(d['fixo']),
            _num(d['comissao_percentual']), _num(d['icms_saida_percentual']), _num(d['pis_saida_percentual']),
            _num(d['cofins_saida_percentual']), _num(d['taxa_percentual']), _num(d['denominador']),
            _num(d['preco_exato']), _num(d['preco_final']), _num(d['frete_usado']),
            _num(d['margem_obtida']), _num(d['margem_meta']),
        ]
        for c, valor in enumerate(valores, start=1):
            celula = ws.cell(row=linha_atual, column=c, value=valor)
            celula.font = _FONTE_CORPO
            celula.border = _BORDA
            if c in (7, 8, 9, 10, 11, 12, 13, 20, 21, 22):
                celula.number_format = 'R$ #,##0.00'
            if c in (14, 15, 16, 17, 23, 24):
                celula.number_format = '0.00"%"'
            if c in (18, 19):
                celula.number_format = '0.0000'
        linha_atual += 1

    ultima_linha = linha_atual - 1
    if ultima_linha >= primeira_linha_dados:
        tabela = TabelaExcel(displayName='TabDetalhamento', ref=f'A1:X{ultima_linha}')
        tabela.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
        ws.add_table(tabela)
        ws.conditional_formatting.add(f'F{primeira_linha_dados}:F{ultima_linha}', CellIsRule(operator='equal', formula=['"OK"'], fill=_FUNDO_OK, font=_FONTE_OK))
        ws.conditional_formatting.add(f'F{primeira_linha_dados}:F{ultima_linha}', CellIsRule(operator='equal', formula=['"ERRO DE ASSERT"'], fill=_FUNDO_ERRO, font=_FONTE_ERRO))
        ws.conditional_formatting.add(f'F{primeira_linha_dados}:F{ultima_linha}', CellIsRule(operator='equal', formula=['"SEM CÁLCULO"'], fill=_FUNDO_SEM_CALCULO, font=_FONTE_SEM_CALCULO))
        ws.conditional_formatting.add(f'M{primeira_linha_dados}:M{ultima_linha}', CellIsRule(operator='lessThan', formula=['0'], fill=_FUNDO_ERRO, font=_FONTE_ERRO))

    _autosize(ws, [14, 30, 14, 13, 11, 16, 13, 10, 12, 12, 11, 13, 13, 11, 12, 11, 13, 13, 11, 14, 13, 13, 14, 11])

    # ---------- Passos ----------
    ws = wb.create_sheet('Passos')
    colunas = ['EAN', 'Produto', 'Marketplace', 'Variação', 'Margem', '#', 'Rótulo', 'Fórmula (valores reais)', 'Resultado']
    for c, cabecalho in enumerate(colunas, start=1):
        ws.cell(row=1, column=c, value=cabecalho)
    _estilo_cabecalho(ws)
    ws.freeze_panes = 'A2'

    linha_atual = 2
    primeira_linha_dados = linha_atual
    for passo in passos_geral:
        try:
            resultado_excel = float(passo['resultado'])
        except (TypeError, ValueError):
            resultado_excel = str(passo['resultado'])
        valores = [
            passo['ean'], passo['titulo'], passo['marketplace'], passo['variacao'],
            MARGEM_LABEL.get(passo['margem'], passo['margem']), passo['ordem'], passo['rotulo'],
            passo['formula'], resultado_excel,
        ]
        for c, valor in enumerate(valores, start=1):
            celula = ws.cell(row=linha_atual, column=c, value=valor)
            celula.font = _FONTE_CORPO
            celula.border = _BORDA
            if c == 9 and isinstance(valor, float):
                celula.number_format = 'R$ #,##0.00'
        linha_atual += 1

    ultima_linha = linha_atual - 1
    if ultima_linha >= primeira_linha_dados:
        tabela = TabelaExcel(displayName='TabPassos', ref=f'A1:I{ultima_linha}')
        tabela.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
        ws.add_table(tabela)

    _autosize(ws, [14, 30, 14, 13, 11, 5, 30, 55, 16])

    ordem = ['Capa', 'Resumo', 'Produtos', 'Impostos Entrada (XML)', 'Créditos Fiscais',
             'Impostos Saída', 'Detalhamento', 'Passos']
    wb._sheets = [wb[nome] for nome in ordem]
    wb.active = 0

    wb.save(caminho)


# ========== Execução — N produtos × 6 marketplaces × variações × 4 margens ==========

for ean in EANS_TESTE:
    try:
        produto = Produto.objects.get(ean=ean)
    except Produto.DoesNotExist:
        console.print(Panel(f'[bold red]Produto EAN {ean} não encontrado no banco — pulado.[/bold red]', border_style='red'))
        continue

    console.rule(f'[bold blue]PRODUTO: {produto.ean} — {produto.titulo}[/bold blue]', style='blue')
    console.print(f'SKU: {produto.sku}  |  Marca: {produto.marca}  |  Custo: R$ {produto.custo:.2f}')
    console.print()

    entrada_xml_por_produto[ean] = _coletar_dados_entrada_xml(produto)
    altura, largura, comprimento, peso = resolver_dimensao_produto(produto)
    produtos_processados[ean] = {
        'ean': produto.ean, 'sku': produto.sku, 'cod_fabricante': produto.cod_fabricante,
        'titulo': produto.titulo, 'marca': produto.marca, 'categoria': produto.categoria,
        'ncm': produto.ncm, 'custo': produto.custo,
        'altura': altura, 'largura': largura, 'comprimento': comprimento, 'peso': peso,
        'icms_saida_sp': produto.icms_saida_sp,
    }

    _processar_mercado_livre(produto)
    _processar_raia(produto)
    _processar_magalu(produto)
    _processar_shopee(produto)
    _processar_tiktok(produto)
    _processar_amazon(produto)


# ========== Resumo geral — 1 linha por combinação, pra ver o panorama de uma vez ==========

console.rule('[bold]RESUMO GERAL[/bold]')

tabela_resumo = Table(title='Todas as combinações rodadas neste Duble')
tabela_resumo.add_column('Produto', style='cyan', max_width=30)
tabela_resumo.add_column('Marketplace')
tabela_resumo.add_column('Variação')
tabela_resumo.add_column('Margem')
tabela_resumo.add_column('Status')
tabela_resumo.add_column('Preço Final', justify='right')
tabela_resumo.add_column('Margem Obtida', justify='right')
tabela_resumo.add_column('Meta', justify='right')

for linha in resumo_geral:
    cor_status = {
        'OK': 'green', 'ERRO DE ASSERT': 'bold red',
        'SEM CÁLCULO': 'yellow', 'MARGEM ABAIXO DA META': 'bold red',
    }.get(linha['status'], 'white')
    tabela_resumo.add_row(
        linha['produto'], linha['marketplace'], linha['variacao'], linha['margem'],
        f'[{cor_status}]{linha["status"]}[/{cor_status}]',
        linha['preco_final'], linha['margem_obtida'], linha['margem_meta'],
    )

console.print(tabela_resumo)

total = len(resumo_geral)
qtd_ok = sum(1 for l in resumo_geral if l['status'] == 'OK')
qtd_erro = sum(1 for l in resumo_geral if l['status'] == 'ERRO DE ASSERT')
qtd_sem_calculo = sum(1 for l in resumo_geral if l['status'] == 'SEM CÁLCULO')
qtd_margem_abaixo = sum(1 for l in resumo_geral if l['status'] == 'MARGEM ABAIXO DA META')

console.print(
    f'\nTotal de combinações: {total}  |  OK: {qtd_ok}  |  '
    f'Erros de assert: {qtd_erro}  |  Sem cálculo: {qtd_sem_calculo}  |  '
    f'Margem abaixo da meta (sem crashar): {qtd_margem_abaixo}'
)

os.makedirs(PASTA_SAIDAS, exist_ok=True)
carimbo = datetime.now().strftime('%Y%m%d_%H%M%S')
CAMINHO_LOG = os.path.join(PASTA_SAIDAS, f'duble_precificacao_{carimbo}.txt')
console.print(f'\n[dim]Log completo salvo em: {CAMINHO_LOG}[/dim]')
console.save_text(CAMINHO_LOG)

CAMINHO_EXCEL = os.path.join(PASTA_SAIDAS, f'duble_precificacao_{carimbo}.xlsx')
_gerar_excel(CAMINHO_EXCEL)
console.print(f'[dim]Auditoria completa (Excel, 8 abas) salva em: {CAMINHO_EXCEL}[/dim]')