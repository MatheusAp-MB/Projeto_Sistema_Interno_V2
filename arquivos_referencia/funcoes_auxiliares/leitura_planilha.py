# arquivos_referencia/funcoes_auxiliares/leitura_planilha.py

# Função Objetivo: Abre um .xlsx ENVIADO PELO USUÁRIO (só na memória — nada é gravado em disco) e entrega
# as linhas de cada aba do mesmo jeito que os importadores do "popular banco" leem (openpyxl em modo
# leitura, com os valores já calculados). Quem audita recebe só as linhas; aqui ficam as proteções
# do arquivo em si.
#
# * [EXPLICAÇÃO] → Proteções (arquivo vem da mão de uma pessoa, não de dentro do sistema):
#   - tamanho máximo do arquivo e do conteúdo descompactado (um .xlsx é um zip; um zip pode "explodir");
#   - formato: só .xlsx de verdade (o .xls antigo, e o .xlsx protegido por senha, não abrem aqui);
#   - quantidade máxima de linhas por aba (tabela de frete não tem milhares de linhas).
# * [IMPORTANTE] → Fórmula SEM resultado salvo (arquivo gerado por programa e nunca aberto no Excel) é lida
#   pelo openpyxl como célula VAZIA. O importador também lê assim, em silêncio — por isso aqui a gente
#   detecta e devolve a lista, para a auditoria apontar como erro.

import io
import zipfile

import openpyxl

TAMANHO_MAXIMO_DO_ARQUIVO = 5 * 1024 * 1024
LIMITE_DE_BYTES_DESCOMPACTADOS = 100 * 1024 * 1024
LIMITE_DE_PARTES_NO_ZIP = 300
LIMITE_DE_LINHAS_POR_ABA = 5000

# * [EXPLICAÇÃO] → Assinatura do formato antigo do Excel (.xls) e também do .xlsx protegido por senha.
ASSINATURA_OLE2 = b'\xd0\xcf\x11\xe0'


class ErroDeLeitura(Exception):

    # Função Objetivo: Problema que impede de ler o arquivo, já em português e com o que fazer.
    def __init__(self, o_que, como_corrigir=''):
        super().__init__(o_que)
        self.o_que = o_que
        self.como_corrigir = como_corrigir


class LeituraDaPlanilha:

    # Função Objetivo: Guarda as abas lidas, qual é a aba ativa e as fórmulas sem resultado.
    def __init__(self, abas, nome_da_aba_ativa, formulas_sem_resultado):
        self.abas = abas
        self.nome_da_aba_ativa = nome_da_aba_ativa
        self.formulas_sem_resultado = formulas_sem_resultado

    @property
    def nomes_das_abas(self):
        return list(self.abas.keys())

    def linhas(self, nome_da_aba):
        return self.abas[nome_da_aba]


# Função Objetivo: Confere o arquivo ANTES de abrir (tamanho, formato, zip seguro).
def _conferir_o_arquivo(conteudo):
    if not conteudo:
        raise ErroDeLeitura('O arquivo está vazio.', 'Escolha o arquivo certo e tente de novo.')
    if len(conteudo) > TAMANHO_MAXIMO_DO_ARQUIVO:
        raise ErroDeLeitura(
            'O arquivo é grande demais para uma tabela de referência (mais de 5 MB).',
            'Confira se é o arquivo certo. Se for, avise o desenvolvedor.',
        )
    if conteudo[:4] == ASSINATURA_OLE2:
        raise ErroDeLeitura(
            'Este arquivo está no formato antigo do Excel (.xls) ou protegido por senha.',
            'Abra no Excel, tire a senha se houver, e use “Salvar como” → “Pasta de Trabalho do Excel (.xlsx)”.',
        )
    if not zipfile.is_zipfile(io.BytesIO(conteudo)):
        raise ErroDeLeitura(
            'Este arquivo não é uma planilha do Excel (.xlsx).',
            'Abra no Excel e use “Salvar como” → “Pasta de Trabalho do Excel (.xlsx)”.',
        )
    with zipfile.ZipFile(io.BytesIO(conteudo)) as zip_:
        partes = zip_.infolist()
        if len(partes) > LIMITE_DE_PARTES_NO_ZIP or sum(p.file_size for p in partes) > LIMITE_DE_BYTES_DESCOMPACTADOS:
            raise ErroDeLeitura(
                'O conteúdo deste arquivo é grande demais para uma tabela de referência.',
                'Confira se é o arquivo certo. Se for, avise o desenvolvedor.',
            )


# Função Objetivo: Lê todas as abas (valores já calculados), do mesmo jeito dos importadores.
def _ler_valores(conteudo, limite_de_linhas):
    try:
        pasta = openpyxl.load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True)
    except Exception:
        raise ErroDeLeitura(
            'Não consegui abrir este arquivo como planilha do Excel.',
            'Abra no Excel, salve de novo como .xlsx e tente outra vez.',
        )
    try:
        abas = {}
        for aba in pasta.worksheets:
            linhas = []
            for linha in aba.iter_rows(values_only=True):
                linhas.append(tuple(linha))
                if len(linhas) > limite_de_linhas:
                    raise ErroDeLeitura(
                        f'A aba “{aba.title}” tem mais de {limite_de_linhas} linhas — não parece uma tabela de referência.',
                        'Confira se é o arquivo certo.',
                    )
            abas[aba.title] = linhas
        ativa = pasta.active.title if pasta.active is not None else (next(iter(abas), None))
    finally:
        pasta.close()
    if not abas:
        raise ErroDeLeitura('O arquivo não tem nenhuma aba.', 'Confira se é o arquivo certo.')
    return abas, ativa


# Função Objetivo: Acha as células que têm FÓRMULA mas nenhum resultado salvo (o importador leria como vazia).
def _achar_formulas_sem_resultado(conteudo, abas_com_valores):
    achadas = {}
    try:
        pasta = openpyxl.load_workbook(io.BytesIO(conteudo), read_only=True, data_only=False)
    except Exception:
        return achadas
    try:
        for aba in pasta.worksheets:
            valores = abas_com_valores.get(aba.title)
            if valores is None:
                continue
            for n_linha, linha in enumerate(aba.iter_rows(values_only=True)):
                if n_linha >= len(valores):
                    break
                linha_de_valores = valores[n_linha]
                for n_coluna, bruto in enumerate(linha):
                    if isinstance(bruto, str) and bruto.startswith('='):
                        resultado = linha_de_valores[n_coluna] if n_coluna < len(linha_de_valores) else None
                        if resultado is None:
                            achadas.setdefault(aba.title, []).append((n_linha + 1, n_coluna))
    finally:
        pasta.close()
    return achadas


# Função Objetivo: Ponto de entrada — do conteúdo do arquivo (bytes) para a leitura pronta.
# Levanta ErroDeLeitura (com texto pronto para a tela) se o arquivo não puder ser lido.
def ler_planilha(conteudo, limite_de_linhas=LIMITE_DE_LINHAS_POR_ABA):
    _conferir_o_arquivo(conteudo)
    abas, ativa = _ler_valores(conteudo, limite_de_linhas)
    formulas = _achar_formulas_sem_resultado(conteudo, abas)
    return LeituraDaPlanilha(abas, ativa, formulas)
