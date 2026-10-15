# arquivos_referencia/funcoes_auxiliares/auditoria_busca_legal.py

# Função Objetivo: Auditoria das 2 planilhas Busca Legal ("TABELA SAIDA POR UF MAGAZINE/SAMVALE.xlsx"): confere se o
# sistema vai ler a planilha corretamente ANTES de ela ir para o Drive. Espelha os importadores do popular banco
# (preencher_CST_produtos, importar_icms_saida_por_ncm_cst_origem, importar_pis_cofins_saida_por_ncm_cst): o que eles
# leriam errado ou ignorariam em silêncio, a auditoria aponta antes.
#
# * [DECISÃO, 14/10/2026] → ERRO impede o envio; AVISO não impede. "EAN que não existe no sistema" e "NCM rejeitado"
#   são só AVISO (decisão de Matheus).
# * [DECISÃO, 15/10/2026] → Fonte única das regras: nomes de coluna, as 27 UFs e a conversão das células vêm DIRETO dos
#   módulos dos importadores — e o "NCM rejeitado" é decidido pelas MESMAS classes de agrupamento que o importador usa
#   (AgrupadorIcmsPorNcm e AgrupadorPisCofinsPorNcmCst), alimentadas com as linhas da planilha enviada. Se o importador
#   mudar uma regra, a auditoria muda junto.
# * [EXPLICAÇÃO] → Como o importador lê a planilha (e por que cada regra existe):
#   - só a aba ATIVA (a que estava aberta ao salvar); linha 1 = título dos grupos (descartada), linha 2 = cabeçalho,
#     dados a partir da linha 3; linha 100% vazia é pulada;
#   - as colunas são lidas pelo NOME do cabeçalho, igualzinho ("Cód Barras", "NCM", "CST", "PIS", "COFINS" e as 27 UFs de
#     "AC" a "TO"). Nome diferente = coluna "some" e o campo fica em branco, sem erro nenhum;
#   - PIS, COFINS e UFs são frações (0,18 = 18%). Texto numa dessas células derruba o importador inteiro; vazio é válido
#     (produto sem aquele imposto);
#   - a linha SEM EAN é ignorada; linha sem NCM não entra no ICMS nem no PIS/COFINS;
#   - o ICMS agrupa por NCM + CST + Origem; PIS/COFINS agrupa por NCM + CST. Se os produtos de um grupo discordam entre si
#     em qualquer UF (ou em PIS/COFINS), o grupo INTEIRO é rejeitado — é o "NCM rejeitado".
# * [IMPORTANTE] → Esta é a única auditoria que LÊ o banco (só leitura, ver consulta_ao_sistema.py): o CST de saída e a
#   Origem vêm do cadastro do produto, não da planilha. Se o banco não responder, a auditoria avisa e segue sem essas
#   duas conferências — o arquivo em si continua sendo conferido por inteiro.
# * [DECISÃO, 15/10/2026] → Coluna com nome REPETIDO no cabeçalho (achado nos arquivos "VERSÃO FINAL (2)", onde CST, PIS e
#   COFINS existem no grupo ENTRADA e de novo no grupo SAÍDA): o importador monta cada linha com dict(zip(cabecalho, linha)),
#   então a ÚLTIMA coluna com o nome vence. Isso só é seguro quando a última é a do grupo SAÍDA e as anteriores são de outro
#   grupo (ENTRADA) — o título do grupo está na linha 1. Nesse caso a auditoria NÃO acusa erro e mostra qual coluna é lida
#   (chip "Colunas lidas do grupo SAÍDA"). Qualquer outra repetição continua sendo ERRO, porque o sistema leria a coluna errada.

import traceback
from collections import defaultdict

from openpyxl.utils import get_column_letter

from core.empresa import EMPRESA_MAGAZINE, EMPRESA_SAMVALE
from core.management.commands.popular_banco_suporte.conversor_celula_excel import ConversorCelulaExcel
from impostos.funcoes_auxiliares.saida.importacao_icms_ncm import (
    COLUNA_NCM, UFS_ORDENADAS, AgrupadorIcmsPorNcm, LinhaIcmsNcm, _normalizar_codigo_celula,
)
from impostos.funcoes_auxiliares.saida.importacao_pis_cofins_ncm_cst import (
    COLUNA_COFINS, COLUNA_PIS, AgrupadorPisCofinsPorNcmCst, LinhaPisCofinsSaidaPorNcmCst,
)
from impostos.funcoes_auxiliares.saida.preenchimento_impostos_saida import COLUNA_CST, COLUNA_EAN

from . import consulta_ao_sistema
from .auditor_base import AuditorDeTabela, celula, descrever_celula, sem_acentos
from .relatorio_auditoria import descrever_local, formatar_numero, limpar_texto, montar_previa
from .valores_de_celula import LeitorDeNumeros

LINHA_DO_CABECALHO = 2
PRIMEIRA_LINHA_DE_DADOS = 3
LINHAS_PROCURADAS_PARA_O_CABECALHO = 15

# Só para a pessoa reconhecer o produto nos avisos — nenhum importador lê esta coluna.
COLUNA_DESCRICAO = 'Descrição do produto'

COLUNAS_DE_PERCENTUAL = [COLUNA_PIS, COLUNA_COFINS] + list(UFS_ORDENADAS)
COLUNAS_OBRIGATORIAS = [COLUNA_EAN, COLUNA_NCM, COLUNA_CST] + COLUNAS_DE_PERCENTUAL

# O que acontece se a coluna não existir (texto que aparece no erro).
CONSEQUENCIA_DA_COLUNA_FALTANDO = {
    COLUNA_EAN: 'Sem ela nenhuma linha se liga a um produto: o EAN é a chave.',
    COLUNA_NCM: 'Sem ela nenhuma linha entra no ICMS por NCM nem no PIS/COFINS por NCM.',
    COLUNA_CST: 'Sem ela o sistema não atualiza o CST de saída dos produtos (fica o que já estava no cadastro).',
    COLUNA_PIS: 'Sem ela o PIS de saída fica em branco para todos os produtos.',
    COLUNA_COFINS: 'Sem ela o COFINS de saída fica em branco para todos os produtos.',
}

LIMITE_DA_FRACAO = 1   # 1 = 100%: acima disso quase certamente digitaram "18" em vez de "0,18"
EXEMPLOS_POR_AVISO = 8
TAMANHO_DA_DESCRICAO_NA_FRASE = 45


# Função Objetivo: Texto curto de um produto para as frases: 'linha 7 · 7891988035770 · PULVERIZADOR BOMBA…'.
def _citar(n_linha, ean=None, descricao=None):
    partes = [f'linha {n_linha}']
    if ean:
        partes.append(ean)
    if descricao:
        texto = limpar_texto(descricao)
        partes.append(texto if len(texto) <= TAMANHO_DA_DESCRICAO_NA_FRASE else texto[:TAMANHO_DA_DESCRICAO_NA_FRASE - 1] + '…')
    return ' · '.join(partes)


# Função Objetivo: "1 produto" / "3 produtos" — o número com a palavra no singular ou no plural.
def _quantos(quantidade, singular, plural):
    return f'{quantidade} {singular if quantidade == 1 else plural}'


# Função Objetivo: "a, b, c e mais N" — a lista curta que cabe numa frase.
def _listar(itens, maximo=EXEMPLOS_POR_AVISO):
    mostrados = '; '.join(itens[:maximo])
    return f'{mostrados} (e mais {len(itens) - maximo})' if len(itens) > maximo else mostrados


# Função Objetivo: Descrição curta para a tabela da prévia (a descrição não é lida pelo sistema: só ajuda a reconhecer o produto).
def _encurtar(descricao, limite=34):
    if not descricao:
        return '—'
    texto = limpar_texto(descricao)
    return texto if len(texto) <= limite else texto[:limite - 1] + '…'


# Função Objetivo: Percentual no jeito brasileiro a partir da fração lida: 0.0165 -> "1,65%".
def _percentual(fracao):
    if fracao is None:
        return '—'
    return f'{formatar_numero(round(fracao * 100, 4), 2)}%'


class AuditorBuscaLegal(AuditorDeTabela):

    EMPRESA = None   # preenchido pelas subclasses (a empresa decide em QUAL banco o EAN é procurado)

    # Função Objetivo: Consulta (só leitura) os produtos da empresa pelo EAN. Chamada por nome do módulo na hora, para os
    # testes poderem trocar a consulta sem tocar no banco.
    def _consultar_produtos(self, eans):
        return consulta_ao_sistema.consultar_produtos_por_ean(self.EMPRESA, eans)

    def auditar(self):
        self.colunas_de_saida_lidas = {}   # {nome: letra} das colunas repetidas em ENTRADA e SAÍDA (o sistema lê a de SAÍDA)
        self.avisar_abas_extras()
        aba = self.leitura.nome_da_aba_ativa
        linhas = self.leitura.linhas(aba)

        if len(linhas) < LINHA_DO_CABECALHO:
            self.relatorio.erro(
                descrever_local(aba),
                'A aba não tem o cabeçalho: o sistema lê o cabeçalho na linha 2 (a linha 1 é o título dos grupos) e a aba tem '
                f'só {_quantos(len(linhas), "linha", "linhas")}.',
                'Confira se é o arquivo certo, com o cabeçalho na linha 2 e os produtos a partir da linha 3.',
            )
            return

        cabecalho = linhas[LINHA_DO_CABECALHO - 1]
        self.memorizar_titulos(cabecalho)
        colunas, abortar = self._conferir_colunas(aba, linhas, cabecalho)
        if abortar:
            return
        faltando = [nome for nome in COLUNAS_OBRIGATORIAS if nome not in colunas]

        # Mesma regra do importador: linha com TODAS as células None é "linha em branco no fim da planilha".
        dados = [
            (n_linha, linha)
            for n_linha, linha in enumerate(linhas[PRIMEIRA_LINHA_DE_DADOS - 1:], start=PRIMEIRA_LINHA_DE_DADOS)
            if not all(valor is None for valor in linha)
        ]
        if not dados:
            self.relatorio.aviso(
                descrever_local(aba),
                'A planilha não tem nenhum produto (da linha 3 em diante está tudo em branco).',
                'Confira se é o arquivo certo.',
            )

        houve_formula_sem_resultado = self._apontar_formulas_nas_colunas_lidas(aba, colunas)
        valores, houve_erro_de_celula = self._conferir_numeros(aba, dados, colunas)
        houve_erro_de_celula = houve_erro_de_celula or houve_formula_sem_resultado
        linhas_do_produto = self._ler_identificacao(dados, colunas)
        self._conferir_ean(aba, linhas_do_produto, colunas, valores)
        self._conferir_ncm_e_cst(aba, linhas_do_produto, colunas)

        produtos, resultado_da_consulta = self._conferir_no_sistema(aba, linhas_do_produto, colunas)
        grupos = self._conferir_grupos_rejeitados(
            cabecalho, dados, linhas_do_produto, produtos, bool(faltando), houve_erro_de_celula,
        )
        self._resumir(linhas_do_produto, produtos, resultado_da_consulta, grupos)
        self._montar_previa(aba, linhas_do_produto, valores, produtos)

    # ---------- Colunas ----------

    # Função Objetivo: Confere o cabeçalho da linha 2. Devolve ({nome: índice da coluna}, abortar). `abortar` é True quando o
    # cabeçalho está inteiro em OUTRA linha (nada mais faz sentido auditar: o importador leria tudo errado).
    def _conferir_colunas(self, aba, linhas, cabecalho):
        posicoes = defaultdict(list)
        for indice, valor in enumerate(cabecalho):
            if valor in COLUNAS_OBRIGATORIAS or valor == COLUNA_DESCRICAO:
                posicoes[valor].append(indice)

        if COLUNA_EAN not in posicoes:
            for n_linha, linha in enumerate(linhas[:LINHAS_PROCURADAS_PARA_O_CABECALHO], start=1):
                if n_linha != LINHA_DO_CABECALHO and COLUNA_EAN in linha:
                    self.relatorio.erro(
                        descrever_local(aba, n_linha),
                        f'O cabeçalho (com “{COLUNA_EAN}”) está na linha {n_linha}, mas o sistema lê o cabeçalho na linha '
                        f'{LINHA_DO_CABECALHO} e os produtos a partir da linha {PRIMEIRA_LINHA_DE_DADOS}. As linhas seriam lidas '
                        'deslocadas, com os nomes de coluna errados.',
                        'Deixe o título dos grupos na linha 1, o cabeçalho na linha 2 e os produtos a partir da linha 3, como no '
                        'arquivo original do Busca Legal.',
                    )
                    return {}, True

        titulos_dos_grupos = linhas[0]
        for nome, onde_aparece in posicoes.items():
            if len(onde_aparece) <= 1 or nome not in COLUNAS_OBRIGATORIAS:
                continue
            grupos = [self._grupo_da_coluna(titulos_dos_grupos, indice) for indice in onde_aparece]
            letra_lida = get_column_letter(onde_aparece[-1] + 1)
            if self._repeticao_de_entrada_e_saida(grupos):
                self.colunas_de_saida_lidas[nome] = letra_lida
                continue
            letras = ', '.join(get_column_letter(i + 1) for i in onde_aparece)
            self.relatorio.erro(
                descrever_local(aba, LINHA_DO_CABECALHO),
                f'A coluna “{nome}” aparece {len(onde_aparece)} vezes no cabeçalho (colunas {letras}). O sistema usa só a '
                f'ÚLTIMA ({letra_lida}) e ignora as outras. {self._frase_dos_grupos(onde_aparece, grupos)}',
                'Deixe uma coluna só com esse nome — ou uma por grupo (ENTRADA e SAÍDA), com a de SAÍDA mais à direita: '
                'é a que o sistema lê.',
            )

        colunas = {nome: onde_aparece[-1] for nome, onde_aparece in posicoes.items()}

        for nome in [n for n in COLUNAS_OBRIGATORIAS if n not in colunas and n not in UFS_ORDENADAS]:
            self.relatorio.erro(
                descrever_local(aba, LINHA_DO_CABECALHO),
                f'Não achei a coluna “{nome}” no cabeçalho. {CONSEQUENCIA_DA_COLUNA_FALTANDO[nome]}{self._parecidas(cabecalho, nome)}',
                f'O sistema procura a coluna pelo nome exato “{nome}”, na linha {LINHA_DO_CABECALHO}. Confira se é o arquivo certo '
                'e se o cabeçalho não foi alterado.',
            )

        ufs_faltando = [uf for uf in UFS_ORDENADAS if uf not in colunas]
        if ufs_faltando:
            self.relatorio.erro(
                descrever_local(aba, LINHA_DO_CABECALHO),
                f'Não achei {_quantos(len(ufs_faltando), "coluna de UF", "colunas de UF")} no cabeçalho: {", ".join(ufs_faltando)}. '
                'Sem a coluna, o ICMS dessa UF fica em branco para todos os produtos.',
                f'O sistema procura as 27 UFs pelo nome exato (AC a TO), na linha {LINHA_DO_CABECALHO}.',
            )
        return colunas, False

    # Função Objetivo: Título do grupo (linha 1) a que a coluna pertence. A linha 1 só traz o título na 1ª coluna do grupo
    # (as demais ficam em branco, como numa célula mesclada); vale o título mais próximo à esquerda. None = sem título.
    def _grupo_da_coluna(self, titulos_dos_grupos, indice):
        for posicao in range(min(indice, len(titulos_dos_grupos) - 1), -1, -1):
            titulo = titulos_dos_grupos[posicao]
            if titulo is not None and limpar_texto(titulo):
                return limpar_texto(titulo)
        return None

    # Função Objetivo: True quando a repetição é a esperada — a ÚLTIMA coluna (a que o importador lê) está no grupo SAÍDA e
    # todas as anteriores estão num outro grupo com título (ENTRADA). Sem título na linha 1 não dá para saber: não aceita.
    def _repeticao_de_entrada_e_saida(self, grupos):
        def eh_saida(titulo):
            return titulo is not None and sem_acentos(titulo).startswith('saida')
        return eh_saida(grupos[-1]) and all(titulo is not None and not eh_saida(titulo) for titulo in grupos[:-1])

    # Função Objetivo: "Linha 1: J = “ENTRADA” · T = “SAÍDA”" — para a pessoa ver em que grupo cada repetição está.
    def _frase_dos_grupos(self, onde_aparece, grupos):
        if all(titulo is None for titulo in grupos):
            return 'Nenhuma delas tem título de grupo na linha 1.'
        partes = [
            f'{get_column_letter(indice + 1)} = {"“" + _encurtar(titulo, 30) + "”" if titulo else "sem grupo"}'
            for indice, titulo in zip(onde_aparece, grupos)
        ]
        return f'Grupos na linha 1: {" · ".join(partes)}.'

    # Função Objetivo: Frase "existe algo parecido": cabeçalho escrito diferente (maiúscula, acento, espaço) do nome que o sistema procura.
    def _parecidas(self, cabecalho, nome):
        alvo = sem_acentos(nome)
        parecidas = [
            f'“{limpar_texto(valor)}” (coluna {get_column_letter(i + 1)})'
            for i, valor in enumerate(cabecalho)
            if valor is not None and limpar_texto(valor) and sem_acentos(limpar_texto(valor)) == alvo
        ]
        if not parecidas:
            return ''
        return f' Existe {", ".join(parecidas)}, escrita de um jeito diferente.'

    # Função Objetivo: Fórmula sem resultado salvo só importa nas colunas que o importador lê (a "ICMS MÉDIA" é recalculada pelo sistema).
    # Devolve True se achou (ERRO registrado).
    def _apontar_formulas_nas_colunas_lidas(self, aba, colunas):
        lidas = set(colunas.values())
        achadas = [
            (n_linha, n_coluna) for n_linha, n_coluna in self.leitura.formulas_sem_resultado.get(aba, [])
            if n_linha >= PRIMEIRA_LINHA_DE_DADOS and n_coluna in lidas
        ]
        if not achadas:
            return False
        exemplos = ', '.join(descrever_celula(linha, coluna) for linha, coluna in achadas[:5])
        mais = f' (e mais {len(achadas) - 5})' if len(achadas) > 5 else ''
        self.relatorio.erro(
            descrever_local(aba),
            f'{_quantos(len(achadas), "célula de coluna que o sistema lê tem", "células de colunas que o sistema lê têm")} fórmula '
            f'mas nenhum resultado salvo: {exemplos}{mais}. '
            'O sistema leria essas células como vazias.',
            'Abra o arquivo no Excel e salve de novo (Ctrl+S) para o Excel gravar o resultado das fórmulas.',
        )
        return True

    # ---------- Valores ----------

    # Função Objetivo: Lê PIS, COFINS e as 27 UFs de cada linha. Devolve ({n_linha: {nome: fração|None}}, houve_erro).
    # Texto, data, VERDADEIRO/FALSO, erro de fórmula, negativo ou acima de 100% = ERRO; vazio é válido.
    def _conferir_numeros(self, aba, dados, colunas):
        leitor = LeitorDeNumeros(self.relatorio, aba)
        erros_antes = self.relatorio.total_de_erros
        valores = {}
        for n_linha, linha in dados:
            da_linha = {}
            for nome in COLUNAS_DE_PERCENTUAL:
                indice = colunas.get(nome)
                if indice is None:
                    continue
                bruto = celula(linha, indice)
                valor, ok = leitor.ler(n_linha, indice, bruto, nome, obrigatorio=False)
                if ok and valor is not None and valor > LIMITE_DA_FRACAO:
                    self.relatorio.erro(
                        self.local(aba, n_linha, indice),
                        f'O valor {formatar_numero(valor)} é maior que 100%: o sistema leria {formatar_numero(valor * 100, 2)}%.',
                        'A planilha traz as alíquotas em fração (0,18 = 18%). Confira se não foi digitado “18” em vez de “0,18”.',
                    )
                    ok = False
                da_linha[nome] = valor if ok else None
            valores[n_linha] = da_linha
        return valores, self.relatorio.total_de_erros > erros_antes

    # ---------- Identificação dos produtos ----------

    # Função Objetivo: EAN, NCM, CST e descrição de cada linha, já no formato que o importador usa (texto sem ".0" e sem espaços).
    def _ler_identificacao(self, dados, colunas):
        lidos = []
        for n_linha, linha in dados:
            descricao = celula(linha, colunas[COLUNA_DESCRICAO]) if COLUNA_DESCRICAO in colunas else None
            lidos.append({
                'linha': n_linha,
                'ean': _normalizar_codigo_celula(celula(linha, colunas[COLUNA_EAN])) if COLUNA_EAN in colunas else None,
                'ncm': _normalizar_codigo_celula(celula(linha, colunas[COLUNA_NCM])) if COLUNA_NCM in colunas else None,
                'cst': _normalizar_codigo_celula(celula(linha, colunas[COLUNA_CST])) if COLUNA_CST in colunas else None,
                'descricao': descricao if isinstance(descricao, str) and descricao.strip() else None,
            })
        return lidos

    # Função Objetivo: EAN vazio (linha ignorada) e EAN repetido. (EAN com letras NÃO é problema por si: produto sem código de barras
    # pode ter um código interno no campo EAN — quem diz se casa é o banco, em _conferir_no_sistema.)
    def _conferir_ean(self, aba, produtos, colunas, valores):
        if COLUNA_EAN not in colunas:
            return
        indice = colunas[COLUNA_EAN]
        por_ean = defaultdict(list)
        for produto in produtos:
            ean = produto['ean']
            if ean is None:
                nome = f' (“{limpar_texto(produto["descricao"])}”)' if produto['descricao'] else ''
                self.relatorio.erro(
                    self.local(aba, produto['linha'], indice),
                    f'A célula do EAN está vazia, mas a linha tem outros dados{nome}. O sistema ignora a linha: sem EAN ela não '
                    'se liga a nenhum produto.',
                    'Preencha o EAN ou apague a linha inteira.',
                )
                continue
            por_ean[ean].append(produto)

        repetidos_iguais = []
        rotulos = ['NCM', 'CST'] + COLUNAS_DE_PERCENTUAL
        for ean, repetidas in por_ean.items():
            if len(repetidas) < 2:
                continue
            assinaturas = [
                (p['ncm'], p['cst']) + tuple(valores[p['linha']].get(nome) for nome in COLUNAS_DE_PERCENTUAL)
                for p in repetidas
            ]
            numeros = ', '.join(str(p['linha']) for p in repetidas)
            if len(set(assinaturas)) == 1:
                repetidos_iguais.append(f'{ean} (linhas {numeros})')
                continue
            primeira = assinaturas[0]
            diferem = [rotulos[i] for i in range(len(rotulos)) if any(a[i] != primeira[i] for a in assinaturas[1:])]
            self.relatorio.erro(
                descrever_local(aba, coluna=indice, titulo_da_coluna=COLUNA_EAN),
                f'O EAN {ean} aparece nas linhas {numeros} com dados DIFERENTES ({_listar(diferem, 6)}). O sistema usa a 1ª linha '
                'para o CST, mas as outras entram no ICMS e no PIS/COFINS: não dá para saber qual vale.',
                'Deixe uma linha só por produto (apague a repetida errada).',
            )
        if repetidos_iguais:
            self.relatorio.aviso(
                descrever_local(aba, coluna=indice, titulo_da_coluna=COLUNA_EAN),
                f'{_quantos(len(repetidos_iguais), "EAN aparece", "EANs aparecem")} em mais de uma linha, com os mesmos dados: '
                f'{_listar(repetidos_iguais)}.',
                'Não atrapalha a leitura, mas vale deixar uma linha só por produto.',
            )

    # Função Objetivo: NCM vazio (a linha não entra em ICMS/PIS/COFINS), NCM fora do padrão de 8 dígitos e CST vazio/estranho.
    def _conferir_ncm_e_cst(self, aba, produtos, colunas):
        sem_ncm, ncm_estranho, sem_cst, cst_estranho = [], [], [], []
        for produto in produtos:
            if produto['ean'] is None:
                continue   # já é erro (linha sem EAN)
            citacao = _citar(produto['linha'], produto['ean'], produto['descricao'])
            if COLUNA_NCM in colunas:
                ncm = produto['ncm']
                if ncm is None:
                    sem_ncm.append(citacao)
                elif not (ncm.isascii() and ncm.isdigit() and len(ncm) == 8):
                    ncm_estranho.append(f'{citacao} → NCM “{ncm}”')
            if COLUNA_CST in colunas:
                cst = produto['cst']
                if cst is None:
                    sem_cst.append(citacao)
                elif not (cst.isascii() and cst.isdigit()):
                    cst_estranho.append(f'{citacao} → CST “{cst}”')

        if sem_ncm:
            self.relatorio.aviso(
                descrever_local(aba, coluna=colunas[COLUNA_NCM], titulo_da_coluna=COLUNA_NCM),
                f'{_quantos(len(sem_ncm), "produto está", "produtos estão")} sem NCM: {_listar(sem_ncm)}. O sistema não usa '
                f'{"essa linha" if len(sem_ncm) == 1 else "essas linhas"} no ICMS nem no PIS/COFINS por NCM.',
                'Preencha o NCM na planilha, se for um produto que precisa de imposto de saída.',
            )
        if ncm_estranho:
            self.relatorio.aviso(
                descrever_local(aba, coluna=colunas[COLUNA_NCM], titulo_da_coluna=COLUNA_NCM),
                f'{_quantos(len(ncm_estranho), "NCM não tem", "NCMs não têm")} 8 dígitos: {_listar(ncm_estranho)}. O sistema não '
                'completa zeros: o NCM dificilmente vai casar com o do cadastro.',
                'Confira se o Excel não tirou um zero à esquerda (formate a coluna como Texto) e corrija na planilha.',
            )
        if sem_cst:
            self.relatorio.aviso(
                descrever_local(aba, coluna=colunas[COLUNA_CST], titulo_da_coluna=COLUNA_CST),
                f'{_quantos(len(sem_cst), "produto está", "produtos estão")} sem CST na planilha: {_listar(sem_cst)}. O sistema '
                f'não mexe no CST {"desse produto" if len(sem_cst) == 1 else "desses produtos"}: vale o que já está no cadastro.',
                'Preencha o CST na planilha para atualizar o cadastro.',
            )
        if cst_estranho:
            self.relatorio.aviso(
                descrever_local(aba, coluna=colunas[COLUNA_CST], titulo_da_coluna=COLUNA_CST),
                f'{_quantos(len(cst_estranho), "CST tem", "CSTs têm")} caracteres que não são dígitos: {_listar(cst_estranho)}.',
                'Confira o CST na planilha: o sistema grava exatamente o que está na célula.',
            )

    # ---------- Banco (só leitura) ----------

    # Função Objetivo: Procura no banco da empresa os EANs da planilha. Devolve ({ean: (cst, origem)} ou None se a consulta
    # falhou, texto para o resumo). "EAN que não existe no sistema" é só AVISO (decisão de 14/10).
    def _conferir_no_sistema(self, aba, produtos, colunas):
        if COLUNA_EAN not in colunas:
            return None, 'não consultado'
        eans = {p['ean'] for p in produtos if p['ean']}
        if not eans:
            return {}, 'nenhum EAN para consultar'
        try:
            no_sistema = self._consultar_produtos(eans)
        except Exception:
            traceback.print_exc()
            self.relatorio.aviso(
                'consulta ao sistema',
                'Não consegui consultar os produtos no banco da empresa, então não conferi “EAN que não existe no sistema” nem '
                '“NCM rejeitado”.',
                'Confira se o banco está no ar e escolha o arquivo de novo. O resto do arquivo foi conferido normalmente.',
            )
            return None, 'não consegui consultar'

        fora = [
            _citar(p['linha'], p['ean'], p['descricao'])
            for p in produtos if p['ean'] in eans and p['ean'] not in no_sistema
        ]
        if fora:
            quantidade_de_eans = len({p['ean'] for p in produtos if p['ean'] in eans and p['ean'] not in no_sistema})
            self.relatorio.aviso(
                descrever_local(aba, coluna=colunas[COLUNA_EAN], titulo_da_coluna=COLUNA_EAN),
                f'{_quantos(quantidade_de_eans, "EAN não existe", "EANs não existem")} no sistema desta empresa: {_listar(fora)}.',
                'Isso não impede o envio: essas linhas só não terão efeito até o produto existir (ou se o EAN estiver diferente '
                'do cadastro, corrija um dos dois).',
            )
        achados = sum(1 for ean in eans if ean in no_sistema)
        return no_sistema, f'{achados} de {len(eans)} EANs'

    # ---------- NCM rejeitado ----------

    # Função Objetivo: Roda o MESMO agrupamento do importador sobre as linhas da planilha e avisa cada grupo rejeitado.
    # O CST usado é o que o sistema terá DEPOIS de atualizar o cadastro com a planilha (o CST da planilha, quando existe;
    # senão o que já está no cadastro) — o importador roda o passo do CST antes do ICMS e do PIS/COFINS. Devolve
    # {'icms': (aceitos, rejeitados), 'pis': (aceitos, rejeitados)} ou None se a conferência foi adiada.
    def _conferir_grupos_rejeitados(self, cabecalho, dados, produtos, no_sistema, faltam_colunas, houve_erro_de_celula):
        if no_sistema is None or faltam_colunas or houve_erro_de_celula:
            return None

        cst_efetivo = {}
        for produto in produtos:
            ean = produto['ean']
            if ean and ean not in cst_efetivo and ean in no_sistema:
                cst_efetivo[ean] = produto['cst'] if produto['cst'] is not None else no_sistema[ean][0]

        conversor = ConversorCelulaExcel(origem='openpyxl')
        agrupador_icms = AgrupadorIcmsPorNcm()
        agrupador_pis = AgrupadorPisCofinsPorNcmCst()
        for (n_linha, linha), produto in zip(dados, produtos):
            bruto = dict(zip(cabecalho, linha))   # igualzinho ao importador
            ean = produto['ean']
            if ean in no_sistema:
                cst, origem = cst_efetivo[ean], no_sistema[ean][1]
            else:
                cst, origem = None, None
            agrupador_icms.adicionar_linha(LinhaIcmsNcm(bruto, conversor).extrair_campos().aplicar_cst_e_origem(cst, origem))
            agrupador_pis.adicionar_linha(LinhaPisCofinsSaidaPorNcmCst(bruto, conversor).extrair_campos().aplicar_cst_saida(cst))
        agrupador_icms.processar()
        agrupador_pis.processar()

        for rejeitado in agrupador_icms.rejeitados:
            self._avisar_icms_rejeitado(rejeitado)
        for rejeitado in agrupador_pis.rejeitados:
            self._avisar_pis_cofins_rejeitado(rejeitado)
        return {
            'icms': (len(agrupador_icms.aceitos), len(agrupador_icms.rejeitados)),
            'pis': (len(agrupador_pis.aceitos), len(agrupador_pis.rejeitados)),
        }

    def _avisar_icms_rejeitado(self, rejeitado):
        origem = rejeitado.origem_mercadoria_cadastro if rejeitado.origem_mercadoria_cadastro is not None else 'em branco'
        ufs = sorted(rejeitado.divergencias_por_uf)
        diagnostico, _ = rejeitado.resumir_diagnostico()
        self.relatorio.aviso(
            f'ICMS · NCM {rejeitado.ncm} · CST {rejeitado.cst} · origem {origem}',
            f'Os {rejeitado.total_eans_no_grupo} produtos deste grupo não têm o mesmo ICMS em {_quantos(len(ufs), "UF", "UFs")}: '
            f'{_listar(ufs, 10)}. O sistema rejeita o grupo inteiro — esses produtos ficam sem ICMS de saída por NCM. '
            f'Diagnóstico: {diagnostico}.',
            'Deixe todos os produtos do grupo com o mesmo valor em cada UF (ou todos em branco). Se o produto suspeito está com '
            'NCM, CST ou origem errados, corrija no cadastro.',
        )

    def _avisar_pis_cofins_rejeitado(self, rejeitado):
        campos = sorted(rejeitado.divergencias_por_campo)
        detalhes = []
        for campo in campos:
            partes = []
            for valor, eans in rejeitado._grupos_por_valor(rejeitado.divergencias_por_campo[campo]):
                texto_do_valor = f'{formatar_numero(float(valor), 2)}%' if valor is not None else 'em branco'
                quem = f' (EAN {", ".join(eans)})' if len(eans) <= 3 else ''
                partes.append(f'{texto_do_valor} em {_quantos(len(eans), "produto", "produtos")}{quem}')
            detalhes.append(f'{campo}: {" · ".join(partes)}')
        self.relatorio.aviso(
            f'PIS/COFINS · NCM {rejeitado.ncm} · CST {rejeitado.cst}',
            f'Os {rejeitado.total_eans_no_grupo} produtos deste grupo não têm o mesmo valor em {" e ".join(campos)}. '
            f'O sistema rejeita o grupo inteiro — esses produtos ficam sem PIS/COFINS por NCM. {" — ".join(detalhes)}.',
            'Deixe todos os produtos do grupo com o mesmo valor (ou todos em branco). Se um produto está com NCM ou CST errados, '
            'corrija no cadastro.',
        )

    # ---------- Resumo e pré-visualização ----------

    def _resumir(self, produtos, no_sistema, resultado_da_consulta, grupos):
        self.relatorio.resumir('Produtos na planilha', len(produtos))
        self.relatorio.resumir('NCMs diferentes', len({p['ncm'] for p in produtos if p['ncm']}))
        self.relatorio.resumir('No sistema', resultado_da_consulta)
        if self.colunas_de_saida_lidas:
            lidas = [(nome, self.colunas_de_saida_lidas[nome]) for nome in COLUNAS_OBRIGATORIAS if nome in self.colunas_de_saida_lidas]
            self.relatorio.resumir('Colunas lidas do grupo SAÍDA', ' · '.join(f'{nome} {letra}' for nome, letra in lidas))
        if grupos is None:
            self.relatorio.resumir('Grupos de ICMS e PIS/COFINS', 'não conferidos — há erros ou o banco não respondeu')
            return
        aceitos, rejeitados = grupos['icms']
        self.relatorio.resumir('Grupos de ICMS', f'{aceitos} aceitos · {rejeitados} rejeitados')
        aceitos, rejeitados = grupos['pis']
        self.relatorio.resumir('Grupos de PIS/COFINS', f'{aceitos} aceitos · {rejeitados} rejeitados')

    # Função Objetivo: O que o sistema vai ler de cada linha (PIS, COFINS e ICMS de SP como percentual, quantas das 27 UFs vieram preenchidas).
    def _montar_previa(self, aba, produtos, valores, no_sistema):
        linhas = []
        for produto in produtos:
            da_linha = valores.get(produto['linha'], {})
            preenchidas = sum(1 for uf in UFS_ORDENADAS if da_linha.get(uf) is not None)
            if no_sistema is None or produto['ean'] is None:
                situacao = '—'
            else:
                situacao = 'Sim' if produto['ean'] in no_sistema else 'Não'
            linhas.append([
                produto['linha'], produto['ean'] or '—', _encurtar(produto['descricao']),
                produto['ncm'] or '—', produto['cst'] or '—',
                _percentual(da_linha.get(COLUNA_PIS)), _percentual(da_linha.get(COLUNA_COFINS)), _percentual(da_linha.get('SP')),
                f'{preenchidas}/{len(UFS_ORDENADAS)}', situacao,
            ])
        self.relatorio.adicionar_previa(montar_previa(
            f'aba “{aba}”',
            ['Linha', 'EAN', 'Descrição', 'NCM', 'CST', 'PIS', 'COFINS', 'ICMS SP', 'UFs preenchidas', 'No sistema'],
            linhas,
        ))


class AuditorBuscaLegalMagazine(AuditorBuscaLegal):
    EMPRESA = EMPRESA_MAGAZINE


class AuditorBuscaLegalSamvale(AuditorBuscaLegal):
    EMPRESA = EMPRESA_SAMVALE
