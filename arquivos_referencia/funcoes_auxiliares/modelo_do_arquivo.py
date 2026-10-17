# arquivos_referencia/funcoes_auxiliares/modelo_do_arquivo.py

# Função Objetivo: Entrega o MODELO de cada arquivo de referência: uma planilha .xlsx IGUAL ao arquivo real (visual e estrutura), para a pessoa
# preencher ou usar de guia do "como o sistema espera receber". É o que o botão "Modelo" do cartão baixa.
#
# * [DECISÃO, 17/10/2026] → Matheus perdeu o arquivo de referência da Busca Legal e não sabia mais "exatamente como o sistema espera receber o
#   arquivo". 1ª versão (montada do zero em código) ficou "nem um pouco estilizada e nem um pouco parecida com as referências reais". Agora o
#   modelo NASCE DO ARQUIVO REAL: scripts_dev/gerar_modelos_base_dos_arquivos_de_referencia.py abre cada arquivo da pasta "Arquivos usados para
#   Popular Banco", esvazia o que não pode ser exemplo e grava em arquivos_referencia/modelos_base/<id>.xlsx. Este módulo só ENTREGA esses arquivos.
#   - [DECISÃO, 17/10/2026, Matheus] O modelo só GUIA: cabeçalho + NO MÁXIMO 1 linha de exemplo (2 no Mercado Livre: o par Sem/Com Frete Grátis Rápido
#     é 1 conjunto) + uma linha cinza escura inteira marcando o fim. Nenhum valor de frete vem preenchido (valor velho esquecido = preço errado).
#   - BUSCA LEGAL: cabeçalho igual ao real (cada empresa com a sua ordem de colunas) + 1 linha de exemplo real (CST 00) + a linha cinza.
# * [IMPORTANTE] → O modelo não pode mentir: antes de entregar, confere se o cabeçalho do modelo ainda bate com o que o importador lê (as MESMAS
#   constantes que as auditorias usam: nome das colunas e as 27 UFs da Busca Legal; posição das colunas e nome das abas do frete). Se um importador
#   mudar e o modelo não, a tela avisa "modelo desatualizado" em vez de entregar um arquivo que o sistema leria errado.

import io
import re
import unicodedata
from pathlib import Path

from openpyxl import load_workbook

PASTA_DOS_MODELOS = Path(__file__).resolve().parent.parent / 'modelos_base'


class ModeloDesatualizado(Exception):
    """O modelo guardado não bate mais com o que o importador lê (ou o arquivo do modelo sumiu)."""


def _caminho_do_modelo(arquivo):
    return PASTA_DOS_MODELOS / f'{arquivo.id}.xlsx'


def _sem_acentos(texto):
    decomposto = unicodedata.normalize('NFD', str(texto))
    return ''.join(c for c in decomposto if unicodedata.category(c) != 'Mn').lower()


# Função Objetivo: A célula da linha do cabeçalho (1 = primeira linha; coluna 0 = A) tem todas as palavras? (sem acento nem maiúscula)
def _cabecalho_tem(aba, linha, coluna, palavras):
    texto = _sem_acentos(aba.cell(row=linha, column=coluna + 1).value or '')
    return all(palavra in texto for palavra in palavras)


def _exigir(condicao, o_que):
    if not condicao:
        raise ModeloDesatualizado(o_que)


# ======================================================================================================================
# Conferência do modelo contra os importadores (mesmas constantes das auditorias)
# ======================================================================================================================
def _conferir_busca_legal(livro):
    from .auditoria_busca_legal import COLUNAS_OBRIGATORIAS, LINHA_DO_CABECALHO

    _exigir(len(livro.sheetnames) == 1, 'a Busca Legal tem de ter uma aba só')
    aba = livro.active
    cabecalho = [aba.cell(row=LINHA_DO_CABECALHO, column=c).value for c in range(1, aba.max_column + 1)]
    faltando = [nome for nome in COLUNAS_OBRIGATORIAS if nome not in cabecalho]
    _exigir(not faltando, f'faltam no cabeçalho (linha {LINHA_DO_CABECALHO}) as colunas: {", ".join(faltando)}')
    repetidas = [nome for nome in COLUNAS_OBRIGATORIAS if cabecalho.count(nome) > 1]
    _exigir(not repetidas, f'colunas repetidas no cabeçalho: {", ".join(repetidas)}')


def _conferir_frete_ml(livro):
    from core.management.commands.popular_banco_suporte import importar_tabela_frete_ml as importador

    _exigir(len(livro.sheetnames) == 1, 'o frete do Mercado Livre tem de ter uma aba só')
    aba = livro.active
    for coluna, palavras in ((importador.COLUNA_PESO_MIN, ['peso', 'min']), (importador.COLUNA_PESO_MAX, ['peso', 'max']), (importador.COLUNA_REGIME, ['regime'])):
        _exigir(_cabecalho_tem(aba, 1, coluna, palavras), f'a coluna {coluna + 1} do cabeçalho não é mais “{" ".join(palavras)}”')
    for coluna in range(importador.COLUNA_INICIO_PRECOS, importador.COLUNA_INICIO_PRECOS + 8):
        texto = str(aba.cell(row=1, column=coluna + 1).value or '')
        _exigir(re.search(r'\d', texto), f'a faixa de preço da coluna {coluna + 1} não tem número no cabeçalho (“{texto}”)')


def _conferir_frete_magalu(livro):
    from core.management.commands.popular_banco_suporte import importar_tabela_frete_magalu as importador

    _exigir(len(livro.sheetnames) == 1, 'o frete do Magalu tem de ter uma aba só')
    aba = livro.active
    esperadas = (
        (importador.COL_VALOR_BAIXA, ['92']), (importador.COL_VALOR_MEDIA, ['92', '97']), (importador.COL_VALOR_ALTA, ['97']),
        (importador.COL_PESO_MIN, ['minimo']), (importador.COL_PESO_MAX, ['maximo']),
    )
    for coluna, palavras in esperadas:
        _exigir(_cabecalho_tem(aba, 1, coluna, palavras), f'a coluna {coluna + 1} do cabeçalho não bate mais com o importador')


def _conferir_frete_tiktok(livro):
    from core.management.commands.popular_banco_suporte import importar_tabela_frete_tiktok as importador

    _exigir(len(livro.sheetnames) == 1, 'o frete do TikTok Shop tem de ter uma aba só')
    aba = livro.active
    for coluna, palavras in ((importador.COL_VALOR, ['valor']), (importador.COL_PESO_MIN, ['minimo']), (importador.COL_PESO_MAX, ['maximo'])):
        _exigir(_cabecalho_tem(aba, 1, coluna, palavras), f'a coluna {coluna + 1} do cabeçalho não bate mais com o importador')


def _conferir_frete_amazon(livro):
    from core.management.commands.popular_banco_suporte import importar_tabela_frete_amazon as importador

    nomes_esperados = list(importador.ABAS_POR_TIPO.values())
    _exigir(sorted(livro.sheetnames) == sorted(nomes_esperados), f'as abas têm de ser {nomes_esperados} e são {livro.sheetnames}')
    esperadas = (
        (importador.COL_PESO_MIN, ['peso', 'min']), (importador.COL_PESO_MAX, ['peso', 'max']), (importador.COL_PRECO_MIN, ['preco', 'min']),
        (importador.COL_PRECO_MAX, ['preco', 'max']), (importador.COL_VALOR, ['valor']),
    )
    for aba in livro.worksheets:
        for coluna, palavras in esperadas:
            _exigir(_cabecalho_tem(aba, 1, coluna, palavras), f'na aba “{aba.title}” a coluna {coluna + 1} do cabeçalho não bate mais com o importador')


_CONFERENCIAS = {
    'frete-mercado-livre': _conferir_frete_ml,
    'frete-magalu': _conferir_frete_magalu,
    'frete-tiktok-shop': _conferir_frete_tiktok,
    'frete-amazon': _conferir_frete_amazon,
    'busca-legal-magazine': _conferir_busca_legal,
    'busca-legal-samvale': _conferir_busca_legal,
}


# Função Objetivo: Há modelo para este arquivo do catálogo? (Um arquivo novo no catálogo sem modelo é um esquecimento que o teste pega.)
def tem_modelo(arquivo):
    return arquivo.id in _CONFERENCIAS and _caminho_do_modelo(arquivo).is_file()


# Função Objetivo: O modelo do arquivo: {'conteudo': bytes do .xlsx, 'nome': nome sugerido para gravar no computador}.
# O nome começa com "Modelo - " e leva o nome fixo que o arquivo tem no Drive: o modelo nunca se confunde com o arquivo de verdade.
# Levanta ModeloDesatualizado se o modelo não existe ou não bate mais com o importador.
def gerar_modelo(arquivo):
    if arquivo.id not in _CONFERENCIAS:
        raise ValueError(f'O arquivo “{arquivo.id}” ainda não tem modelo.')
    caminho = _caminho_do_modelo(arquivo)
    if not caminho.is_file():
        raise ModeloDesatualizado(
            f'o arquivo {caminho.name} não está na pasta modelos_base. Rode scripts_dev/gerar_modelos_base_dos_arquivos_de_referencia.py.'
        )
    conteudo = caminho.read_bytes()
    _CONFERENCIAS[arquivo.id](load_workbook(io.BytesIO(conteudo)))
    return {'conteudo': conteudo, 'nome': f'Modelo - {arquivo.nome_no_drive}'}
