# arquivos_referencia/funcoes_auxiliares/uniao_da_busca_legal.py

# Função Objetivo: O modo EXPANDIR no ENVIO da planilha Busca Legal ao Drive. Em vez de trocar o arquivo do Drive por inteiro, JUNTA a planilha nova ao que
# já está lá: a linha cujo EAN já existe no arquivo do Drive é atualizada com os valores da planilha nova, e a linha com EAN novo vai para o FINAL do arquivo
# (como um append). O que só existe no arquivo do Drive fica como está. Devolve os bytes do .xlsx resultante e os números da junção.
#
# * [DECISÃO, 16/10/2026] → Ideia do Matheus: Expandir no envio = "expande a planilha do Drive: atualiza as existentes e adiciona as novas no final";
#   Substituir no envio = o arquivo do Drive é trocado por inteiro (é o que o envio sempre fez, ver envio_ao_drive.py). O modo é UMA configuração da empresa
#   (card no topo da tela) que vale também para o "Atualizar o sistema".
# * [EXPLICAÇÃO] → A chave da linha é o EAN (a coluna "Cód Barras"), a mesma que o sistema usa. O cabeçalho (linha 2) das duas planilhas precisa ser IGUAL,
#   coluna por coluna: é assim que se sabe que cada valor da planilha nova cai na coluna certa do arquivo do Drive (a planilha repete colunas de ENTRADA e SAÍDA
#   com o mesmo nome, então juntar "por nome" seria ambíguo). Cabeçalho diferente = não junta, nada é enviado, e a pessoa usa o Substituir.
# * [EXPLICAÇÃO] → Só a aba ATIVA é juntada (a mesma que a auditoria e o sistema leem); as outras abas do arquivo do Drive ficam como estão. A linha 1 (títulos
#   dos grupos) e o cabeçalho do arquivo do Drive não mudam. As linhas novas copiam o formato (cor, número, borda) da última linha de produto do arquivo do Drive.
# * [IMPORTANTE] → O arquivo é reescrito com os VALORES: uma célula com fórmula passa a guardar o último resultado calculado (o sistema só lê valores, então a
#   leitura não muda). `formulas_viraram_valor` conta quantas foram, para a tela avisar. Itens que o openpyxl não guarda (imagens, gráficos) não sobrevivem; a
#   versão anterior do arquivo continua no histórico de versões do Drive.
# * [IMPORTANTE] → Esta função só junta bytes: não fala com o Drive nem com o banco. Quem chama (a tela) é que lê o arquivo do Drive, confere o resultado com
#   a auditoria de sempre e só então envia.

import io
from copy import copy

from openpyxl import load_workbook

LIMITE_DE_LINHAS_DA_ABA = 20000   # só para não varrer uma aba com milhões de linhas formatadas; a auditoria limita muito antes (5000 linhas)


class ErroDaUniao(Exception):
    """Problema PREVISTO ao juntar as planilhas (cabeçalho diferente, arquivo que não abre…). `args[0]` é o texto para a tela."""


def _texto(valor):
    return ' '.join(str(valor).split()) if valor is not None else ''


# O cabeçalho como lista de textos, sem as colunas vazias do fim.
def _cabecalho(valores):
    textos = [_texto(valor) for valor in valores]
    while textos and not textos[-1]:
        textos.pop()
    return textos


# O EAN como texto comparável: 7891234567890, 7891234567890.0 e " 7891234567890 " são o mesmo EAN. Vazio = None (linha sem EAN não entra).
def _ean(valor):
    if valor is None or isinstance(valor, bool):
        return None
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    texto = str(valor).strip()
    if texto.endswith('.0') and texto[:-2].isdigit():
        texto = texto[:-2]
    return texto or None


# Dois valores de célula são iguais se são o mesmo valor ou se os dois estão vazios (None e texto vazio contam como vazio).
def _valor_igual(a, b):
    return (a in (None, '') and b in (None, '')) or a == b


def _abrir(conteudo, quem):
    try:
        return load_workbook(io.BytesIO(conteudo), data_only=True)
    except Exception as erro:
        raise ErroDaUniao(f'Não consegui abrir {quem} para juntar as planilhas. Confira se é um .xlsx válido.') from erro


# Quantas células de fórmula o arquivo do Drive tem (em todas as abas): elas passam a guardar o valor calculado.
def _contar_formulas(conteudo):
    try:
        livro = load_workbook(io.BytesIO(conteudo), read_only=True)
        total = 0
        for aba in livro.worksheets:
            for linha in aba.iter_rows(values_only=True):
                total += sum(1 for valor in linha if isinstance(valor, str) and valor.startswith('='))
        livro.close()
        return total
    except Exception:
        return 0


def _cabecalho_da_aba(aba, linha_do_cabecalho):
    if aba.max_row < linha_do_cabecalho:
        raise ErroDaUniao(f'A aba “{aba.title}” não tem o cabeçalho (linha {linha_do_cabecalho}).')
    return _cabecalho([aba.cell(row=linha_do_cabecalho, column=coluna).value for coluna in range(1, aba.max_column + 1)])


# Função Objetivo: Junta `conteudo_novo` ao `conteudo_do_drive` pelo EAN (`coluna_chave`). Devolve
# {'conteudo' (bytes), 'atualizadas', 'novas', 'iguais', 'sem_chave', 'formulas_viraram_valor'}.
# `atualizadas` = EANs que já existiam e mudaram; `iguais` = existiam e já estavam iguais; `novas` = EANs acrescentados no final;
# `sem_chave` = linhas da planilha nova sem EAN (ignoradas, como o sistema faz).
def juntar_planilhas(conteudo_do_drive, conteudo_novo, coluna_chave, linha_do_cabecalho=2):
    livro_do_drive = _abrir(conteudo_do_drive, 'o arquivo que está no Drive')
    livro_novo = _abrir(conteudo_novo, 'a planilha escolhida')
    aba_do_drive = livro_do_drive.active
    aba_nova = livro_novo.active
    if aba_do_drive.max_row > LIMITE_DE_LINHAS_DA_ABA or aba_nova.max_row > LIMITE_DE_LINHAS_DA_ABA:
        raise ErroDaUniao('Uma das planilhas tem linhas demais para juntar. Confira se é o arquivo certo.')

    cabecalho_do_drive = _cabecalho_da_aba(aba_do_drive, linha_do_cabecalho)
    cabecalho_novo = _cabecalho_da_aba(aba_nova, linha_do_cabecalho)
    if cabecalho_do_drive != cabecalho_novo:
        for posicao in range(max(len(cabecalho_do_drive), len(cabecalho_novo))):
            no_drive = cabecalho_do_drive[posicao] if posicao < len(cabecalho_do_drive) else '(vazia)'
            na_nova = cabecalho_novo[posicao] if posicao < len(cabecalho_novo) else '(vazia)'
            if no_drive != na_nova:
                raise ErroDaUniao(
                    f'As colunas da planilha escolhida não são as mesmas do arquivo que está no Drive (a coluna {posicao + 1} é “{no_drive}” no Drive e '
                    f'“{na_nova}” na planilha escolhida), então não dá para juntar. Para trocar o arquivo inteiro, use o modo Substituir.'
                )
    largura = len(cabecalho_do_drive)
    chave = _texto(coluna_chave)
    if chave not in cabecalho_do_drive:
        raise ErroDaUniao(f'O arquivo não tem a coluna “{coluna_chave}” (o EAN) na linha {linha_do_cabecalho}, então não dá para juntar.')
    posicao_da_chave = cabecalho_do_drive.index(chave)

    # Onde está cada EAN no arquivo do Drive (pode haver o mesmo EAN em mais de uma linha) e onde termina a última linha de produto.
    linhas_por_ean = {}
    ultima_linha = linha_do_cabecalho
    for numero in range(linha_do_cabecalho + 1, aba_do_drive.max_row + 1):
        valores = [aba_do_drive.cell(row=numero, column=coluna).value for coluna in range(1, largura + 1)]
        if all(valor is None for valor in valores):
            continue
        ultima_linha = numero
        ean = _ean(valores[posicao_da_chave])
        if ean:
            linhas_por_ean.setdefault(ean, []).append(numero)

    proxima_linha = ultima_linha + 1
    estado = {}   # ean -> 'nova' | 'atualizada' | 'igual'
    sem_chave = 0
    for linha in aba_nova.iter_rows(min_row=linha_do_cabecalho + 1, values_only=True):
        valores = list(linha[:largura]) + [None] * (largura - len(linha))
        if all(valor is None for valor in valores):
            continue
        ean = _ean(valores[posicao_da_chave])
        if not ean:
            sem_chave += 1
            continue
        if ean in linhas_por_ean:
            mudou = False
            for numero in linhas_por_ean[ean]:
                for coluna, valor in enumerate(valores, start=1):
                    if coluna == posicao_da_chave + 1:
                        continue   # é o mesmo EAN por construção: mantém como o arquivo do Drive o guarda (número ou texto)
                    celula = aba_do_drive.cell(row=numero, column=coluna)
                    if not _valor_igual(celula.value, valor):
                        celula.value = valor
                        mudou = True
            anterior = estado.get(ean)
            if anterior is None:
                estado[ean] = 'atualizada' if mudou else 'igual'
            elif anterior == 'igual' and mudou:
                estado[ean] = 'atualizada'
        else:
            for coluna, valor in enumerate(valores, start=1):
                destino = aba_do_drive.cell(row=proxima_linha, column=coluna, value=valor)
                if ultima_linha > linha_do_cabecalho:
                    origem = aba_do_drive.cell(row=ultima_linha, column=coluna)
                    if origem.has_style:
                        destino._style = copy(origem._style)
            linhas_por_ean[ean] = [proxima_linha]
            estado[ean] = 'nova'
            proxima_linha += 1

    formulas = _contar_formulas(conteudo_do_drive)
    saida = io.BytesIO()
    livro_do_drive.save(saida)
    estados = list(estado.values())
    return {
        'conteudo': saida.getvalue(), 'atualizadas': estados.count('atualizada'), 'novas': estados.count('nova'),
        'iguais': estados.count('igual'), 'sem_chave': sem_chave, 'formulas_viraram_valor': formulas,
    }
