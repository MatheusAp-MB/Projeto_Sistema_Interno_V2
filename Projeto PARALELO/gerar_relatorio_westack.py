# gerar_relatorio_westack.py
#
# Le o HTML salvo da tela "Catalogo da loja" do WeStack (com ou sem filtro)
# e gera um relatorio .xlsx numa ABA UNICA, agrupado por SKU, com TODOS os
# dados que o HTML traz de cada anuncio: tags, preco final e preco base,
# promocao, piso/teto, Auto, Catalogo, PMA, estoque, custo, margem (% e R$),
# status, alertas ("Perdendo R$/venda", "Sem piso", "PMA acima"), visitas,
# foto, e qualquer texto novo que a WeStack passe a mostrar (coluna
# "Outros dados do HTML"). Nao usa API nenhuma, so le o que ja esta no HTML.
#
# IMPORTANTE: a pagina carrega os anuncios de 20 em 20 ("80 de 1205 ·
# carregar mais 20"). O HTML salvo so tem o que ja apareceu na tela. Antes
# de salvar (Ctrl+S), role ate o fim / clique em "carregar mais" ate o
# contador chegar em "1205 de 1205". O script avisa se o HTML estiver
# incompleto.
#
# Limites (nao existem no HTML salvo, so aparecem ao clicar na tela):
#   - a janelinha "Detalhar margem e custo de entrada";
#   - o painel do botao de expandir (seta ao lado do status).
#
# Uso:
#   python gerar_relatorio_westack.py caminho\para\pagina.html
#   python gerar_relatorio_westack.py caminho\para\pagina.html -o relatorio.xlsx
#   python gerar_relatorio_westack.py caminho\para\pagina.html --conta Samvale
#
# Dependencias: pip install beautifulsoup4 openpyxl

from __future__ import annotations

import argparse
import re
import sys
from collections import defaultdict
from textwrap import wrap
from datetime import datetime
from pathlib import Path

from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# Linha monoespaçada "MLB4753881553 · F7898336431864.001 · 0 visitas/30d" (acima de mil vem "1.497 visitas/30d")
PADRAO_LINHA_MONO = re.compile(r"(MLB\d+)\s*·\s*([^\s·]+)(?:\s*·\s*(\d[\d.]*)\s*visitas?/30d)?")
PADRAO_DINHEIRO = re.compile(r"(-)?\s*R\$\s*(-)?\s*([\d.,]+)")
PADRAO_PERCENTUAL = re.compile(r"(-?\d+(?:[.,]\d+)?)\s*%")

SEM_SKU = "(sem SKU)"

FONTE = "Arial"

COR_TITULO = "1F2937"
COR_BORDA = "D1D5DB"
COR_GRUPO_FUNDO = "BFDBFE"
COR_ZEBRA = "F3F4F6"
COR_DIVISORIA_FAIXA = "4B5563"
COR_GRUPO_TEXTO = "1E3A8A"

# Faixas de cor do cabeçalho: (cor escura da faixa, cor clara dos títulos das colunas).
FAIXAS = {
    "Identificação": ("374151", "E5E7EB"),
    "Preço e promoção": ("1E40AF", "DBEAFE"),
    "Regras do robô": ("5B21B6", "EDE9FE"),
    "Estoque e custo": ("047857", "D1FAE5"),
    "Margem": ("B45309", "FEF3C7"),
    "Status e alertas": ("B91C1C", "FEE2E2"),
    "Textos completos": ("4B5563", "F3F4F6"),
}

# Cada coluna: (chave no dict do anúncio, título, largura mínima, largura máxima,
#               formato, alinhamento, faixa, quebra)
#   formato:     None | "int" | "dinheiro" | "pct"
#   alinhamento: "esq" | "centro" | "dir"   (números e R$ à direita, textos à esquerda)
#   quebra:      True = o texto quebra em várias linhas e a altura da linha cresce pra caber
# A largura real de cada coluna é calculada pelo maior valor dela, entre o mínimo e o máximo.
COLUNAS = [
    ("mlb", "MLB", 22, 22, None, "esq", "Identificação", False),
    ("sku", "SKU", 20, 26, None, "esq", "Identificação", False),
    ("titulo", "Título do anúncio", 40, 60, None, "esq", "Identificação", True),
    ("tipo", "Tipo", 12, 14, None, "centro", "Identificação", False),
    ("logistica", "Logística", 11, 14, None, "centro", "Identificação", False),
    ("codigo_logistica", "Código da logística", 14, 40, None, "centro", "Identificação", False),
    ("frete_gratis", "Frete grátis", 10, 12, None, "centro", "Identificação", False),
    ("outras_tags", "Outras tags", 14, 24, None, "centro", "Identificação", True),
    ("visitas", "Visitas/30d", 10, 12, "int", "dir", "Identificação", False),
    ("foto", "Arquivo da foto", 24, 46, None, "esq", "Identificação", False),
    ("renovar", "Botão Renovar", 10, 12, None, "centro", "Identificação", False),

    ("preco_final", "Preço final", 12, 16, "dinheiro", "dir", "Preço e promoção", False),
    ("preco_base", "Preço base (riscado)", 13, 16, "dinheiro", "dir", "Preço e promoção", False),
    ("promocao", "Promoção", 16, 28, None, "centro", "Preço e promoção", True),
    ("preco_promocao", "Preço da promoção", 13, 16, "dinheiro", "dir", "Preço e promoção", False),

    ("piso", "Piso (margem mín.)", 11, 14, "pct", "dir", "Regras do robô", False),
    ("teto", "Teto (margem máx.)", 11, 14, "pct", "dir", "Regras do robô", False),
    ("auto", "Auto", 10, 12, None, "centro", "Regras do robô", False),
    ("auto_obs", "Auto — observação", 22, 32, None, "esq", "Regras do robô", True),
    ("prioriza_catalogo", "Prioriza catálogo", 11, 14, None, "centro", "Regras do robô", False),
    ("situacao_catalogo", "Situação no catálogo", 13, 24, None, "centro", "Regras do robô", False),
    ("pma", "PMA", 10, 12, None, "centro", "Regras do robô", False),
    ("valor_pma", "Valor PMA", 11, 14, "dinheiro", "dir", "Regras do robô", False),

    ("estoque", "Estoque", 10, 12, "int", "dir", "Estoque e custo", False),
    ("custo", "Custo", 12, 16, "dinheiro", "dir", "Estoque e custo", False),
    ("custo_manual", "Custo travado (manual)", 12, 14, None, "centro", "Estoque e custo", False),

    ("margem_pct", "Margem líq. %", 12, 14, "pct", "dir", "Margem", False),
    ("margem_rs", "Margem líq. R$", 12, 16, "dinheiro", "dir", "Margem", False),
    ("margem_obs", "Margem — observação", 13, 18, None, "centro", "Margem", True),

    ("status", "Status", 12, 20, None, "centro", "Status e alertas", False),
    ("status_obs", "Status — observação", 13, 20, None, "centro", "Status e alertas", True),
    ("sem_piso", "Alerta: sem piso", 10, 12, None, "centro", "Status e alertas", False),
    ("motor_inerte", "Alerta: motor inerte", 10, 12, None, "centro", "Status e alertas", False),
    ("pma_desvio", "Alerta PMA: diferença R$", 13, 16, "dinheiro", "dir", "Status e alertas", False),
    ("pma_preco_efetivo", "Alerta PMA: preço efetivo", 13, 16, "dinheiro", "dir", "Status e alertas", False),
    ("perda_campanha", "Perda: campanha do ML", 16, 28, None, "centro", "Status e alertas", True),
    ("perda_preco_campanha", "Perda: preço na campanha", 13, 16, "dinheiro", "dir", "Status e alertas", False),
    ("perda_preco_bruto", "Perda: preço bruto", 12, 16, "dinheiro", "dir", "Status e alertas", False),
    ("perda_custo", "Perda: custo", 11, 16, "dinheiro", "dir", "Status e alertas", False),
    ("perda_por_venda", "Perda por venda", 12, 16, "dinheiro", "dir", "Status e alertas", False),
    ("perda_vendas_30d", "Perda: vendas 30d", 10, 12, "int", "dir", "Status e alertas", False),
    ("perda_total_30d", "Perda total 30d", 12, 16, "dinheiro", "dir", "Status e alertas", False),

    ("textos_completos", "Textos completos das tags e do status", 50, 90, None, "esq", "Textos completos", True),
    ("outros", "Outros dados do HTML", 30, 60, None, "esq", "Textos completos", True),
]

# Cores do Status (fundo, texto) — iguais às da tela: verde / amarelo / vermelho.
CORES_STATUS = {
    "Saudável": ("C6EFCE", "006100"),
    "Atenção": ("FFEB9C", "7F6000"),
    "Crítico": ("FFC7CE", "9C0006"),
    "Imposto pendente": ("E9D5FF", "6B21A8"),
}

FORMATOS = {"int": "0", "dinheiro": '"R$" #,##0.00', "pct": "0.0%"}

# Textos fixos da interface (botões, rótulos, dicas que só explicam o controle).
# Nao sao dados do anúncio, então não entram em "Outros dados do HTML".
RUIDO = {
    "", "—", "-", "%", "R$", "!", "▲", "▼",
    "Preço final", "Estoque", "Custo", "Margem líq.", "Automático", "PMA", "Catálogo",
    "Valor PMA", "Renovar",
    "Selecionar anúncio", "Expandir", "Ver mais detalhes",
    "Alterar preço base", "Alterar preço base no Mercado Livre",
    "Preço base no Mercado Livre (o preço com desconto segue pela promoção ativa).",
    "Alterar preço de custo",
    "Ativar entrada automática em promoção",
    "Preço mínimo anunciado neste anúncio",
    "Priorizar ganho de catálogo neste anúncio",
    "Detalhar margem e custo de entrada",
    "Cria um anúncio novo (cópia fiel + textos reescritos pela IA).",
}
RUIDO_PREFIXOS = (
    "Margem líquida MÍNIMA",
    "Margem líquida MÁXIMA",
    "Priorizar catálogo",
    "PMA desligado",
    "PMA ligado",
)


# ---------------------------------------------------------------- utilitários

def _norm(texto) -> str:
    """
    Tira espaços repetidos, o espaço 'não separável' (\\xa0) que o WeStack usa no R$
    e caracteres de controle que o Excel não aceita.
    """
    texto = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", texto or "")
    return re.sub(r"\s+", " ", texto.replace("\xa0", " ")).strip()


def _numero(texto):
    """
    Converte '1.234,56' (tela, padrão BR), '44.08' (dicas, ponto decimal) ou
    '65.00' (campos) em float. Devolve None se não achar número.
    """
    s = re.sub(r"[^\d.,-]", "", _norm(texto))
    negativo = "-" in s
    s = s.replace("-", "").strip(".,")
    if not s:
        return None
    if "," in s and "." in s:
        if s.rfind(",") > s.rfind("."):          # 1.234,56
            s = s.replace(".", "").replace(",", ".")
        else:                                     # 1,234.56
            s = s.replace(",", "")
    elif "," in s:                                # 198,94
        s = s.replace(",", ".")
    try:
        valor = float(s)
    except ValueError:
        return None
    return -valor if negativo else valor


def _dinheiro(texto):
    m = PADRAO_DINHEIRO.search(_norm(texto))
    if not m:
        return None
    valor = _numero(m.group(3))
    if valor is None:
        return None
    return -valor if (m.group(1) or m.group(2)) else valor


def _fracao(texto):
    """'14.1%' -> 0.141 (o Excel guarda percentual como fração)."""
    m = PADRAO_PERCENTUAL.search(_norm(texto))
    if not m:
        return None
    valor = _numero(m.group(1))
    return None if valor is None else round(valor / 100, 6)


def _fracao_de_numero(texto):
    """'14.0' (valor de um campo, sem o sinal de %) -> 0.14."""
    valor = _numero(texto)
    return None if valor is None else round(valor / 100, 6)


def _inteiro(texto):
    digitos = re.sub(r"[^\d]", "", _norm(texto))
    return int(digitos) if digitos else None


def _sim_nao(valor):
    if valor is None:
        return None
    return "Sim" if valor else "Não"


def _ligado(valor):
    if valor is None:
        return None
    return "Ligado" if valor else "Desligado"


def _celula_do_rotulo(row, rotulo):
    """
    Acha o rótulo pequeno (classe md:hidden, existe no HTML mesmo escondido no
    desktop) e devolve (célula que o contém, textos que vêm depois dele).
    """
    for el in row.find_all(["p", "span"], class_="md:hidden"):
        if _norm(el.get_text()) == rotulo:
            celula = el.parent
            textos = [_norm(t) for t in celula.stripped_strings]
            if textos and textos[0] == rotulo:
                textos = textos[1:]
            return celula, textos
    return None, []


def _tooltip_do_toggle(botao):
    """O 'title' do toggle (Auto / PMA / Catálogo) fica na div que envolve o botão."""
    dono = botao.find_parent(attrs={"title": True})
    return _norm(dono.get("title")) if dono is not None else ""


# ----------------------------------------------------------- leitura das tags

def _ler_tags(bloco_produto, usar):
    """
    Lê as tags do topo do bloco do produto e separa pelo que cada uma É
    (não pela posição, porque as tags variam de anúncio pra anúncio).
    """
    r = {
        "tipo": None, "logistica": None, "codigo_logistica": None, "frete_gratis": "Não",
        "promocoes": [], "outras_tags": [], "textos_completos": [],
        "sem_piso": "Não", "pma_desvio": None, "pma_preco_efetivo": None,
        "perda_campanha": None, "perda_preco_campanha": None, "perda_preco_bruto": None,
        "perda_custo": None, "perda_por_venda": None, "perda_vendas_30d": None,
        "perda_total_30d": None,
    }
    tags = bloco_produto.select("div.flex.items-center.flex-wrap > span")
    for i, tag in enumerate(tags):
        texto = usar(tag.get_text(" ", strip=True))
        titulo = usar(tag.get("title"))
        baixo = texto.lower()

        if titulo.startswith("Logística:"):
            r["logistica"] = texto
            r["codigo_logistica"] = titulo.split("Logística:", 1)[1].strip() or None
        elif baixo == "frete grátis":
            r["frete_gratis"] = "Sim"
        elif titulo.startswith("Promoção:"):
            sem_mudanca = "sem mudança de preço" in titulo
            m = re.search(r"preço final\s*(R\$\s*[\d.,]+)", titulo)
            m_nome = re.match(r"Promoção:\s*(.+?)\s*·\s*(?:preço final|sem mudança)", titulo)
            nome_promo = m_nome.group(1) if m_nome else texto   # a tag corta nomes longos com "…"
            r["promocoes"].append((nome_promo, _dinheiro(m.group(1)) if m else None, sem_mudanca))
            r["textos_completos"].append(f"[{texto}] {titulo}")
        elif baixo.startswith("perdendo"):
            r["textos_completos"].append(f"[{texto}] {titulo}")
            m = re.search(r"numa campanha (.+?) a (R\$\s*[\d.,]+?)\.\s", titulo + " ")
            if m:
                r["perda_campanha"] = m.group(1).strip()
                r["perda_preco_campanha"] = _dinheiro(m.group(2))
            m = re.search(r"Preço bruto (R\$\s*[\d.,]+?)[,.]\s", titulo + " ")
            r["perda_preco_bruto"] = _dinheiro(m.group(1)) if m else None
            m = re.search(r"custo (R\$\s*[\d.,]+?)[,.]\s", titulo + " ")
            r["perda_custo"] = _dinheiro(m.group(1)) if m else None
            m = re.search(r"Perda de (R\$\s*[\d.,]+?) por venda", titulo)
            r["perda_por_venda"] = _dinheiro(m.group(1)) if m else _dinheiro(texto)
            m = re.search(r"×\s*(\d+)\s*vendas? em 30d\s*=\s*(R\$\s*[\d.,]+?)\s", titulo + " ")
            if m:
                r["perda_vendas_30d"] = int(m.group(1))
                r["perda_total_30d"] = _dinheiro(m.group(2))
        elif baixo.startswith("sem piso"):
            r["sem_piso"] = "Sim"
            r["textos_completos"].append(f"[{texto}] {titulo}")
        elif baixo.startswith("pma") and re.search(r"\b(acima|abaixo)\b", baixo):
            r["textos_completos"].append(f"[{texto}] {titulo}")
            m = re.search(
                r"Preço efetivo (R\$\s*[\d.,]+?) está (R\$\s*[\d.,]+?) (acima|abaixo) do PMA", titulo)
            if m:
                r["pma_preco_efetivo"] = _dinheiro(m.group(1))
                diferenca = _dinheiro(m.group(2))
                if diferenca is not None:
                    r["pma_desvio"] = diferenca if m.group(3) == "acima" else -diferenca
            else:
                r["pma_desvio"] = _dinheiro(texto)
        elif i == 0 and not titulo:
            r["tipo"] = texto                      # "Tradicional" / "Catálogo"
        else:
            r["outras_tags"].append(texto)
            if titulo:
                r["textos_completos"].append(f"[{texto}] {titulo}")
    return r


# ------------------------------------------------------------ leitura da linha

def extrair_anuncio(row):
    """
    Lê TODOS os dados de uma linha (1 anúncio/MLB) do grid do WeStack.
    Devolve o dict do anúncio, ou None se a linha não bater com o padrão.
    """
    bloco = row.select_one("div.min-w-0.flex-1")
    if bloco is None:
        return None
    divs = bloco.find_all("div")
    if not divs:
        return None

    usados = set()     # tudo que foi lido vira "usado" — o que sobrar vai pra "Outros"

    def usar(texto):
        t = _norm(texto)
        if t:
            usados.add(t)
        return t

    # --- Produto: MLB · SKU · visitas (última div do bloco) -------------------
    linha_mono = usar(divs[-1].get_text(" ", strip=True))
    m = PADRAO_LINHA_MONO.search(linha_mono)
    if not m:
        return None
    mlb = m.group(1)
    sku = None if m.group(2) in ("—", "-") else m.group(2)
    visitas = _inteiro(m.group(3)) if m.group(3) else None

    # Título vem no atributo title= (evita o texto cortado por CSS).
    titulo_div = bloco.find("div", title=True)
    titulo = ""
    if titulo_div is not None:
        titulo = (titulo_div.get("title") or "").strip()
        usar(titulo)
        usar(titulo_div.get_text(" ", strip=True))

    # Foto: só o nome do arquivo (o HTML guarda o caminho da pasta _files).
    img = row.find("img")
    foto = None
    if img is not None and img.get("src"):
        foto = Path(usar(img.get("src"))).name

    tags = _ler_tags(bloco, usar)

    botao_renovar = None
    for b in row.find_all("button"):
        if _norm(b.get_text()) == "Renovar":
            botao_renovar = b
            break

    # --- Preço final e preço base (riscado) ----------------------------------
    celula_preco, textos_preco = _celula_do_rotulo(row, "Preço final")
    preco_final = _dinheiro(textos_preco[0]) if textos_preco else None
    for t in textos_preco:
        usar(t)
    if preco_final is None:
        p_antigo = row.select_one("p.font-semibold")
        if p_antigo is not None:
            preco_final = _dinheiro(usar(p_antigo.get_text(" ", strip=True)))
    preco_base = None
    if celula_preco is not None:
        riscado = celula_preco.select_one("span.line-through")
        if riscado is not None:
            preco_base = _dinheiro(riscado.get_text())
        for el in celula_preco.find_all(attrs={"title": True}):
            usar(el.get("title"))

    # --- Piso / Teto  ou  Valor PMA ------------------------------------------
    piso = teto = valor_pma = None
    for campo in row.find_all("input"):
        tit = _norm(campo.get("title"))
        valor = usar(campo.get("value"))
        if tit.startswith("Margem líquida MÍNIMA"):
            piso = _fracao_de_numero(valor) if valor else None
        elif tit.startswith("Margem líquida MÁXIMA"):
            teto = _fracao_de_numero(valor) if valor else None
        elif tit.startswith("PMA"):
            valor_pma = _numero(valor) if valor else None

    # --- Toggles: Auto, Catálogo, PMA ----------------------------------------
    auto = auto_obs = None
    motor_inerte = False
    celula_auto = None
    botao = row.find("button", attrs={"aria-label": "Ativar entrada automática em promoção"})
    if botao is not None:
        auto = botao.get("aria-checked") == "true"
        dica = usar(_tooltip_do_toggle(botao))
        if dica and not dica.startswith("Automático — margem"):
            auto_obs = dica
        if dica.startswith("Motor inerte"):
            motor_inerte = True
            usar("Motor inerte")
        celula_auto = botao
        while celula_auto.parent is not None and celula_auto.parent is not row:
            celula_auto = celula_auto.parent

    prioriza = None
    botao = row.find("button", attrs={"aria-label": "Priorizar ganho de catálogo neste anúncio"})
    if botao is not None:
        prioriza = botao.get("aria-checked") == "true"

    # A célula do Catálogo é a que vem logo depois da célula do Auto. O indicador dela
    # (▲ ▼ —) traz o motivo em aria-label: "Ganhando o catálogo", "Perdendo o catálogo",
    # "Sem disputa de catálogo", "Dividindo 1º lugar", "Fora do buy box"...
    situacao_catalogo = None
    celula_catalogo = celula_auto.find_next_sibling("div") if celula_auto is not None else None
    area_catalogo = celula_catalogo if celula_catalogo is not None else row
    for el in area_catalogo.find_all(attrs={"aria-label": True}):
        rotulo = _norm(el.get("aria-label"))
        if rotulo in ("Priorizar ganho de catálogo neste anúncio",
                      "Ativar entrada automática em promoção",
                      "Preço mínimo anunciado neste anúncio"):
            continue
        if rotulo in ("Ganhando o catálogo", "Perdendo o catálogo", "Sem disputa de catálogo",
                      "Dividindo 1º lugar", "Fora do buy box") or celula_catalogo is not None:
            usar(rotulo)
            situacao_catalogo = {"Ganhando o catálogo": "Ganhando",
                                 "Perdendo o catálogo": "Perdendo",
                                 "Sem disputa de catálogo": "Sem disputa"}.get(rotulo, rotulo)
            break

    pma = None
    botao = row.find("button", attrs={"aria-label": "Preço mínimo anunciado neste anúncio"})
    if botao is not None:
        pma = botao.get("aria-checked") == "true"

    # --- Estoque, custo, margem ----------------------------------------------
    _, textos = _celula_do_rotulo(row, "Estoque")
    estoque = _inteiro(textos[0]) if textos else None
    for t in textos:
        usar(t)

    celula_custo, textos = _celula_do_rotulo(row, "Custo")
    custo = _dinheiro(textos[0]) if textos else None
    for t in textos:
        usar(t)
    custo_manual = None
    if celula_custo is not None:
        custo_manual = celula_custo.find(attrs={"aria-label": "Custo fixado manualmente"}) is not None
        if custo_manual:
            usar("Custo fixado manualmente")

    _, textos = _celula_do_rotulo(row, "Margem líq.")
    margem_pct = margem_rs = None
    margem_obs = []
    for t in textos:
        usar(t)
        if "%" in t and margem_pct is None:
            margem_pct = _fracao(t)
        elif "R$" in t and margem_rs is None:
            margem_rs = _dinheiro(t)
        else:
            margem_obs.append(t)      # ex.: "Pendente" (imposto pendente, sem margem calculada)

    # --- Status (célula logo antes do botão de expandir) ----------------------
    status = status_obs = None
    botao_exp = row.find("button", attrs={"aria-label": "Expandir"})
    celula_status = botao_exp.find_previous_sibling("div") if botao_exp is not None else None
    if celula_status is not None:
        textos = [usar(t) for t in celula_status.stripped_strings]
        textos = [t for t in textos if t != "!"]
        status = textos[0] if textos else None
        status_obs = " | ".join(textos[1:]) or None      # ex.: "Definir custo"
        for el in celula_status.find_all(attrs={"title": True}):
            dica = usar(el.get("title"))
            tags["textos_completos"].append(f"[Status {status}] {dica}")

    promocoes = tags["promocoes"]
    nomes_promo = []
    for nome, preco, sem_mudanca in promocoes:
        nomes_promo.append(f"{nome} (sem mudança de preço)" if sem_mudanca else nome)
    preco_promocao = promocoes[0][1] if len(promocoes) == 1 else None

    # --- Tudo que sobrou no HTML e não foi lido acima --------------------------
    achados = [_norm(t) for t in row.stripped_strings]
    for el in [row] + row.find_all(True):
        for atributo in ("title", "aria-label", "alt", "value"):
            v = el.get(atributo)
            if isinstance(v, str):
                achados.append(_norm(v))
        for atributo, v in el.attrs.items():
            if atributo.startswith("data-") and isinstance(v, str) and v.strip():
                achados.append(f"{atributo}={_norm(v)}")
        if el.name == "a" and el.get("href"):
            achados.append(f"link={el.get('href')}")
    outros = []
    for t in achados:
        if t in usados or t in RUIDO or t.startswith(RUIDO_PREFIXOS) or t in outros:
            continue
        outros.append(t)

    return {
        "mlb": mlb,
        "sku": sku,
        "titulo": titulo,
        "tipo": tags["tipo"],
        "logistica": tags["logistica"],
        "codigo_logistica": tags["codigo_logistica"],
        "frete_gratis": tags["frete_gratis"],
        "outras_tags": " | ".join(tags["outras_tags"]) or None,
        "visitas": visitas,
        "foto": foto,
        "renovar": _sim_nao(botao_renovar is not None),

        "preco_final": preco_final,
        "preco_base": preco_base,
        "promocao": " | ".join(nomes_promo) or None,
        "preco_promocao": preco_promocao,

        "piso": piso,
        "teto": teto,
        "auto": _ligado(auto),
        "auto_obs": auto_obs,
        "prioriza_catalogo": _sim_nao(prioriza),
        "situacao_catalogo": situacao_catalogo,
        "pma": _ligado(pma),
        "valor_pma": valor_pma,

        "estoque": estoque,
        "custo": custo,
        "custo_manual": _sim_nao(custo_manual),

        "margem_pct": margem_pct,
        "margem_rs": margem_rs,
        "margem_obs": " | ".join(margem_obs) or None,

        "status": status,
        "status_obs": status_obs,
        "sem_piso": tags["sem_piso"],
        "motor_inerte": _sim_nao(motor_inerte),
        "pma_desvio": tags["pma_desvio"],
        "pma_preco_efetivo": tags["pma_preco_efetivo"],
        "perda_campanha": tags["perda_campanha"],
        "perda_preco_campanha": tags["perda_preco_campanha"],
        "perda_preco_bruto": tags["perda_preco_bruto"],
        "perda_custo": tags["perda_custo"],
        "perda_por_venda": tags["perda_por_venda"],
        "perda_vendas_30d": tags["perda_vendas_30d"],
        "perda_total_30d": tags["perda_total_30d"],

        "textos_completos": " || ".join(tags["textos_completos"]) or None,
        "outros": " | ".join(outros) or None,
    }


def extrair_linhas(soup: BeautifulSoup) -> tuple[list[dict], int]:
    """
    Percorre cada 'linha' (1 anúncio/MLB) do grid do WeStack. Retorna os
    anúncios lidos e a contagem de linhas encontradas mas que não bateram
    com o padrão esperado (pra avisar, nunca falhar silenciosamente).
    """
    linhas_ok = []
    nao_reconhecidas = 0
    for row in soup.select("div.group.grid.items-center.transition-colors"):
        anuncio = extrair_anuncio(row)
        if anuncio is None:
            nao_reconhecidas += 1
        else:
            linhas_ok.append(anuncio)
    return linhas_ok, nao_reconhecidas


# ------------------------------------------------------- dados do topo da tela

def _int_ou_none(texto):
    return int(texto.replace(".", "")) if texto and texto.replace(".", "").isdigit() else None


def extrair_contexto_pagina(soup: BeautifulSoup) -> dict:
    """
    Lê os contadores do topo/rodapé da tela (total de anúncios, saúde do
    catálogo, alertas, 'N de M carregados'). Chamar DEPOIS de tirar as linhas
    do soup, senão o texto dos anúncios atrapalha.
    """
    for tag in soup(["script", "style", "svg", "noscript"]):
        tag.decompose()
    textos = [_norm(t) for t in soup.stripped_strings]
    junto = " ".join(textos)
    ctx = {}

    m = re.search(r"(\d[\d.]*)\s*anúncios ativos\s*·\s*(\d[\d.]*)\s*catálogo\s*·\s*(\d[\d.]*)\s*tradicional", junto)
    if m:
        ctx["ativos"], ctx["catalogo"], ctx["tradicional"] = (_int_ou_none(g) for g in m.groups())

    def contador(rotulo):
        for i, t in enumerate(textos[:-1]):
            if t == rotulo and textos[i + 1].replace(".", "").isdigit():
                return _int_ou_none(textos[i + 1])
        return None

    for chave, rotulo in (("saudavel", "Saudável"), ("atencao", "Atenção"), ("critico", "Crítico"),
                          ("aguardando", "Aguardando aprovação"), ("motor_parado", "Motor parado"),
                          ("sem_custo", "Sem custo")):
        valor = contador(rotulo)
        if valor is not None:
            ctx[chave] = valor

    for t in textos:
        m = re.match(r"^(\d[\d.]*)\s+de\s+(\d[\d.]*)\b", t)
        if m:
            ctx["carregados"], ctx["total_filtro"] = _int_ou_none(m.group(1)), _int_ou_none(m.group(2))
            break
    return ctx


def texto_contexto(ctx: dict) -> list[str]:
    """Monta as linhas de resumo do topo da planilha a partir do contexto lido."""
    linhas = []
    if "ativos" in ctx:
        linhas.append(f"Na tela: {ctx['ativos']} anúncios ativos  ·  {ctx.get('catalogo', '?')} catálogo  ·  "
                      f"{ctx.get('tradicional', '?')} tradicional")
    saude = [f"{nome} {ctx[k]}" for k, nome in (("saudavel", "Saudável"), ("atencao", "Atenção"),
                                               ("critico", "Crítico")) if k in ctx]
    if saude:
        linhas.append("Saúde do catálogo (total da loja): " + "  ·  ".join(saude))
    alertas = [f"{nome} {ctx[k]}" for k, nome in (("aguardando", "Aguardando aprovação"),
                                                 ("motor_parado", "Motor parado"),
                                                 ("sem_custo", "Sem custo")) if k in ctx]
    if alertas:
        linhas.append("Alertas (total da loja): " + "  ·  ".join(alertas))
    return linhas


def html_incompleto(ctx: dict, lidos: int) -> str | None:
    """Se a tela diz 'N de M' e N < M, o HTML salvo não tem todos os anúncios."""
    total = ctx.get("total_filtro")
    if total is not None and lidos < total:
        return (f"HTML INCOMPLETO: tem {lidos} anúncios, mas a tela mostra {total}. "
                f"Role a página até o fim (ou clique em 'carregar mais') até o contador "
                f"chegar em '{total} de {total}' e salve o HTML de novo.")
    return None


# ------------------------------------------------------------------ planilha

def agrupar_por_sku(linhas: list[dict]) -> dict[str, list[dict]]:
    grupos = defaultdict(list)
    for linha in linhas:
        grupos[linha["sku"] or SEM_SKU].append(linha)
    return dict(sorted(grupos.items()))


HORIZONTAL = {"esq": "left", "centro": "center", "dir": "right"}


def _exibicao(valor, formato) -> str:
    """Como o valor aparece na célula (só pra medir o tamanho da coluna)."""
    if valor is None:
        return ""
    if formato == "dinheiro":
        return f"R$ {valor:,.2f}"
    if formato == "pct":
        return f"{valor * 100:.1f}%"
    return str(valor)


def _linhas_do_texto(texto, largura, fator=1.05) -> int:
    """
    Quantas linhas o texto ocupa numa coluna dessa largura (estimativa com folga
    pro recuo). Texto em negrito (cabeçalho) é mais largo, então usa fator menor.
    """
    if not texto:
        return 1
    capacidade = max(int(largura * fator) - 2, 1)
    return sum(max(len(wrap(trecho, capacidade)), 1) for trecho in str(texto).split("\n"))


def calcular_larguras(grupos: dict[str, list[dict]]) -> list[float]:
    """Largura de cada coluna = maior valor dela (ou maior palavra do título), entre o mínimo e o máximo."""
    itens = [item for itens_do_grupo in grupos.values() for item in itens_do_grupo]
    larguras = []
    for chave, titulo, largura_min, largura_max, formato, _, _, _ in COLUNAS:
        maior_valor = max((len(_exibicao(i.get(chave), formato)) for i in itens), default=0)
        maior_palavra = max((len(p) for p in titulo.split()), default=0)
        larguras.append(min(max(maior_valor + 3, maior_palavra + 3, largura_min), largura_max))
    return larguras


def gerar_xlsx(grupos: dict[str, list[dict]], nome_conta: str,
               total_linhas: int, ctx: dict, aviso: str | None):
    fonte_titulo = Font(name=FONTE, size=14, bold=True, color=COR_TITULO)
    fonte_subtitulo = Font(name=FONTE, size=10, italic=True, color="6B7280")
    fonte_aviso = Font(name=FONTE, size=10, bold=True, color="B91C1C")
    fonte_faixa = Font(name=FONTE, size=10, bold=True, color="FFFFFF")
    fonte_cabecalho = Font(name=FONTE, size=10, bold=True, color=COR_TITULO)
    fonte_grupo_sku = Font(name=FONTE, size=11, bold=True, color=COR_GRUPO_TEXTO)
    fonte_grupo_info = Font(name=FONTE, size=10, italic=True, color="3B5488")
    fonte_dado = Font(name=FONTE, size=10, color="111827")

    fill_grupo = PatternFill("solid", fgColor=COR_GRUPO_FUNDO)
    fill_zebra = PatternFill("solid", fgColor=COR_ZEBRA)

    # Grade: toda célula tem as 4 bordas; a 1ª coluna de cada faixa tem divisória mais grossa.
    lado = Side(style="thin", color=COR_BORDA)
    lado_divisoria = Side(style="medium", color=COR_DIVISORIA_FAIXA)
    lado_grupo = Side(style="thin", color="60A5FA")

    n_cols = len(COLUNAS)
    ultima_letra = get_column_letter(n_cols)
    inicios_de_faixa = {c for c in range(1, n_cols + 1)
                        if c == 1 or COLUNAS[c - 1][6] != COLUNAS[c - 2][6]}

    def borda_dado(c, cabecalho=False):
        return Border(
            left=lado_divisoria if c in inicios_de_faixa else lado,
            right=lado_divisoria if c == n_cols else lado,
            top=lado,
            bottom=Side(style="medium", color=COR_TITULO) if cabecalho else lado,
        )

    larguras = calcular_larguras(grupos)

    wb = Workbook()
    ws = wb.active
    ws.title = "Catálogo WeStack"

    for i, largura in enumerate(larguras, start=1):
        ws.column_dimensions[get_column_letter(i)].width = largura

    # --- Topo: título, resumo, avisos (sem mesclar: o texto transborda pra direita) ---
    linha = 1
    ws.cell(row=linha, column=1, value=f"WeStack — Catálogo da loja ({nome_conta})").font = fonte_titulo
    ws.row_dimensions[linha].height = 22
    linha += 1

    resumo = (f"Gerado em {datetime.now().strftime('%d/%m/%Y %H:%M')}  ·  "
              f"{len(grupos)} SKUs  ·  {total_linhas} MLBs neste arquivo")
    if ctx.get("total_filtro") is not None:
        resumo += f" (a tela mostrava {ctx['carregados']} de {ctx['total_filtro']})"
    notas = [resumo]
    contexto = texto_contexto(ctx)
    if contexto:
        notas.append("   |   ".join(contexto))
    notas.append("Quando o PMA está ligado, a WeStack não mostra piso/teto (o robô ignora) — as células ficam vazias.")
    for texto in notas:
        ws.cell(row=linha, column=1, value=texto).font = fonte_subtitulo
        linha += 1
    if aviso:
        ws.cell(row=linha, column=1, value=aviso).font = fonte_aviso
        linha += 1
    linha += 1     # respiro antes da tabela

    # --- Faixas coloridas (1 por tipo de dado) + títulos das colunas -----------
    linha_faixa = linha
    linha_cabecalho = linha + 1
    inicio = 1
    while inicio <= n_cols:
        faixa = COLUNAS[inicio - 1][6]
        fim = inicio
        while fim < n_cols and COLUNAS[fim][6] == faixa:
            fim += 1
        cor_escura, _ = FAIXAS[faixa]
        for c in range(inicio, fim + 1):
            ws.cell(row=linha_faixa, column=c).fill = PatternFill("solid", fgColor=cor_escura)
        ws.merge_cells(start_row=linha_faixa, start_column=inicio, end_row=linha_faixa, end_column=fim)
        cel = ws.cell(row=linha_faixa, column=inicio, value=faixa.upper())
        cel.font = fonte_faixa
        cel.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        inicio = fim + 1
    ws.row_dimensions[linha_faixa].height = 20

    linhas_cabecalho = 1
    for c, col in enumerate(COLUNAS, start=1):
        _, cor_clara = FAIXAS[col[6]]
        cel = ws.cell(row=linha_cabecalho, column=c, value=col[1])
        cel.font = fonte_cabecalho
        cel.fill = PatternFill("solid", fgColor=cor_clara)
        cel.alignment = Alignment(horizontal="left" if col[5] == "esq" else "center",
                                  vertical="center", wrap_text=True,
                                  indent=1 if col[5] == "esq" else 0)
        cel.border = borda_dado(c, cabecalho=True)
        linhas_cabecalho = max(linhas_cabecalho, _linhas_do_texto(col[1], larguras[c - 1], fator=1.0))
    ws.row_dimensions[linha_cabecalho].height = max(34, 14 * linhas_cabecalho + 8)

    linha_atual = linha_cabecalho + 1
    ws.freeze_panes = f"C{linha_atual}"     # trava cabeçalho + colunas MLB e SKU

    # --- Dados: 1 faixa de título por SKU + 1 linha por MLB --------------------
    n_dado = 0
    for sku, itens in grupos.items():
        for c in range(1, n_cols + 1):
            cel = ws.cell(row=linha_atual, column=c)
            cel.fill = fill_grupo
            cel.border = Border(top=lado_grupo, bottom=lado_grupo,
                                left=lado_divisoria if c == 1 else None,
                                right=lado_divisoria if c == n_cols else None)
        # Coluna A: SKU sozinho, sem rótulo nem sufixo — pra copiar/colar direto.
        cel_sku = ws.cell(row=linha_atual, column=1, value=sku)
        cel_sku.font = fonte_grupo_sku
        cel_sku.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        # Contexto na coluna do título (sem mesclar, pra não atrapalhar filtro/ordenação).
        rotulo_mlb = "MLB" if len(itens) == 1 else "MLBs"
        cel_info = ws.cell(row=linha_atual, column=3, value=f"{len(itens)} {rotulo_mlb} neste SKU")
        cel_info.font = fonte_grupo_info
        cel_info.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws.row_dimensions[linha_atual].height = 24
        linha_atual += 1

        for item in itens:
            n_dado += 1
            cor_status = CORES_STATUS.get(item.get("status"))
            linhas_da_linha = 1
            for c, col in enumerate(COLUNAS, start=1):
                chave, _, _, _, formato, alinhamento, _, quebra = col
                valor = item.get(chave)
                if chave == "sku" and not valor:
                    valor = sku
                cel = ws.cell(row=linha_atual, column=c, value=valor)
                if isinstance(valor, str) and valor.startswith("="):
                    cel.data_type = "s"     # texto que começa com "=" não pode virar fórmula
                cel.font = fonte_dado
                cel.border = borda_dado(c)
                cel.alignment = Alignment(horizontal=HORIZONTAL[alinhamento], vertical="center",
                                          wrap_text=quebra,
                                          indent=0 if alinhamento == "centro" else 1)
                if formato:
                    cel.number_format = FORMATOS[formato]
                if n_dado % 2 == 0:
                    cel.fill = fill_zebra
                if cor_status and chave == "status":
                    cel.fill = PatternFill("solid", fgColor=cor_status[0])
                    cel.font = Font(name=FONTE, size=10, bold=True, color=cor_status[1])
                elif cor_status and chave in ("margem_pct", "margem_rs"):
                    cel.font = Font(name=FONTE, size=10, bold=True, color=cor_status[1])
                if quebra and valor:
                    linhas_da_linha = max(linhas_da_linha, _linhas_do_texto(valor, larguras[c - 1]))
            ws.row_dimensions[linha_atual].height = min(max(24, 13.5 * linhas_da_linha + 8), 220)
            linha_atual += 1

    ws.auto_filter.ref = f"A{linha_cabecalho}:{ultima_letra}{linha_atual - 1}"
    ws.sheet_view.showGridLines = False
    return wb


def salvar_planilha(wb, caminho_saida: Path) -> Path:
    """
    Salva o .xlsx. Se o arquivo estiver aberto no Excel (o Windows trava ele),
    salva com a hora no nome em vez de falhar.
    """
    try:
        wb.save(caminho_saida)
        return caminho_saida
    except PermissionError:
        alternativo = caminho_saida.with_name(
            f"{caminho_saida.stem}_{datetime.now().strftime('%H%M%S')}{caminho_saida.suffix}")
        wb.save(alternativo)
        print(f"AVISO: '{caminho_saida.name}' está aberto no Excel (ou sem permissão). "
              f"Salvei com outro nome.")
        return alternativo


def main():
    parser = argparse.ArgumentParser(
        description="Gera relatório Excel (aba única, agrupado por SKU, com todos os dados) a partir do HTML do WeStack.")
    parser.add_argument("html", type=Path, help="Caminho do HTML salvo da tela do WeStack.")
    parser.add_argument("-o", "--saida", type=Path, default=None, help="Caminho do .xlsx de saída (default: mesmo nome do HTML, extensão .xlsx).")
    parser.add_argument("--conta", type=str, default="Samvale", help="Nome da conta, só pro cabeçalho do relatório (default: Samvale).")
    args = parser.parse_args()

    if not args.html.exists():
        print(f"Arquivo não encontrado: {args.html}", file=sys.stderr)
        sys.exit(1)

    caminho_saida = args.saida or args.html.with_suffix(".xlsx")

    print("Lendo o HTML (com 1000+ anúncios pode levar alguns segundos)...")
    with open(args.html, encoding="utf-8") as f:
        soup = BeautifulSoup(f.read(), "html.parser")

    linhas, nao_reconhecidas = extrair_linhas(soup)
    if not linhas:
        print("Nenhuma linha de anúncio reconhecida no HTML — confira se é a página certa.", file=sys.stderr)
        sys.exit(1)

    # Tira as linhas do soup pra ler só os contadores do topo/rodapé da tela.
    for row in soup.select("div.group.grid.items-center.transition-colors"):
        (row.parent or row).decompose()
    ctx = extrair_contexto_pagina(soup)
    aviso = html_incompleto(ctx, len(linhas))

    grupos = agrupar_por_sku(linhas)
    wb = gerar_xlsx(grupos, args.conta, len(linhas), ctx, aviso)
    caminho_final = salvar_planilha(wb, caminho_saida)

    print(f"OK — {len(linhas)} MLBs em {len(grupos)} SKUs.")
    if aviso:
        print(f"\n*** {aviso} ***\n")
    if nao_reconhecidas:
        print(f"AVISO: {nao_reconhecidas} linha(s) do HTML não bateram com o padrão esperado e foram ignoradas — vale conferir manualmente.")
    com_outros = [a for a in linhas if a["outros"]]
    if com_outros:
        print(f"AVISO: {len(com_outros)} anúncio(s) têm dados que o script não conhecia (coluna 'Outros dados do HTML'). Exemplos:")
        vistos = []
        for a in com_outros:
            for parte in a["outros"].split(" | "):
                if parte not in vistos:
                    vistos.append(parte)
        for parte in vistos[:8]:
            print(f"   - {parte[:120]}")
    print(f"Salvo em: {caminho_final}")


if __name__ == "__main__":
    main()
