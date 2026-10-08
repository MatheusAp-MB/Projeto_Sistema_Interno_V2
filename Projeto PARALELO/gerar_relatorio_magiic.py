# gerar_relatorio_magiic.py
#
# Le o HTML salvo da tela "Envio para o FULL" da Magiic e gera um relatorio
# .xlsx numa ABA UNICA, agrupado por SKU (mesmo padrao do relatorio da WeStack),
# com TODOS os dados que o HTML traz de cada anuncio:
#   - Estoque ERP da empresa 1 e da empresa 4 (o numero da coluna "Estoque ERP"
#     mostra so um dos modos; o detalhamento com Emp 1 / Emp 4 / Total ja vem
#     escondido no HTML, dentro de cada numero);
#   - kits (composicao, Emp 1 / Emp 4 / Total calculados pelo componente que limita);
#   - vendas 30/28/21/14/7 dias, tendencia, vendas do catalogo, anuncios que
#     dividem o mesmo inventario;
#   - Full: transito, estoque alvo, estoque atual, dias de cobertura, minimo do
#     Mercado Livre, detalhamento (indisponivel, transferencia, perdido...);
#   - envio: "Enviar hoje", sugestao do Mercado Livre, alertas, previsao;
#   - qualidade do anuncio e experiencia de compra;
#   - e qualquer texto novo que a Magiic passe a mostrar (coluna "Outros dados
#     do HTML"). Nao usa API nenhuma, so le o que ja esta no HTML.
#
# IMPORTANTE: o HTML salvo so tem o que estava na tela na hora do Ctrl+S
# (com os filtros daquele momento). O script mostra no topo da planilha os
# filtros que estavam ativos e confere a soma do "Enviar hoje" com o
# "Qtd do Envio" da tela.
#
# Limite (nao existe no HTML salvo, so aparece ao clicar na tela):
#   - o detalhe do Transito ("Detalhes de Trânsito"): vem so o total.
#
# Uso:
#   python gerar_relatorio_magiic.py caminho\para\pagina.html
#   python gerar_relatorio_magiic.py caminho\para\pagina.html -o relatorio.xlsx
#   python gerar_relatorio_magiic.py caminho\para\pagina.html --conta MB
#
# Dependencias: pip install beautifulsoup4 openpyxl

from __future__ import annotations

import argparse
import json
import math
import re
import sys
import warnings
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from textwrap import wrap
from urllib.parse import unquote_plus

from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

try:  # o bs4 reclama quando um texto parece nome de arquivo/URL; aqui nao importa
    from bs4 import MarkupResemblesLocatorWarning
    warnings.filterwarnings("ignore", category=MarkupResemblesLocatorWarning)
except ImportError:  # versoes antigas do bs4
    pass

SEM_SKU = "(sem SKU)"

FONTE = "Arial"

COR_TITULO = "1F2937"
COR_BORDA = "D1D5DB"
COR_GRUPO_FUNDO = "BFDBFE"
COR_ZEBRA = "F3F4F6"
COR_DIVISORIA_FAIXA = "4B5563"
COR_GRUPO_TEXTO = "1E3A8A"

# Faixas de cor do cabeçalho: (cor escura da faixa, cor clara dos títulos das colunas).
FAIXAS = {
    "Identificação": ("374151", "E5E7EB"),
    "Estoque ERP": ("047857", "D1FAE5"),
    "Vendas": ("1E40AF", "DBEAFE"),
    "Full (Mercado Livre)": ("0F766E", "CCFBF1"),
    "Envio": ("5B21B6", "EDE9FE"),
    "Alertas": ("B91C1C", "FEE2E2"),
    "Qualidade": ("B45309", "FEF3C7"),
    "Textos completos": ("4B5563", "F3F4F6"),
}

# Cada coluna: (chave no dict do anúncio, título, largura mínima, largura máxima,
#               formato, alinhamento, faixa, quebra)
#   formato:     None | "int" | "dec" | "pct" | "link"
#   alinhamento: "esq" | "centro" | "dir"   (números à direita, textos à esquerda)
#   quebra:      True = o texto quebra em várias linhas e a altura da linha cresce pra caber
# A largura real de cada coluna é calculada pelo maior valor dela, entre o mínimo e o máximo.
COLUNAS = [
    ("mlb", "MLB", 22, 22, None, "esq", "Identificação", False),
    ("sku", "SKU", 20, 26, None, "esq", "Identificação", False),
    ("titulo", "Título do anúncio", 40, 60, None, "esq", "Identificação", True),
    ("status", "Status do anúncio", 11, 14, None, "centro", "Identificação", False),
    ("tipo", "Tipo (Simples / Kit / Sem ERP)", 12, 14, None, "centro", "Identificação", False),
    ("familia", "Família", 30, 46, None, "esq", "Identificação", True),
    ("familia_id", "ID da família", 18, 20, None, "centro", "Identificação", False),
    ("inventory_id", "Inventory ID (Full)", 14, 16, None, "centro", "Identificação", False),
    ("curva_anuncio", "Curva ABC do anúncio", 10, 12, None, "centro", "Identificação", False),
    ("curva_item", "Curva ABC do item", 10, 12, None, "centro", "Identificação", False),
    ("reposicao", "Reposição", 11, 12, None, "centro", "Identificação", False),
    ("reposicao_obs", "Reposição — quem e quando", 20, 30, None, "esq", "Identificação", True),
    ("produto_estrela", "Produto estrela (ML)", 10, 12, None, "centro", "Identificação", False),
    ("catalogo_mlb", "Anúncio de catálogo vinculado", 16, 18, None, "centro", "Identificação", False),
    ("catalogo_na_planilha", "Catálogo na planilha de envio", 12, 14, None, "centro", "Identificação", False),
    ("anotacoes", "Anotações no anúncio (qtd)", 11, 13, "int", "dir", "Identificação", False),
    ("n_anuncios_inv", "Anúncios no mesmo inventário", 11, 13, "int", "dir", "Identificação", False),
    ("anuncios_inv", "Anúncios do inventário (status, vendas 30d)", 36, 80, None, "esq", "Identificação", True),
    ("link", "Link do anúncio", 24, 40, "link", "esq", "Identificação", False),
    ("link_edicao", "Link de edição", 24, 40, "link", "esq", "Identificação", False),
    ("foto", "Foto (URL)", 24, 40, "link", "esq", "Identificação", False),

    ("erp_emp1", "Estoque Emp 1", 10, 12, "int", "dir", "Estoque ERP", False),
    ("erp_emp4", "Estoque Emp 4", 10, 12, "int", "dir", "Estoque ERP", False),
    ("erp_total", "Total (Emp 1 + Emp 4)", 11, 13, "int", "dir", "Estoque ERP", False),
    ("erp_exibido", "Número exibido na tela", 11, 13, "int", "dir", "Estoque ERP", False),
    ("kit_composicao", "Composição do kit", 40, 64, None, "esq", "Estoque ERP", True),
    ("erp_obs", "Estoque ERP — observação", 28, 46, None, "esq", "Estoque ERP", True),

    ("vendas_30", "Vendas 30 dias", 10, 12, "int", "dir", "Vendas", False),
    ("vendas_28", "Vendas 28 dias", 10, 12, "int", "dir", "Vendas", False),
    ("vendas_21", "Vendas 21 dias", 10, 12, "int", "dir", "Vendas", False),
    ("vendas_14", "Vendas 14 dias", 10, 12, "int", "dir", "Vendas", False),
    ("vendas_7", "Vendas 7 dias", 10, 12, "int", "dir", "Vendas", False),
    ("inclui_catalogo", "Inclui vendas do catálogo", 11, 13, None, "centro", "Vendas", False),
    ("vendas_catalogo", "Vendas do catálogo (dica da tela)", 22, 34, None, "esq", "Vendas", True),
    ("tendencia", "Tendência", 22, 34, None, "esq", "Vendas", True),
    ("tendencia_pct", "Variação %", 10, 12, "pct", "dir", "Vendas", False),

    ("transito", "Trânsito Full", 10, 12, "int", "dir", "Full (Mercado Livre)", False),
    ("alvo", "Estoque alvo Full", 10, 12, "int", "dir", "Full (Mercado Livre)", False),
    ("estoque_full", "Estoque atual Full", 10, 12, "int", "dir", "Full (Mercado Livre)", False),
    ("dias_cobertura", "Dias de cobertura", 10, 12, "int", "dir", "Full (Mercado Livre)", False),
    ("minimo_ml", "Mínimo do ML para distribuir", 11, 13, "int", "dir", "Full (Mercado Livre)", False),
    ("full_total_meli", "Estoque total (Meli)", 10, 12, "int", "dir", "Full (Mercado Livre)", False),
    ("full_indisponivel", "Total indisponível", 10, 12, "int", "dir", "Full (Mercado Livre)", False),
    ("full_disponivel", "Disponível (Meli)", 10, 12, "int", "dir", "Full (Mercado Livre)", False),
    ("full_status", "Indisponível por status (considerado?)", 28, 40, None, "esq", "Full (Mercado Livre)", True),
    ("full_considerado", "Estoque considerado (Magiic)", 11, 13, "int", "dir", "Full (Mercado Livre)", False),
    ("ml_dado_em", "Dado do ML coletado em", 16, 18, None, "centro", "Full (Mercado Livre)", False),

    ("enviar_hoje", "Enviar hoje", 10, 12, "int", "dir", "Envio", False),
    ("envio_original", "Envio original (antes do limite)", 11, 13, "int", "dir", "Envio", False),
    ("envio_bloqueado", "Campo Enviar bloqueado", 10, 12, None, "centro", "Envio", False),
    ("base", "Base do envio (data-base)", 10, 12, "int", "dir", "Envio", False),
    ("teto", "Teto do envio (data-teto)", 10, 12, "int", "dir", "Envio", False),
    ("venda_diaria", "Venda diária", 10, 12, "dec", "dir", "Envio", False),
    ("qtd_caixa", "Qtd por caixa", 10, 12, "int", "dir", "Envio", False),
    ("qtd_pallet", "Qtd por pallet", 10, 12, "int", "dir", "Envio", False),
    ("estoque_limitado", "Envio limitado pelo estoque (campo)", 12, 14, None, "centro", "Envio", False),
    ("sugestao_ml", "Sugestão do ML", 10, 12, "int", "dir", "Envio", False),
    ("sugestao_prazo", "Sugestão — prazo", 12, 20, None, "centro", "Envio", True),
    ("alerta_ml", "Alerta do ML", 18, 28, None, "centro", "Envio", True),
    ("alerta_ml_detalhe", "Alerta do ML — detalhe", 36, 56, None, "esq", "Envio", True),
    ("previsao", "Previsão após o envio", 10, 12, "int", "dir", "Envio", False),
    ("previsao_obs", "Previsão — observação", 28, 44, None, "esq", "Envio", True),

    ("alerta_ruptura", "Ruptura / ruptura prevista", 11, 13, None, "centro", "Alertas", False),
    ("alerta_ponto_reposicao", "Ponto de reposição", 11, 13, None, "centro", "Alertas", False),
    ("alerta_envio_limitado", "Envio limitado pelo estoque local", 12, 14, None, "centro", "Alertas", False),
    ("alerta_sem_estoque", "Anúncio/variação sem estoque", 12, 14, None, "centro", "Alertas", False),
    ("nao_repor", "Não repor no Full", 10, 12, None, "centro", "Alertas", False),
    ("nao_repor_obs", "Não repor — quem e quando", 20, 30, None, "esq", "Alertas", True),

    ("qualidade_pontos", "Qualidade do anúncio (pontos)", 11, 13, "int", "dir", "Qualidade", False),
    ("qualidade_situacao", "Qualidade — situação", 20, 28, None, "esq", "Qualidade", True),
    ("qualidade_pendencias", "Qualidade — pendências", 40, 90, None, "esq", "Qualidade", True),
    ("experiencia_pontos", "Experiência de compra (pontos)", 11, 13, "int", "dir", "Qualidade", False),
    ("experiencia_texto", "Experiência de compra — texto", 40, 90, None, "esq", "Qualidade", True),

    ("textos_completos", "Textos completos das dicas da tela (clique na célula pra ler tudo)", 40, 60, None, "esq", "Textos completos", False),
    ("outros", "Outros dados do HTML", 30, 60, None, "esq", "Textos completos", False),
]

# Cores de células pela tela (fundo, texto).
VERDE = ("C6EFCE", "006100")
AMARELO = ("FFEB9C", "7F6000")
VERMELHO = ("FFC7CE", "9C0006")
ROXO = ("E9D5FF", "6B21A8")
CINZA = ("E5E7EB", "374151")
AZUL = ("DBEAFE", "1E3A8A")
LARANJA = ("FFE4C4", "7C2D12")
CORES_VALOR = {
    "status": {"Ativo": VERDE, "Pausado": AMARELO},
    "curva_anuncio": {"A": VERDE, "B": AMARELO, "C": VERMELHO, "0": ROXO},
    "curva_item": {"A": VERDE, "B": AMARELO, "C": VERMELHO, "0": ROXO},
    "tipo": {"Kit": AZUL, "Sem ERP": VERMELHO},
    "reposicao": {"Inativa": CINZA},
    "produto_estrela": {"Sim": AMARELO},
    "alerta_ruptura": {"Sim": VERMELHO},
    "alerta_ponto_reposicao": {"Sim": AMARELO},
    "alerta_envio_limitado": {"Sim": LARANJA},
    "alerta_sem_estoque": {"Sim": VERMELHO},
    "nao_repor": {"Sim": CINZA},
}

FORMATOS = {"int": "0", "dec": "0.00", "pct": "0.0%"}

STATUS_ML = {"active": "Ativo", "closed": "Encerrado", "paused": "Pausado", "removed": "Removido",
             "under_review": "Em revisão", "inactive": "Inativo"}
TIPO_ML = {"fulfillment": "Full"}
MODOS_ESTOQUE = {"1": "Total (Emp 1 + Emp 4)", "2": "Apenas Estoque Emp 1", "3": "Apenas Estoque Emp 4"}

# Textos fixos da interface (botões, rótulos, dicas que só explicam o controle).
# Não são dados do anúncio, então não entram em "Outros dados do HTML".
RUIDO_EXATO = {
    "", "—", "-", "%", "R$", "!", "A", "B", "C", "0", "ADS",
    "Filtrar apenas este anúncio", "Dados do Anúncio:", "Dados do Anúncio", "Checar promoções",
    "Ver mais detalhes", "Salvo com sucesso!", "Carregando dados...", "Detalhes de Trânsito",
    "Consulta de Operações de Estoque", "Configurar Arredondamento", "Arredondar para caixa/pallet",
    "Qualidade do Anúncio", "Experiência de Compra",
    "selecione aqui o nome do produto", "selecione aqui o código ML do produto",
    "clique aqui para optar ou não por mostrar o anúncio catálogo na planilha de envio para o Meli",
    "Planejamento do Full: recomendação de reposição gravada pela coleta e estoque por localização "
    "consultado agora no Mercado Livre",
}
RUIDO_REGEX = [re.compile(p) for p in (
    r"^Visualizar Elasticidade do Anúncio MLB\d+$",
    r"^Exibir anotações do anúncio - MLB\d+$",
    r"^(Qualidade do Anúncio|Experiência de Compra) - Clique para visualizar mais informações$",
    r"^.* - ir para a página de edição do Anúncio no MELI$",
    r"^.* - link para checar Promoções no MELI$",
    r"^SKU secundário: insira aqui o SKU cadastrado em seu ERP, para o produto: MLB\d+$",
    r"^Caso seu produto MLB\d+ seja um KIT, ajuste aqui a formação$",
)]

# Atributos data-* que o script já conhece (ou que só repetem outro dado). Qualquer
# data-* diferente destes aparece em "Outros dados do HTML".
DATA_CONHECIDOS = {
    "data-toggle", "data-html", "data-placement", "data-trigger", "data-boundary", "data-container",
    "data-variant", "data-template", "data-original-title", "data-content", "data-title", "data-id",
    "data-id-prod", "data-cid", "data-item-id", "data-anuncio", "data-count", "data-type", "data-up-id",
    "data-products", "data-total-sales", "data-sku", "data-inventory-id", "data-product-name",
    "data-family-id", "data-src", "data-base", "data-teto", "data-venda-diaria", "data-qtd-caixa",
    "data-qtd-pallet", "data-estoque-limitado", "data-estoque", "data-minimo", "data-transito",
    "data-sugestao-ml", "data-alerta-ml", "data-label-id", "data-catalog-product-id", "data-product-id",
    "data-product", "data-thumb", "data-url",       # botão de anotações: repete MLB, foto e link do anúncio
}

PADRAO_DADO_ML = re.compile(r"Dado do Mercado Livre de\s*(\d{2}/\d{2}/\d{4}\s+\d{2}:\d{2})")
PADRAO_PERCENTUAL = re.compile(r"(-?\d+(?:[.,]\d+)?)\s*%")
PADRAO_MLB = re.compile(r"MLB\d+")


# ---------------------------------------------------------------- utilitários

def _norm(texto) -> str:
    """Tira espaços repetidos, espaço 'não separável' e caracteres de controle que o Excel não aceita."""
    texto = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", texto or "")
    return re.sub(r"\s+", " ", texto.replace("\xa0", " ")).strip()


def _plano(valor) -> str:
    """Texto puro de um atributo que traz HTML (dica / popover)."""
    if not valor:
        return ""
    if "<" in valor:
        valor = BeautifulSoup(valor, "html.parser").get_text(" ")
    return _norm(valor)


def _dica(el) -> str:
    """Texto da dica (tooltip) de um elemento."""
    if el is None:
        return ""
    return _plano(el.get("data-original-title") or el.get("title") or "")


def _inteiro(texto):
    m = re.search(r"-?\d[\d.]*", _norm(texto))
    if not m:
        return None
    digitos = m.group().replace(".", "")
    try:
        return int(digitos)
    except ValueError:
        return None


def _numero(valor):
    """'50.00' -> 50 ; '6.466' -> 6.47 ; vazio -> None."""
    if valor is None or _norm(str(valor)) == "":
        return None
    try:
        n = float(str(valor).replace(",", "."))
    except ValueError:
        return None
    return int(n) if n.is_integer() else round(n, 2)


def _sim(valor):
    return "Sim" if valor else None


def _fracao(texto):
    m = PADRAO_PERCENTUAL.search(_norm(texto))
    if not m:
        return None
    try:
        return round(float(m.group(1).replace(",", ".")) / 100, 6)
    except ValueError:
        return None


class Uso:
    """Guarda os textos que o script já leu de cada anúncio; o que sobrar vai pra 'Outros dados do HTML'."""

    def __init__(self):
        self.textos = set()

    def add(self, *textos):
        for t in textos:
            if hasattr(t, "stripped_strings"):          # elemento: usa o texto todo e cada pedaço dele
                self.add(t.get_text(), *t.stripped_strings)
                continue
            t = _norm(t) if t else ""
            if t:
                self.textos.add(t)

    def foi(self, texto) -> bool:
        return texto in self.textos


def _e_ruido(texto: str) -> bool:
    return texto in RUIDO_EXATO or any(p.match(texto) for p in RUIDO_REGEX)


def _sobras(raizes, uso: Uso) -> list[str]:
    """Textos, dicas e atributos data-* dos elementos que nenhuma regra do script usou."""
    achados = []

    def candidato(t):
        t = _norm(t)
        if t and not uso.foi(t) and not _e_ruido(t) and t not in achados:
            achados.append(t)

    for raiz in raizes:
        for t in raiz.stripped_strings:
            candidato(t)
        for el in [raiz] + raiz.find_all(True):
            for atributo in ("data-original-title", "data-content", "title"):
                valor = el.get(atributo)
                if valor:
                    candidato(_plano(valor))
            for atributo, valor in el.attrs.items():
                if atributo.startswith("data-") and atributo not in DATA_CONHECIDOS:
                    candidato(f"{atributo}={valor if not isinstance(valor, list) else ' '.join(valor)}")
    return achados


# ----------------------------------------------------------- leitura por célula

def _ler_cinza(cinza, uso: Uso) -> dict:
    """Linha cinza do anúncio: status, MLB, título, qualidade, experiência, curva ABC."""
    tds = cinza.find_all("td", recursive=False)
    d = {}
    el = cinza.find(class_="product-status-excel")
    d["status"] = _norm(el.get_text()) if el else None
    uso.add(d["status"])

    d["mlb_cinza"] = None
    if len(tds) > 1:
        m = PADRAO_MLB.search(tds[1].get_text())
        d["mlb_cinza"] = m.group() if m else None
        uso.add(d["mlb_cinza"])
    d["titulo_cinza"] = None
    if len(tds) > 2:
        span = tds[2].find("span")
        d["titulo_cinza"] = _norm(span.get_text()) if span else None
        uso.add(d["titulo_cinza"])
        for link in tds[2].find_all("a"):               # dica do "Checar promoções" = título + frase fixa
            uso.add(_dica(link))

    # selo vermelho do botão de anotações = quantidade de anotações do anúncio
    d["anotacoes"] = None
    selo = cinza.select_one("a.btn-anotacoes span.badge")
    if selo is not None:
        d["anotacoes"] = _inteiro(selo.get_text())
        uso.add(selo)

    d["curva_anuncio"] = _norm(tds[5].get_text()) if len(tds) > 5 else None
    uso.add(d["curva_anuncio"])

    # Qualidade do anúncio
    q = cinza.find(class_="qualidade-span")
    d.update(qualidade_pontos=None, qualidade_situacao=None, qualidade_pendencias=None)
    if q is not None:
        texto_visivel = _norm(q.get_text())
        uso.add(texto_visivel)
        d["qualidade_pontos"] = _inteiro(texto_visivel) if re.search(r"\d", texto_visivel) else None
        conteudo = BeautifulSoup(q.get("data-content") or "", "html.parser")
        itens = [_norm(li.get_text()) for li in conteudo.find_all("li")]
        forte = conteudo.find("strong")
        completo = _plano(q.get("data-content"))
        plano = re.sub(r"\s*Ver mais detalhes\s*$", "", completo)
        uso.add(completo, plano, *itens)
        if forte is not None:
            d["qualidade_situacao"] = _norm(forte.get_text()).rstrip(":")
        elif plano:
            d["qualidade_situacao"] = plano
        d["qualidade_pendencias"] = "\n".join(f"• {i}" for i in itens) or None
        for tip in q.find_all(attrs={"data-original-title": True}):
            t = _dica(tip)
            if t.startswith("Qualidade:") and not d["qualidade_situacao"]:
                d["qualidade_situacao"] = t
            uso.add(t)
        if not d["qualidade_situacao"] and texto_visivel and d["qualidade_pontos"] is None:
            d["qualidade_situacao"] = texto_visivel

    # Experiência de compra
    r = cinza.find(class_="reputacao-span")
    d.update(experiencia_pontos=None, experiencia_texto=None)
    if r is not None:
        texto_visivel = _norm(r.get_text())
        uso.add(texto_visivel)
        d["experiencia_pontos"] = _inteiro(texto_visivel) if re.search(r"\d", texto_visivel) else None
        completo = _plano(r.get("data-content"))
        plano = re.sub(r"\s*Ver mais detalhes\s*$", "", completo)
        uso.add(completo, plano)
        d["experiencia_texto"] = plano or None
        for tip in r.find_all(attrs={"data-original-title": True}):
            t = _dica(tip)
            if t.startswith("Experiência:") and not d["experiencia_texto"]:
                d["experiencia_texto"] = t
            uso.add(t)
        if not d["experiencia_texto"] and texto_visivel and d["experiencia_pontos"] is None:
            d["experiencia_texto"] = texto_visivel
    return d


def _ler_inventario(td2, mlb: str, uso: Uso) -> dict:
    """'(N itens)' e a lista de anúncios que dividem o mesmo inventário (JSON dentro do HTML)."""
    d = dict(n_anuncios_inv=None, anuncios_inv=None)
    up = td2.find(class_="user-product-merged")
    if up is None:
        return d
    uso.add(up.get_text())
    d["n_anuncios_inv"] = _inteiro(up.get("data-count") or "")
    try:
        produtos = json.loads(up.get("data-products") or "[]")
    except ValueError:
        produtos = []
    linhas = []
    for p in produtos:
        partes = [str(p.get("id") or "")]
        status = STATUS_ML.get(p.get("status"), p.get("status") or "")
        if status:
            partes.append(status)
        tipo = TIPO_ML.get(p.get("tipo"), p.get("tipo") or "")
        if tipo:
            partes.append(tipo)
        if p.get("is_catalog"):
            partes.append("catálogo")
        if p.get("has_catalog"):
            partes.append(f"catálogo vinculado {p['has_catalog']}")
        if p.get("ignored"):
            partes.append("ignorado")
        partes.append(f"vendas 30d: {p.get('sales_30', 0)}")
        linha = " · ".join(partes)
        if p.get("id") == mlb:
            linha += "  ← este"
        linhas.append(linha)
    d["anuncios_inv"] = "\n".join(linhas) or None
    return d


def _ler_td2(td2, uso: Uso, textos: list) -> dict:
    """Identificador do MELI: MLB, título, inventário, estrela, reposição, catálogo vinculado."""
    d = {}
    a = td2.find("a", class_="id-meli-excel-1")
    d["mlb"] = _norm(a.get_text()) if a else None
    d["link_edicao"] = a.get("href") if a else None
    uso.add(d["mlb"], _dica(a))
    nome = td2.find(class_="selecionavel text-white") or td2.find(class_="text-white")
    d["titulo"] = _norm(nome.get_text()) if nome else None
    uso.add(d["titulo"])
    inv = td2.find(class_="id-meli-excel-2")
    d["inventory_id"] = _norm(inv.get_text()) if inv else None
    uso.add(d["inventory_id"])

    caixa = td2.find(class_="caixa-master-icon")
    d["sku_icone"] = caixa.get("data-sku") if caixa else None

    # Produto estrela do Mercado Livre
    estrela = td2.find(class_="repl-ml-estrela")
    d["produto_estrela"] = _sim(estrela is not None)
    d["ml_dado_em_estrela"] = None
    if estrela is not None:
        dono = estrela.find_parent("a")
        tip = _dica(dono)
        uso.add(tip)
        textos.append(tip)
        m = PADRAO_DADO_ML.search(tip)
        d["ml_dado_em_estrela"] = m.group(1) if m else None

    # Reposição ativa / inativa (e quem mexeu)
    tog = td2.find(class_="toggle-reposicao")
    d["reposicao"] = d["reposicao_obs"] = None
    if tog is not None:
        classes = tog.get("class") or []
        d["reposicao"] = "Ativa" if "ativa" in classes else ("Inativa" if "inativa" in classes else None)
        tip = _dica(tog)
        uso.add(tip)
        textos.append(tip)
        m = re.search(r"(Ativado|Desativado) por (.+?) Data (\d{2}/\d{2}/\d{4}) às (\d{2}:\d{2})", tip)
        if m:
            d["reposicao_obs"] = f"{m.group(1)} por {m.group(2)} em {m.group(3)} às {m.group(4)}"

    # Anúncio de catálogo vinculado (o "C MLB..." com caixinha)
    d["catalogo_mlb"] = d["catalogo_na_planilha"] = None
    cx = td2.find("input", attrs={"type": "checkbox"})
    if cx is not None:
        d["catalogo_mlb"] = cx.get("data-catalog-product-id")
        d["catalogo_na_planilha"] = "Sim" if cx.has_attr("checked") else "Não"
        uso.add(d["catalogo_mlb"], _dica(cx))
    d.update(_ler_inventario(td2, d["mlb"] or "", uso))
    return d


def _ler_erp(td4, td3, uso: Uso, textos: list) -> dict:
    """Estoque ERP: popover simples (Emp 1 / Emp 4 / Total) ou de kit (componentes)."""
    d = dict(tipo=None, erp_emp1=None, erp_emp4=None, erp_total=None, erp_exibido=None,
             kit_composicao=None, modo=None, erp_obs=[], extras=[])
    # SKU não localizado no ERP (aviso na célula do ERP e na do SKU)
    for celula in (td4, td3):
        for el in celula.find_all(attrs={"data-original-title": True}):
            tip = _dica(el)
            if "SKU não localizado no ERP" in tip:
                d["tipo"] = "Sem ERP"
                uso.add(tip)
    span = td4.find("span", class_="estoque-com-detalhes")
    if span is None:
        if d["tipo"] == "Sem ERP":
            d["erp_obs"].append("SKU não localizado no ERP (a Magiic não trouxe estoque)")
        return d

    classes = span.get("class") or []
    conteudo = BeautifulSoup(span.get("data-content") or "", "html.parser")
    uso.add(_plano(span.get("data-content")), span.get_text())
    uso.add(span.get("data-title") or "")
    d["erp_exibido"] = _inteiro(span.get_text())
    opcao = conteudo.select_one("option[selected]")
    d["modo"] = opcao.get("value") if opcao else None

    if "estoque-kit-popover" in classes:
        d["tipo"] = "Kit"
        componentes = []
        for tr in conteudo.select("tbody tr"):
            c = [_norm(td.get_text()) for td in tr.find_all("td")]
            if len(c) < 5:
                continue
            qtd = _numero(c[1]) or 0
            componentes.append(dict(sku=c[0], qtd=qtd, emp1=_inteiro(c[2]) or 0, emp4=_inteiro(c[3]) or 0,
                                    disp=_inteiro(c[4]) or 0, limitante="limitante" in (tr.get("class") or [])))
        validos = [c for c in componentes if c["qtd"] > 0]
        if validos:
            d["erp_emp1"] = min(math.floor(c["emp1"] / c["qtd"]) for c in validos)
            d["erp_emp4"] = min(math.floor(c["emp4"] / c["qtd"]) for c in validos)
            d["erp_total"] = min(math.floor((c["emp1"] + c["emp4"]) / c["qtd"]) for c in validos)
        linhas = []
        for c in componentes:
            qtd = int(c["qtd"]) if float(c["qtd"]).is_integer() else c["qtd"]
            linha = f"{c['sku']} ×{qtd} · Emp 1: {c['emp1']} · Emp 4: {c['emp4']} · disponível p/ o kit: {c['disp']}"
            if c["limitante"]:
                linha += "  ← limita o kit"
            linhas.append(linha)
        d["kit_composicao"] = "\n".join(linhas) or None
        caixa = conteudo.find(class_="kit-total-box")
        d["erp_obs"].append("Kit: Emp 1, Emp 4 e Total calculados pelo componente que limita (quantidade do kit ÷ qtd de cada componente)")
        esperado = {"1": d["erp_total"], "2": d["erp_emp1"], "3": d["erp_emp4"]}.get(d["modo"])
        if caixa is not None:
            informado = _inteiro(caixa.get_text().split("unidades")[0])
            if informado is not None and esperado is not None and informado != esperado:
                d["erp_obs"].append(f"ATENÇÃO: a Magiic informa {informado} para o kit; o cálculo deu {esperado}")
    else:
        d["tipo"] = "Simples"
        for linha in conteudo.select(".estoque-info-row"):
            rotulo = linha.select_one(".estoque-info-label")
            valor = linha.select_one(".estoque-info-value")
            if rotulo is None or valor is None:
                continue
            rot = _norm(rotulo.get_text())
            n = _inteiro(valor.get_text())
            m = re.search(r"Emp\s*(\d+)", rot)
            if m and m.group(1) == "1":
                d["erp_emp1"] = n
            elif m and m.group(1) == "4":
                d["erp_emp4"] = n
            elif m:
                d["extras"].append(f"Estoque Emp {m.group(1)}: {n}")
            elif "Total" in rot:
                d["erp_total"] = n
            else:
                d["extras"].append(f"{rot} {n}")
        if d["erp_total"] is None and d["erp_emp1"] is not None and d["erp_emp4"] is not None:
            d["erp_total"] = d["erp_emp1"] + d["erp_emp4"]
        esperado = {"1": d["erp_total"], "2": d["erp_emp1"], "3": d["erp_emp4"]}.get(d["modo"])
        if d["erp_exibido"] is not None and esperado is not None and d["erp_exibido"] != esperado:
            d["erp_obs"].append(f"A tela mostra {d['erp_exibido']}, mas o estoque real é {esperado}")

    # ícone ao lado do número (ex.: "Estoque local subtraído no cálculo do envio atual")
    for el in td4.find_all(attrs={"data-original-title": True}):
        tip = _dica(el)
        if tip and not tip.startswith("SKU não localizado"):
            uso.add(tip)
            textos.append(tip)
            d["erp_obs"].append(tip)
    return d


def _ler_vendas(td5, uso: Uso, textos: list) -> dict:
    d = dict(vendas_30=None, vendas_28=None, vendas_21=None, vendas_14=None, vendas_7=None,
             inclui_catalogo=None, vendas_catalogo=None, tendencia=None, tendencia_pct=None)
    col = td5.select_one("div.col.text-end") or td5
    linhas = [_norm(x) for x in col.get_text("\n").split("\n") if _norm(x)]
    if linhas:
        d["vendas_30"] = _inteiro(linhas[0])
        uso.add(linhas[0])
    if len(linhas) > 1:
        m = re.fullmatch(r"(\d+)/(\d+)/(\d+)/(\d+)", linhas[1])
        if m:
            d["vendas_28"], d["vendas_21"], d["vendas_14"], d["vendas_7"] = (int(g) for g in m.groups())
            uso.add(linhas[1])
    for el in td5.find_all(attrs={"data-original-title": True}):
        tip = _dica(el)
        if not tip:
            continue
        if "período comparado" in tip:
            m = re.match(r"^(.*?)\s*no período comparado\s*\d+\s*/\s*\d+$", tip, flags=re.I)
            d["tendencia"] = m.group(1) if m else tip
            d["tendencia_pct"] = _fracao(d["tendencia"])
        elif tip.startswith("Incluindo vendas"):
            d["inclui_catalogo"] = "Sim"
            d["vendas_catalogo"] = re.sub(r"^Incluindo vendas de Anúncio Catálogo:\s*", "", tip)
        uso.add(tip)
        textos.append(tip)
    for el in td5.find_all(["span", "div"]):
        if _norm(el.get_text()) == "C" and not el.find(True):
            d["inclui_catalogo"] = "Sim"
    return d


def _ler_full(td6, td7, td8, uso: Uso, textos: list) -> dict:
    d = dict(transito=None, alvo=None, estoque_full=None, dias_cobertura=None, minimo_ml=None,
             full_total_meli=None, full_indisponivel=None, full_disponivel=None, full_status=None,
             full_considerado=None, alerta_sem_estoque=None, ml_dado_em_min=None)
    el = td6.find(class_="transit-value")
    d["transito"] = _inteiro(el.get_text()) if el else _inteiro(td6.get_text())
    uso.add(el.get_text() if el else td6.get_text())
    d["alvo"] = _inteiro(td7.get_text())
    uso.add(td7.get_text())

    atual = td8.select_one(".estoque-atual-excel")
    d["estoque_full"] = _inteiro(atual.get_text()) if atual else None
    uso.add(atual.get_text() if atual else "")
    for bloco in td8.select("div.text-end.w-100"):
        classes = bloco.get("class") or []
        if "estoque-atual-excel" in classes or "repl-ml-minimo-slot" in classes:
            continue
        txt = _norm(bloco.get_text())
        if not txt:
            continue
        m = re.fullmatch(r"(\d+)\s*dias?", txt)
        d["dias_cobertura"] = int(m.group(1)) if m else txt      # "N/D" quando não há venda
        uso.add(txt)
    minimo = td8.select_one(".repl-ml-minimo-slot a")
    if minimo is not None:
        d["minimo_ml"] = _inteiro(minimo.get_text())
        uso.add(minimo.get_text())
        tip = _dica(minimo)
        uso.add(tip)
        textos.append(tip)
        m = PADRAO_DADO_ML.search(tip)
        d["ml_dado_em_min"] = m.group(1) if m else None

    for el in td8.find_all(attrs={"data-original-title": True}):
        tip = _dica(el)
        if not tip:
            continue
        if tip.startswith("Detalhamento de Estoque Full"):
            d["full_total_meli"] = _inteiro((re.search(r"Estoque Total \(Meli\)\s*(\d+)", tip) or [None, ""])[1])
            d["full_indisponivel"] = _inteiro((re.search(r"Total Indisponível\s*(\d+)", tip) or [None, ""])[1])
            d["full_disponivel"] = _inteiro((re.search(r"Estoque Disponível \(Meli\)\s*(\d+)", tip) or [None, ""])[1])
            trecho = re.search(r"Considerado\?\s*(.*?)\s*Estoque Considerado", tip)
            if trecho:
                linhas = []
                for nome, qtd, ok in re.findall(r"(.+?)\s+(\d+)\s+[✔✘]\s*(Sim|Não)\s*", trecho.group(1) + " "):
                    linhas.append(f"{_norm(nome)}: {qtd} ({'considerado' if ok == 'Sim' else 'não considerado'})")
                d["full_status"] = "\n".join(linhas) or None
            m = re.search(r"Estoque Considerado \(Magiic\).*?\)\s*(\d+)\s*unidade", tip)
            d["full_considerado"] = int(m.group(1)) if m else None
            textos.append(tip)
        elif tip.startswith("Anúncios ou Variações sem Estoque"):
            d["alerta_sem_estoque"] = "Sim"
            textos.append(tip)
        uso.add(tip)
    return d


def _ler_envio(td9, uso: Uso, textos: list) -> dict:
    d = dict(enviar_hoje=None, envio_bloqueado=None, base=None, teto=None, venda_diaria=None,
             qtd_caixa=None, qtd_pallet=None, estoque_limitado=None, sugestao_ml=None, sugestao_prazo=None,
             alerta_ml=None, alerta_ml_detalhe=None, nao_repor=None, nao_repor_obs=None, ml_dado_em_envio=None)
    inp = td9.find("input")
    if inp is not None:
        d["enviar_hoje"] = _inteiro(inp.get("value") or "")
        d["envio_bloqueado"] = _sim(inp.has_attr("disabled"))
        d["base"] = _numero(inp.get("data-base"))
        d["teto"] = _numero(inp.get("data-teto"))
        d["venda_diaria"] = _numero(inp.get("data-venda-diaria"))
        d["qtd_caixa"] = _numero(inp.get("data-qtd-caixa"))
        d["qtd_pallet"] = _numero(inp.get("data-qtd-pallet"))
        limitado = inp.get("data-estoque-limitado")
        d["estoque_limitado"] = {"true": "Sim", "false": "Não"}.get(limitado)
    sug = td9.find(attrs={"data-sugestao-ml": True})
    if sug is not None:
        d["sugestao_ml"] = _inteiro(sug.get("data-sugestao-ml"))
        visivel = _norm(sug.get_text())
        uso.add(sug)
        if "·" in visivel:
            d["sugestao_prazo"] = _norm(visivel.split("·", 1)[1])
        tip = _dica(sug)
        uso.add(tip)
        textos.append(tip)
        m = PADRAO_DADO_ML.search(tip)
        d["ml_dado_em_envio"] = d["ml_dado_em_envio"] or (m.group(1) if m else None)
    alerta = td9.find(attrs={"data-alerta-ml": True})
    if alerta is not None:
        d["alerta_ml"] = _norm(alerta.get_text())
        uso.add(alerta)
        tip = _dica(alerta)
        d["alerta_ml_detalhe"] = tip or None
        uso.add(tip)
        textos.append(tip)
        m = PADRAO_DADO_ML.search(tip)
        d["ml_dado_em_envio"] = d["ml_dado_em_envio"] or (m.group(1) if m else None)
    nr = td9.find("i", class_="tooltip-nao-repor")
    if nr is not None:
        d["nao_repor"] = "Sim"
        tip = _dica(nr)
        uso.add(tip)
        textos.append(tip)
        m = re.search(r"Desativado por (.+?) Data (\d{2}/\d{2}/\d{4}) às (\d{2}:\d{2})", tip)
        d["nao_repor_obs"] = f"Desativado por {m.group(1)} em {m.group(2)} às {m.group(3)}" if m else None
    return d


def _ler_alertas(td10, uso: Uso, textos: list) -> dict:
    d = dict(previsao=None, previsao_obs=None, alerta_ruptura=None, alerta_ponto_reposicao=None,
             alerta_envio_limitado=None, envio_original=None)
    for el in td10.find_all(attrs={"data-original-title": True}):
        tip = _dica(el)
        if not tip:
            continue
        if tip.startswith("Previsão após o envio"):
            m = re.match(r"Previsão após o envio:\s*(\d+)\s*unidades?\s*\d+ no Full \+ \d+ em trânsito \+ \d+ a enviar\.?\s*(.*)$", tip)
            if m:
                d["previsao"] = int(m.group(1))
                d["previsao_obs"] = m.group(2) or None
            else:
                d["previsao_obs"] = tip
        elif tip.startswith("Item em ruptura"):
            d["alerta_ruptura"] = "Sim"
        elif tip.startswith("Item no Ponto de Reposição"):
            d["alerta_ponto_reposicao"] = "Sim"
        elif tip.startswith("Envio limitado pelo estoque local"):
            d["alerta_envio_limitado"] = "Sim"
            m = re.search(r"Envio original\s*=\s*(\d+)", tip)
            d["envio_original"] = int(m.group(1)) if m else None
        uso.add(tip)
        textos.append(tip)
    return d


# ------------------------------------------------------------------ anúncios

def extrair_anuncio(cinza, det, familias: dict) -> dict | None:
    """Junta a linha cinza (anúncio) e a linha de detalhe (SKU, estoque, vendas, envio) num único dict."""
    tds = det.find_all("td", recursive=False)
    if len(tds) < 10:
        return None
    uso = Uso()
    textos: list[str] = []

    d = _ler_cinza(cinza, uso)
    d["curva_item"] = _norm(tds[0].get_text())
    uso.add(d["curva_item"])

    # foto e link do anúncio
    d["link"] = d["foto"] = None
    link = tds[1].find("a", href=True)
    d["link"] = link["href"] if link else None
    img = tds[1].find("img")
    if img is not None:
        d["foto"] = img.get("data-src") or img.get("src")
        uso.add(img.get("alt"), _dica(img))

    d.update(_ler_td2(tds[2], uso, textos))
    d["mlb"] = d["mlb"] or d.pop("mlb_cinza")
    d.pop("mlb_cinza", None)
    d["titulo"] = d["titulo"] or d.pop("titulo_cinza", None)
    d.pop("titulo_cinza", None)
    if not d["mlb"]:
        return None

    sku_el = tds[3].find(class_="sku-excel")
    d["sku"] = _norm(sku_el.get_text()) if sku_el else ""
    uso.add(d["sku"])

    erp = _ler_erp(tds[4], tds[3], uso, textos)
    d.update({k: v for k, v in erp.items() if k not in ("modo", "erp_obs", "extras")})
    d["modo"] = erp["modo"]
    d.update(_ler_vendas(tds[5], uso, textos))
    d.update(_ler_full(tds[6], tds[7], tds[8], uso, textos))
    d.update(_ler_envio(tds[9], uso, textos))
    d.update(_ler_alertas(tds[10], uso, textos) if len(tds) > 10 else
             dict(previsao=None, previsao_obs=None, alerta_ruptura=None, alerta_ponto_reposicao=None,
                  alerta_envio_limitado=None, envio_original=None))

    # a dica "Planejamento do Full" e os de ML usam a mesma data de coleta; guarda a primeira que achar
    d["ml_dado_em"] = d.pop("ml_dado_em_min") or d.pop("ml_dado_em_envio") or d.pop("ml_dado_em_estrela")
    d.pop("ml_dado_em_envio", None)
    d.pop("ml_dado_em_estrela", None)

    fam = familias.get(cinza.get("data-family-id"))
    d["familia"], d["familia_id"] = (fam if fam else (None, cinza.get("data-family-id")))

    obs = list(erp["erp_obs"])
    if d.get("sku_icone") and d["sku"] and d["sku_icone"] != d["sku"]:
        obs.append(f"SKU do ícone da caixa master: {d['sku_icone']}")
    d["erp_obs"] = obs                      # lista; a planilha junta depois (precisa saber do SKU compartilhado)

    outros = list(erp["extras"]) + _sobras([cinza, det], uso)
    d["outros"] = " | ".join(outros) or None
    d["textos_completos"] = " || ".join(dict.fromkeys(t for t in textos if t)) or None
    return d


def localizar_pares(soup: BeautifulSoup):
    """Acha os pares (linha cinza do anúncio, linha de detalhe) e as famílias. Conta o que não reconheceu."""
    tabela = soup.find("table", id="main_table")
    if tabela is None:
        return [], {}, 0, None
    cabecalho = [_norm(th.get_text(" ")) for th in tabela.find_all("th")]
    esperado = ["Identificador do MELI", "SKU", "Estoque ERP", "Vendas 30 dias", "Trânsito FULL",
                "Estoque alvo FULL", "Estoque atual FULL", "Enviar hoje"]
    aviso_cab = None
    if not all(any(e in c for c in cabecalho) for e in esperado):
        aviso_cab = ("O cabeçalho da tabela mudou (colunas diferentes do esperado). "
                     "Confira o relatório com cuidado: o script lê as colunas pela posição.")
    familias, pares, nao_reconhecidas = {}, [], 0
    usadas = set()
    for tr in tabela.find_all("tr"):
        classes = tr.get("class") or []
        if "family-header-row" in classes:
            tds = tr.find_all("td", recursive=False)
            if len(tds) > 1:
                nome = tds[1].find("span")
                ident = tds[1].find("a", class_="text-muted")
                familias[tr.get("data-family-id")] = (_norm(nome.get_text()) if nome else None,
                                                     _norm(ident.get_text()) if ident else None)
        elif "family-anuncio-row" in classes:
            det = tr.find_next_sibling("tr")
            if det is not None and not (det.get("class") or []) and len(det.find_all("td", recursive=False)) >= 10:
                pares.append((tr, det))
                usadas.add(id(det))
            else:
                nao_reconhecidas += 1
    for tr in tabela.find_all("tr"):
        if not (tr.get("class") or []) and len(tr.find_all("td", recursive=False)) >= 10 and id(tr) not in usadas:
            nao_reconhecidas += 1
    return pares, familias, nao_reconhecidas, aviso_cab


def extrair_linhas(soup: BeautifulSoup):
    pares, familias, nao_reconhecidas, aviso_cab = localizar_pares(soup)
    anuncios, falhas = [], []
    for cinza, det in pares:
        try:
            a = extrair_anuncio(cinza, det, familias)
        except Exception as erro:       # uma linha estranha não derruba o relatório inteiro
            falhas.append(f"{type(erro).__name__}: {erro}")
            a = None
        if a is None:
            nao_reconhecidas += 1
        else:
            anuncios.append(a)
    return anuncios, nao_reconhecidas, aviso_cab, falhas


# ------------------------------------------------------- dados do topo da tela

def extrair_contexto(soup: BeautifulSoup, html_bruto: str) -> dict:
    """Conta, filtros ativos, resumo do envio e modo de exibição do estoque ERP, como estavam na tela."""
    ctx = {}
    m = re.search(r"customer_name=([^&\"'\s]+)", html_bruto)
    if m:
        ctx["conta"] = unquote_plus(m.group(1)).title()

    # filtros que estavam marcados (só os que restringem: nem todas as opções marcadas)
    filtros = []
    for sel in soup.select("select"):
        if sel.get("form") != "formFiltros":
            continue
        opcoes = [o for o in sel.find_all("option") if _norm(o.get_text())]
        marcadas = [_norm(o.get_text()) for o in opcoes if o.has_attr("selected")]
        if not marcadas or len(marcadas) == len(opcoes):
            continue
        rotulo = sel.find_previous("p", class_="small")
        nome = _norm(rotulo.get_text()) if rotulo else (sel.get("name") or "filtro")
        filtros.append(f"{nome} = {', '.join(marcadas)}")
    ctx["filtros"] = filtros

    tabela = soup.find("table", id="main_table")
    if tabela is not None:
        tabela.decompose()
    for tag in soup(["script", "style", "svg", "noscript"]):
        tag.decompose()
    junto = " ".join(_norm(t) for t in soup.stripped_strings)
    m = re.search(r"Filtrando por:\s*(.*?)\s*(?:Padrão|Salvar como Padrão)", junto)
    if m:
        ctx["filtrando"] = _norm(m.group(1)).replace(" Período", "  ·  Período")
    m = re.search(r"Qtd do Envio\s*(\d[\d.]*)\s*\(\s*(\d[\d.]*)\s*\)", junto)
    if m:
        ctx["qtd_envio"] = int(m.group(1).replace(".", ""))
        ctx["qtd_envio_entre_parenteses"] = int(m.group(2).replace(".", ""))
    m = re.search(r"Resumo do envio\s*Dias de Estoque\s*(\d+)\s*Sem venda\s*(\d+)\s*Sem Estoque\s*(\d+)\s*Coleta estimada\s*(\d{2}/\d{2}/\d{4})", junto)
    if m:
        ctx["resumo"] = (f"Dias de Estoque {m.group(1)}  ·  Sem venda {m.group(2)}  ·  "
                         f"Sem Estoque {m.group(3)}  ·  Coleta estimada {m.group(4)}")
    return ctx


# ------------------------------------------------------------------ planilha

def agrupar_por_sku(linhas: list[dict]) -> dict[str, list[dict]]:
    grupos = defaultdict(list)
    for linha in linhas:
        grupos[linha["sku"] or SEM_SKU].append(linha)
    return dict(sorted(grupos.items()))


def finalizar_observacoes(grupos: dict[str, list[dict]]):
    """Junta as observações do estoque e avisa quando 2+ anúncios dividem o mesmo SKU (e o mesmo estoque)."""
    for sku, itens in grupos.items():
        for item in itens:
            obs = list(item["erp_obs"]) if isinstance(item["erp_obs"], list) else ([item["erp_obs"]] if item["erp_obs"] else [])
            if len(itens) > 1 and item.get("tipo") in ("Simples", "Kit"):
                obs.append(f"SKU compartilhado por {len(itens)} anúncios — é o mesmo estoque, não some")
            item["erp_obs"] = " | ".join(obs) or None


def resumo_estoque_sku(itens: list[dict]) -> str:
    """Texto do estoque do SKU pra faixa de título do grupo (o estoque é do SKU, igual pra todos os anúncios)."""
    ref = next((i for i in itens if i.get("tipo")), None)
    if ref is None:
        return ""
    if ref["tipo"] == "Sem ERP":
        return "SKU não localizado no ERP"
    prefixo = "KIT  ·  " if ref["tipo"] == "Kit" else ""
    return (f"{prefixo}Estoque ERP — Emp 1: {ref['erp_emp1']}  ·  Emp 4: {ref['erp_emp4']}  ·  "
            f"Total: {ref['erp_total']}")


def soma_estoque_simples(grupos: dict[str, list[dict]]):
    """Soma Emp 1 / Emp 4 / Total dos SKUs simples, contando cada SKU uma vez só."""
    e1 = e4 = tot = n = 0
    for itens in grupos.values():
        ref = next((i for i in itens if i.get("tipo") == "Simples"), None)
        if ref is None:
            continue
        n += 1
        e1 += ref["erp_emp1"] or 0
        e4 += ref["erp_emp4"] or 0
        tot += ref["erp_total"] or 0
    return n, e1, e4, tot


HORIZONTAL = {"esq": "left", "centro": "center", "dir": "right"}


def _exibicao(valor, formato) -> str:
    """Como o valor aparece na célula (só pra medir o tamanho da coluna)."""
    if valor is None:
        return ""
    if formato == "pct":
        return f"{valor * 100:.1f}%"
    if formato == "link":
        return str(valor)
    return str(valor)


def _linhas_do_texto(texto, largura, fator=1.05) -> int:
    """Quantas linhas o texto ocupa numa coluna dessa largura (estimativa com folga pro recuo)."""
    if not texto:
        return 1
    capacidade = max(int(largura * fator) - 2, 1)
    return sum(max(len(wrap(trecho, capacidade)), 1) for trecho in str(texto).split("\n"))


def calcular_larguras(grupos: dict[str, list[dict]]) -> list[float]:
    """Largura de cada coluna = maior valor dela (ou maior palavra do título), entre o mínimo e o máximo."""
    itens = [item for itens_do_grupo in grupos.values() for item in itens_do_grupo]
    larguras = []
    for chave, titulo, largura_min, largura_max, formato, _, _, quebra in COLUNAS:
        if quebra:
            maior_valor = max((max((len(p) for p in str(i.get(chave) or "").split("\n")), default=0)
                               for i in itens), default=0)
        else:
            maior_valor = max((len(_exibicao(i.get(chave), formato)) for i in itens), default=0)
        maior_palavra = max((len(p) for p in titulo.split()), default=0)
        larguras.append(min(max(maior_valor + 3, maior_palavra + 3, largura_min), largura_max))
    return larguras


def gerar_xlsx(grupos: dict[str, list[dict]], nome_conta: str, notas: list[str], avisos: list[str]):
    fonte_titulo = Font(name=FONTE, size=14, bold=True, color=COR_TITULO)
    fonte_subtitulo = Font(name=FONTE, size=10, italic=True, color="6B7280")
    fonte_aviso = Font(name=FONTE, size=10, bold=True, color="B91C1C")
    fonte_faixa = Font(name=FONTE, size=10, bold=True, color="FFFFFF")
    fonte_cabecalho = Font(name=FONTE, size=10, bold=True, color=COR_TITULO)
    fonte_grupo_sku = Font(name=FONTE, size=11, bold=True, color=COR_GRUPO_TEXTO)
    fonte_grupo_info = Font(name=FONTE, size=10, italic=True, color="3B5488")
    fonte_dado = Font(name=FONTE, size=10, color="111827")
    fonte_link = Font(name=FONTE, size=10, color="1D4ED8", underline="single")

    fill_grupo = PatternFill("solid", fgColor=COR_GRUPO_FUNDO)
    fill_zebra = PatternFill("solid", fgColor=COR_ZEBRA)

    # Grade: toda célula tem as 4 bordas; a 1ª coluna de cada faixa tem divisória mais grossa.
    lado = Side(style="thin", color=COR_BORDA)
    lado_divisoria = Side(style="medium", color=COR_DIVISORIA_FAIXA)
    lado_grupo = Side(style="thin", color="60A5FA")

    n_cols = len(COLUNAS)
    ultima_letra = get_column_letter(n_cols)
    inicios_de_faixa = {c for c in range(1, n_cols + 1)
                        if c == 1 or COLUNAS[c - 1][6] != COLUNAS[c - 2][6]}

    def borda_dado(c, cabecalho=False):
        return Border(
            left=lado_divisoria if c in inicios_de_faixa else lado,
            right=lado_divisoria if c == n_cols else lado,
            top=lado,
            bottom=Side(style="medium", color=COR_TITULO) if cabecalho else lado,
        )

    larguras = calcular_larguras(grupos)

    wb = Workbook()
    ws = wb.active
    ws.title = "Envio para o FULL"

    for i, largura in enumerate(larguras, start=1):
        ws.column_dimensions[get_column_letter(i)].width = largura

    # --- Topo: título, resumo, avisos (sem mesclar: o texto transborda pra direita) ---
    linha = 1
    ws.cell(row=linha, column=1, value=f"Magiic — Envio para o FULL ({nome_conta})").font = fonte_titulo
    ws.row_dimensions[linha].height = 22
    linha += 1
    for texto in notas:
        ws.cell(row=linha, column=1, value=texto).font = fonte_subtitulo
        linha += 1
    for aviso in avisos:
        ws.cell(row=linha, column=1, value=aviso).font = fonte_aviso
        linha += 1
    ws.row_dimensions[linha].height = 6
    linha += 1     # respiro antes da tabela

    # --- Faixas coloridas (1 por tipo de dado) + títulos das colunas -----------
    linha_faixa = linha
    linha_cabecalho = linha + 1
    inicio = 1
    while inicio <= n_cols:
        faixa = COLUNAS[inicio - 1][6]
        fim = inicio
        while fim < n_cols and COLUNAS[fim][6] == faixa:
            fim += 1
        cor_escura, _ = FAIXAS[faixa]
        for c in range(inicio, fim + 1):
            ws.cell(row=linha_faixa, column=c).fill = PatternFill("solid", fgColor=cor_escura)
        ws.merge_cells(start_row=linha_faixa, start_column=inicio, end_row=linha_faixa, end_column=fim)
        cel = ws.cell(row=linha_faixa, column=inicio, value=faixa.upper())
        cel.font = fonte_faixa
        cel.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        inicio = fim + 1
    ws.row_dimensions[linha_faixa].height = 20

    linhas_cabecalho = 1
    for c, col in enumerate(COLUNAS, start=1):
        _, cor_clara = FAIXAS[col[6]]
        cel = ws.cell(row=linha_cabecalho, column=c, value=col[1])
        cel.font = fonte_cabecalho
        cel.fill = PatternFill("solid", fgColor=cor_clara)
        cel.alignment = Alignment(horizontal="left" if col[5] == "esq" else "center",
                                  vertical="center", wrap_text=True,
                                  indent=1 if col[5] == "esq" else 0)
        cel.border = borda_dado(c, cabecalho=True)
        linhas_cabecalho = max(linhas_cabecalho, _linhas_do_texto(col[1], larguras[c - 1], fator=1.0))
    ws.row_dimensions[linha_cabecalho].height = max(34, 14 * linhas_cabecalho + 8)

    linha_atual = linha_cabecalho + 1
    ws.freeze_panes = f"C{linha_atual}"     # trava cabeçalho + colunas MLB e SKU

    # --- Dados: 1 faixa de título por SKU + 1 linha por anúncio ----------------
    n_dado = 0
    for sku, itens in grupos.items():
        for c in range(1, n_cols + 1):
            cel = ws.cell(row=linha_atual, column=c)
            cel.fill = fill_grupo
            cel.border = Border(top=lado_grupo, bottom=lado_grupo,
                                left=lado_divisoria if c == 1 else None,
                                right=lado_divisoria if c == n_cols else None)
        # Coluna A: SKU sozinho, sem rótulo nem sufixo — pra copiar/colar direto.
        cel_sku = ws.cell(row=linha_atual, column=1, value=sku)
        cel_sku.font = fonte_grupo_sku
        cel_sku.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        # Contexto na coluna do título (sem mesclar, pra não atrapalhar filtro/ordenação).
        rotulo = "anúncio" if len(itens) == 1 else "anúncios"
        info = f"{len(itens)} {rotulo} neste SKU"
        estoque = resumo_estoque_sku(itens)
        if estoque:
            info += f"   ·   {estoque}"
        cel_info = ws.cell(row=linha_atual, column=3, value=info)
        cel_info.font = fonte_grupo_info
        cel_info.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws.row_dimensions[linha_atual].height = 24
        linha_atual += 1

        for item in itens:
            n_dado += 1
            linhas_da_linha = 1
            for c, col in enumerate(COLUNAS, start=1):
                chave, _, _, _, formato, alinhamento, _, quebra = col
                valor = item.get(chave)
                if chave == "sku" and not valor:
                    valor = sku
                cel = ws.cell(row=linha_atual, column=c, value=valor)
                if isinstance(valor, str) and valor.startswith("="):
                    cel.data_type = "s"     # texto que começa com "=" não pode virar fórmula
                cel.font = fonte_dado
                cel.border = borda_dado(c)
                cel.alignment = Alignment(horizontal=HORIZONTAL[alinhamento], vertical="center",
                                          wrap_text=quebra,
                                          indent=0 if alinhamento == "centro" else 1)
                if formato in FORMATOS:
                    cel.number_format = FORMATOS[formato]
                if n_dado % 2 == 0:
                    cel.fill = fill_zebra
                if formato == "link" and valor and str(valor).startswith("http"):
                    cel.hyperlink = str(valor)
                    cel.font = fonte_link
                cores = CORES_VALOR.get(chave)
                if cores and valor in cores:
                    cel.fill = PatternFill("solid", fgColor=cores[valor][0])
                    cel.font = Font(name=FONTE, size=10, bold=True, color=cores[valor][1])
                if quebra and valor:
                    linhas_da_linha = max(linhas_da_linha, _linhas_do_texto(valor, larguras[c - 1]))
            ws.row_dimensions[linha_atual].height = min(max(24, 13.5 * linhas_da_linha + 8), 220)
            linha_atual += 1

    ws.auto_filter.ref = f"A{linha_cabecalho}:{ultima_letra}{linha_atual - 1}"
    ws.sheet_view.showGridLines = False
    return wb


def salvar_planilha(wb, caminho_saida: Path) -> Path:
    """
    Salva o .xlsx. Se o arquivo estiver aberto no Excel (o Windows trava ele),
    salva com a hora no nome em vez de falhar.
    """
    try:
        wb.save(caminho_saida)
        return caminho_saida
    except PermissionError:
        alternativo = caminho_saida.with_name(
            f"{caminho_saida.stem}_{datetime.now().strftime('%H%M%S')}{caminho_saida.suffix}")
        wb.save(alternativo)
        print(f"AVISO: '{caminho_saida.name}' está aberto no Excel (ou sem permissão). "
              f"Salvei com outro nome.")
        return alternativo


def main():
    parser = argparse.ArgumentParser(
        description="Gera relatório Excel (aba única, agrupado por SKU, com estoque Emp 1 / Emp 4 e todos os dados) "
                    "a partir do HTML da tela 'Envio para o FULL' da Magiic.")
    parser.add_argument("html", type=Path, help="Caminho do HTML salvo da tela da Magiic.")
    parser.add_argument("-o", "--saida", type=Path, default=None,
                        help="Caminho do .xlsx de saída (default: mesmo nome do HTML, extensão .xlsx).")
    parser.add_argument("--conta", type=str, default=None,
                        help="Nome da conta, só pro cabeçalho (default: o que vem no HTML).")
    args = parser.parse_args()

    if not args.html.exists():
        print(f"Arquivo não encontrado: {args.html}", file=sys.stderr)
        sys.exit(1)

    caminho_saida = args.saida or args.html.with_suffix(".xlsx")

    print("Lendo o HTML...")
    with open(args.html, encoding="utf-8", errors="replace") as f:
        html_bruto = f.read()
    soup = BeautifulSoup(html_bruto, "html.parser")

    linhas, nao_reconhecidas, aviso_cab, falhas = extrair_linhas(soup)
    if not linhas:
        print("Nenhuma linha de anúncio reconhecida no HTML — confira se é a página 'Envio para o FULL' da Magiic.",
              file=sys.stderr)
        sys.exit(1)

    ctx = extrair_contexto(soup, html_bruto)
    nome_conta = args.conta or ctx.get("conta") or "Magiic"

    grupos = agrupar_por_sku(linhas)
    finalizar_observacoes(grupos)

    # --- conferências e avisos --------------------------------------------------
    avisos = []
    if aviso_cab:
        avisos.append(aviso_cab)
    soma_envio = sum(a["enviar_hoje"] or 0 for a in linhas)
    if ctx.get("qtd_envio") is not None and soma_envio != ctx["qtd_envio"]:
        avisos.append(f"CONFERIR: a soma do 'Enviar hoje' nas linhas é {soma_envio}, mas a tela mostra "
                      f"Qtd do Envio {ctx['qtd_envio']}. O HTML pode estar incompleto ou ter sido salvo com outro filtro.")
    modos = {a["modo"] for a in linhas if a.get("modo")}
    nomes_modos = ", ".join(MODOS_ESTOQUE.get(m, f"modo {m}") for m in sorted(modos))

    n_skus, e1, e4, tot = soma_estoque_simples(grupos)
    notas = [f"Gerado em {datetime.now().strftime('%d/%m/%Y %H:%M')}  ·  {len(grupos)} SKUs  ·  "
             f"{len(linhas)} anúncios (MLBs) neste arquivo"]
    na_tela = []
    if ctx.get("filtrando"):
        na_tela.append(f"Na tela: {ctx['filtrando']}")
    if ctx.get("filtros"):
        na_tela.append("Filtros marcados: " + "; ".join(ctx["filtros"]))
    if na_tela:
        notas.append("   |   ".join(na_tela))
    resumo = []
    if ctx.get("resumo"):
        resumo.append("Resumo do envio na tela: " + ctx["resumo"])
    if ctx.get("qtd_envio") is not None:
        confere = "confere" if soma_envio == ctx["qtd_envio"] else "NÃO confere"
        resumo.append(f"Qtd do Envio: {ctx['qtd_envio']} ({ctx['qtd_envio_entre_parenteses']})  ·  "
                      f"soma do 'Enviar hoje' nas linhas: {soma_envio} ({confere})")
    if resumo:
        notas.append("   |   ".join(resumo))
    if nomes_modos:
        notas.append(f"Estoque ERP: a tela mostrava '{nomes_modos}'; aqui vão Emp 1 e Emp 4 lidos do detalhamento. "
                     f"SKUs simples (cada SKU 1x): Emp 1 = {e1}  ·  Emp 4 = {e4}  ·  Total = {tot}. "
                     f"Mesmo SKU em 2+ anúncios = mesmo estoque (não some a coluna).")
    notas.append("Trânsito: só o total (o detalhe carrega ao clicar e não existe no HTML). "
                 "Vendas: total do item, como a tela mostra; a lista 'Anúncios do inventário' traz as vendas de cada anúncio, que podem não somar o total.")

    wb = gerar_xlsx(grupos, nome_conta, notas, avisos)
    caminho_final = salvar_planilha(wb, caminho_saida)

    tipos = defaultdict(int)
    for a in linhas:
        tipos[a.get("tipo") or "?"] += 1
    print(f"OK — {len(linhas)} anúncios (MLBs) em {len(grupos)} SKUs "
          f"({tipos.get('Simples', 0)} simples, {tipos.get('Kit', 0)} kits, {tipos.get('Sem ERP', 0)} sem ERP).")
    print(f"Estoque ERP dos SKUs simples (cada SKU 1x): Emp 1 = {e1} | Emp 4 = {e4} | Total = {tot}")
    if ctx.get("qtd_envio") is not None:
        marca = "OK" if soma_envio == ctx["qtd_envio"] else "DIFERENTE"
        print(f"Conferência 'Enviar hoje': soma das linhas = {soma_envio} | tela (Qtd do Envio) = {ctx['qtd_envio']} -> {marca}")
    for aviso in avisos:
        print(f"\n*** {aviso} ***\n")
    if nao_reconhecidas:
        print(f"AVISO: {nao_reconhecidas} linha(s) do HTML não bateram com o padrão esperado e foram ignoradas — vale conferir manualmente.")
    for falha in falhas[:3]:
        print(f"   (erro ao ler uma linha: {falha})")
    com_outros = [a for a in linhas if a["outros"]]
    if com_outros:
        print(f"AVISO: {len(com_outros)} anúncio(s) têm dados que o script não conhecia (coluna 'Outros dados do HTML'). Exemplos:")
        vistos = []
        for a in com_outros:
            for parte in a["outros"].split(" | "):
                if parte not in vistos:
                    vistos.append(parte)
        for parte in vistos[:8]:
            print(f"   - {parte[:120]}")
    print(f"Salvo em: {caminho_final}")


if __name__ == "__main__":
    main()
