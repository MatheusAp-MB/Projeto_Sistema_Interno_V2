# gestao_full/funcoes_auxiliares/comparativo_estoque_exportacao.py

# Função Objetivo: Gera o .xlsx do botão "Exportar Excel" da tela "Full — Comparativo de estoque (debug)": 1 linha por SKU, com as 3 fontes lado a lado (Arquivo do ERP,
# API Sysemp, API Mercado Livre), as contas (diferenças) e os achados. É feito para ANALISAR os números fora do sistema (filtrar, ordenar, cruzar, ou mandar para
# alguém olhar), então é uma tabela "plana": uma linha de títulos, uma coluna por informação, números como número (não como texto) e células vazias onde a tela mostra "—".
#
# * [EXPLICAÇÃO] → Quem monta os DADOS é a própria tela (comparativo_estoque.montar_comparativo, com todas_as_linhas=True): a lista sai inteira, depois da busca, do filtro e da
#                  ordem que estavam na tela. Este arquivo só desenha o Excel a partir desse resultado (por isso não importa comparativo_estoque: recebe o que ele montou).
# * [EXPLICAÇÃO] → Vazio NÃO é zero. Célula vazia = "não sei" (a tela mostra "—"). Já o "zero implícito" é um zero de verdade, deduzido: a API Sysemp só manda linhas com estoque
#                  maior que zero, então, com a coleta COMPLETA, o SKU que não veio vale 0. Ele fica em cinza itálico e as colunas dele aparecem em "API Sysemp | Locais com zero implícito".
# * [EXPLICAÇÃO] → Duas abas: "Comparativo" (os dados) e "Leia-me" (de onde veio cada fonte, que filtro estava ligado, o que cada coluna e cada marca de achado significam).

import io
from decimal import Decimal

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from gestao_full.funcoes_auxiliares import comparativo_estoque_colunas as colunas_mod

NOME_DA_ABA_DADOS = 'Comparativo'
NOME_DA_ABA_LEIA_ME = 'Leia-me'
FORMATO_DATA_HORA = 'dd/mm/yyyy hh:mm'

# O grupo de cada coluna (produto, uma das 3 fontes, conta ou achados) e a cor do cabeçalho dele (o mesmo jeito da tela: cada fonte tem a sua).
GRUPO_PRODUTO, GRUPO_CONTA, GRUPO_ACHADOS = 'produto', 'conta', 'achados'
COR_POR_GRUPO = {GRUPO_PRODUTO: '1E3A5F', colunas_mod.FONTE_ARQUIVO: '7B52AB', colunas_mod.FONTE_SYSEMP: '2E7BB5', colunas_mod.FONTE_ML: 'C0396B',
                 GRUPO_CONTA: '3A4F6B', GRUPO_ACHADOS: '8A6A1E'}
NOME_DO_GRUPO_POR_FONTE = {**colunas_mod.NOME_DA_FONTE, GRUPO_CONTA: 'Conta'}

FONTE_CORPO = Font(name='Arial', size=10)
FONTE_ZERO_IMPLICITO = Font(name='Arial', size=10, italic=True, color='8A94A3')
FONTE_CABECALHO = Font(name='Arial', size=10, bold=True, color='FFFFFF')
FONTE_TITULO = Font(name='Arial', size=14, bold=True, color='1E3A5F')
FONTE_SECAO = Font(name='Arial', size=11, bold=True, color='FFFFFF')
FONTE_ROTULO = Font(name='Arial', size=10, bold=True)
PREENCHIMENTO_SECAO = PatternFill('solid', fgColor='1E3A5F')
PREENCHIMENTO_DIFERENCA = PatternFill('solid', fgColor='FFF1C7')

# De onde vem cada marca de achado, em uma palavra (as duas listas "SKU fora do padrão", a da Sysemp e a do arquivo, têm o MESMO nome: o prefixo desfaz a dúvida).
PREFIXO_DA_MARCA = (('dif_', 'Diferença'), ('so_', 'Presença'), ('sysemp_', 'Sysemp'), ('arquivo_', 'Arquivo'), ('ml_', 'ML'))


# Função Objetivo: O nome do arquivo que o navegador baixa: Comparativo_Estoque_Full_SAMVALE_10_10_26_1530.xlsx (mesmo jeito de data das outras exportações).
def nome_do_arquivo(empresa, agora):
    return f"Comparativo_Estoque_Full_{empresa}_{timezone.localtime(agora).strftime('%d_%m_%y_%H%M')}.xlsx"


# ---------------------------------------------------------------------------
# PEQUENAS AJUDAS
# ---------------------------------------------------------------------------
# Função Objetivo: Decimal -> int (se for inteiro) ou float; None continua None (célula vazia).
def _numero(valor):
    if valor is None:
        return None
    valor = Decimal(valor)
    return int(valor) if valor == valor.to_integral_value() else float(valor)


# Função Objetivo: A data/hora do banco (UTC) como data/hora de Brasília, sem fuso (o Excel não guarda fuso).
def _data_hora(momento):
    return timezone.localtime(momento).replace(tzinfo=None) if momento else None


def _sim_nao(valor):
    return 'Sim' if valor else 'Não'


# Função Objetivo: Escreve um texto na célula sem deixar o Excel tratar como fórmula (um título que comece com "=" viraria conta).
def _escrever_texto(celula, texto):
    celula.value = texto
    if isinstance(texto, str) and texto.startswith('='):
        celula.data_type = 's'


# Função Objetivo: "[Sysemp] SKU fora do padrão": a marca de achado com a origem na frente.
def _rotulo_da_marca(chave, rotulo):
    prefixo = next((nome for inicio, nome in PREFIXO_DA_MARCA if chave.startswith(inicio)), '')
    return f'[{prefixo}] {rotulo}' if prefixo else rotulo


# Função Objetivo: As marcas de achado da empresa, na ordem dos filtros da tela: [(chave, "[Origem] rótulo", dica)].
def _marcas_da_empresa(pagina):
    return [(filtro['chave'], _rotulo_da_marca(filtro['chave'], filtro['rotulo']), filtro['dica'])
            for grupo in pagina['grupos_filtros'] for filtro in grupo['filtros'] if filtro['chave'] != pagina['filtro_padrao']]


# ---------------------------------------------------------------------------
# AS COLUNAS DA ABA "COMPARATIVO"
# ---------------------------------------------------------------------------
# Função Objetivo: Cada coluna do Excel é um dicionário: grupo (decide a cor do cabeçalho), titulo, largura, formato (do número/data), "obter" (a função que tira o valor da
# linha), "explicacao" (vai para a aba Leia-me) e, se for conta, "diferenca" (para pintar o que deu diferente de zero).
def _coluna_excel(grupo, titulo, largura, obter, explicacao, formato=None, **extras):
    return {'grupo': grupo, 'titulo': titulo, 'largura': largura, 'obter': obter, 'explicacao': explicacao, 'formato': formato, **extras}


# Função Objetivo: O título de uma coluna de estoque/conta do Excel, sem ambiguidade: "API Sysemp | Estoque Samvale · id 2 (Flex)".
def _titulo_da_coluna(coluna):
    return f"{NOME_DO_GRUPO_POR_FONTE[coluna['fonte']]} | {coluna['nome']} ({coluna['papel']})"


# Função Objetivo: Monta todas as colunas do Excel, na ordem: produto → fontes e contas (as mesmas da tabela da tela) → o que a tela só mostra na parte que abre → achados.
def _montar_colunas_excel(pagina, linhas):
    colunas = []

    # ① PRODUTO
    colunas += [
        _coluna_excel(GRUPO_PRODUTO, 'SKU', 24, lambda linha: linha['sku'] or '(linha da Sysemp sem SKU)', 'O SKU (o código do produto), já sem espaços nas pontas.'),
        _coluna_excel(GRUPO_PRODUTO, 'Título', 48, lambda linha: linha['titulo'], 'O título do produto (do arquivo do ERP; se o SKU não está nele, do banco ou da Sysemp: ver a coluna ao lado).'),
        _coluna_excel(GRUPO_PRODUTO, 'Marca', 18, lambda linha: linha['marca'], 'A marca do produto (do arquivo do ERP ou do banco).'),
        _coluna_excel(GRUPO_PRODUTO, 'Título e marca vêm de', 14, lambda linha: linha['identificacao_origem'],
                      '"arquivo" (o normal), "banco" (SKU que não está no arquivo, ex.: produto inativo) ou "sysemp" (só a Sysemp conhece o SKU; sem marca).'),
        _coluna_excel(GRUPO_PRODUTO, 'SKU no arquivo do ERP', 11, lambda linha: _sim_nao(linha['no_arquivo']), 'Sim = o SKU está no arquivo do ERP (relatório de produtos ativos).'),
        _coluna_excel(GRUPO_PRODUTO, 'SKU na API Sysemp', 11, lambda linha: _sim_nao(linha['no_sysemp']), 'Sim = o SKU veio na última coleta da API Sysemp (a API só manda SKU com estoque maior que zero).'),
        _coluna_excel(GRUPO_PRODUTO, 'SKU com Código ML do Full', 11, lambda linha: _sim_nao(linha['no_ml']), 'Sim = o SKU tem Código ML do Full no banco (as consultas do Mercado Livre).'),
    ]

    # ② FONTES e ③ CONTAS: as mesmas colunas da tabela da tela, na mesma ordem (a célula i da linha é a coluna i).
    for posicao, coluna in enumerate(pagina['colunas']['todas']):
        eh_conta = coluna['tipo'] == colunas_mod.TIPO_DELTA
        grupo = GRUPO_CONTA if eh_conta else coluna['fonte']
        colunas.append(_coluna_excel(
            grupo, _titulo_da_coluna(coluna), 17 if not eh_conta else 19, (lambda i: lambda linha: linha['celulas'][i]['valor'])(posicao),
            coluna['dica_cabecalho'], formato='+#,##0;-#,##0;0' if eh_conta else '#,##0', diferenca=eh_conta,
            implicito=(lambda i: lambda linha: linha['celulas'][i]['implicito'])(posicao)))

    # O que a tela só mostra na parte que abre: as outras colunas de estoque do arquivo e os locais da Sysemp que a tabela não tem.
    rotulos_do_arquivo = list(dict.fromkeys(item['rotulo'] for linha in linhas for item in linha['arquivo_fora_da_tabela']))
    for rotulo in rotulos_do_arquivo:
        colunas.append(_coluna_excel(
            colunas_mod.FONTE_ARQUIVO, f'Arquivo do ERP | {rotulo} · fora da tabela', 17,
            (lambda nome: lambda linha: next((item['valor'] for item in linha['arquivo_fora_da_tabela'] if item['rotulo'] == nome), None))(rotulo),
            f'Coluna «{rotulo}» do arquivo do ERP. A tela só mostra uma coluna do arquivo na tabela; as outras ficam na parte que abre (aqui vêm todas).', formato='#,##0'))
    locais_fora = {}
    for linha in linhas:
        for local in linha['sysemp_fora_da_tabela']:
            locais_fora.setdefault(local['id'], local['rotulo'])
    for id_local, rotulo in sorted(locais_fora.items(), key=lambda par: par[0]):
        colunas.append(_coluna_excel(
            colunas_mod.FONTE_SYSEMP, f'API Sysemp | {rotulo} · fora da tabela', 17,
            (lambda ident: lambda linha: next((local['valor'] for local in linha['sysemp_fora_da_tabela'] if local['id'] == ident), None))(id_local),
            f'Estoque do local «{rotulo}» (id {id_local}) na API Sysemp. Local que não entra nas contas (sem movimento ou fora do mapa da empresa).', formato='#,##0'))

    # API SYSEMP: detalhes
    def zeros_implicitos(linha):
        ids = [f"id {coluna['origem'][1]}" for coluna, celula in zip(pagina['colunas']['todas'], linha['celulas'])
               if coluna['origem'][0] == 'sysemp_local' and celula['implicito']]
        ids += [f"id {local['id']}" for local in linha['sysemp_fora_da_tabela'] if local['implicito']]
        return ', '.join(ids)

    colunas += [
        _coluna_excel(colunas_mod.FONTE_SYSEMP, 'API Sysemp | Locais com zero implícito', 20, zeros_implicitos,
                      'Locais em que o SKU NÃO veio na API e o zero foi deduzido (a coleta estava completa). Nas colunas, esses zeros aparecem em cinza itálico. Vazio = nenhum.'),
        _coluna_excel(colunas_mod.FONTE_SYSEMP, 'API Sysemp | id_produto', 14, lambda linha: ', '.join(linha['sysemp_ids_produto']),
                      'O id_produto que a Sysemp mandou para o SKU (se vier mais de um, aparece a lista, e o SKU recebe a marca "SKU em mais de um id_produto").'),
        _coluna_excel(colunas_mod.FONTE_SYSEMP, 'API Sysemp | SKU como veio (se diferente)', 20,
                      lambda linha: linha['sysemp_sku_bruto'] if linha['no_sysemp'] and linha['sysemp_sku_bruto'] != linha['sku'] else '',
                      'O SKU exatamente como a API mandou, só quando ele veio diferente do SKU limpo (ex.: com espaço nas pontas).'),
    ]

    # API MERCADO LIVRE: detalhes da consulta
    def do_ml(campo, transformar=lambda valor: valor):
        return lambda linha: transformar(linha['ml'][campo]) if linha['ml'] else None

    colunas += [
        _coluna_excel(colunas_mod.FONTE_ML, 'API Mercado Livre | Códigos ML do SKU', 12, do_ml('n_codigos'), 'Quantos Códigos ML do Full o SKU tem no banco.', formato='#,##0'),
        _coluna_excel(colunas_mod.FONTE_ML, 'API Mercado Livre | Códigos consultados com sucesso', 14, do_ml('n_ok'), 'Quantos desses Códigos ML a última consulta respondeu.', formato='#,##0'),
        _coluna_excel(colunas_mod.FONTE_ML, 'API Mercado Livre | Códigos com erro', 12, do_ml('n_erro'), 'Quantos Códigos ML a última consulta respondeu com erro.', formato='#,##0'),
        _coluna_excel(colunas_mod.FONTE_ML, 'API Mercado Livre | Códigos sem consulta', 12, do_ml('n_sem_consulta'), 'Quantos Códigos ML nunca foram consultados.', formato='#,##0'),
        _coluna_excel(colunas_mod.FONTE_ML, 'API Mercado Livre | Consulta mais antiga', 17, do_ml('mais_antiga', _data_hora), 'A consulta mais antiga entre os Códigos ML do SKU (hora de Brasília).',
                      formato=FORMATO_DATA_HORA),
        _coluna_excel(colunas_mod.FONTE_ML, 'API Mercado Livre | Consulta mais recente', 17, do_ml('mais_recente', _data_hora), 'A consulta mais recente entre os Códigos ML do SKU (hora de Brasília).',
                      formato=FORMATO_DATA_HORA),
        _coluna_excel(colunas_mod.FONTE_ML, 'API Mercado Livre | Aviso do Flex', 30, do_ml('flex_aviso'), 'O aviso do cálculo do Flex (depósito) do SKU, quando existe.'),
    ]

    # ④ ACHADOS
    marcas = _marcas_da_empresa(pagina)
    rotulo_da_marca = {chave: rotulo for chave, rotulo, _ in marcas}
    colunas += [
        _coluna_excel(GRUPO_ACHADOS, 'Achados | Quantidade', 11, lambda linha: linha['n_achados'],
                      'Quantos achados o SKU tem (coisas estranhas nos dados: SKU fora do padrão, repetido, sem consulta etc.). As diferenças entre as fontes e a presença do SKU NÃO entram na conta.',
                      formato='#,##0'),
        _coluna_excel(GRUPO_ACHADOS, 'Achados | Quais (e diferenças entre as fontes)', 70,
                      lambda linha: '; '.join(rotulo_da_marca[chave] for chave, _, _ in marcas if chave in linha['tags']),
                      'As marcas do SKU, cada uma com a origem na frente: [Diferença] = alguma conta deu diferente de zero; [Presença] = o SKU só existe em uma fonte; [Sysemp], [Arquivo] e [ML] = '
                      'achados nos dados de cada fonte. O significado de cada marca está mais abaixo.'),
    ]
    return colunas


# ---------------------------------------------------------------------------
# A ABA "COMPARATIVO"
# ---------------------------------------------------------------------------
def _escrever_aba_dados(planilha, pagina, colunas):
    planilha.title = NOME_DA_ABA_DADOS
    for numero, coluna in enumerate(colunas, start=1):
        celula = planilha.cell(row=1, column=numero, value=coluna['titulo'])
        celula.font = FONTE_CABECALHO
        celula.fill = PatternFill('solid', fgColor=COR_POR_GRUPO[coluna['grupo']])
        celula.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        planilha.column_dimensions[get_column_letter(numero)].width = coluna['largura']
    planilha.row_dimensions[1].height = 80

    for numero_da_linha, linha in enumerate(pagina['linhas'], start=2):
        for numero, coluna in enumerate(colunas, start=1):
            valor = coluna['obter'](linha)
            celula = planilha.cell(row=numero_da_linha, column=numero)
            celula.font = FONTE_CORPO
            if isinstance(valor, (Decimal, int, float)) and not isinstance(valor, bool):
                celula.value = _numero(valor)
                celula.number_format = coluna['formato'] or '#,##0'
                if coluna.get('implicito') and coluna['implicito'](linha):
                    celula.font = FONTE_ZERO_IMPLICITO
                if coluna.get('diferenca') and valor != 0:
                    celula.fill = PREENCHIMENTO_DIFERENCA
            elif valor is not None and valor != '':
                if hasattr(valor, 'year'):
                    celula.value = valor
                    celula.number_format = coluna['formato'] or FORMATO_DATA_HORA
                else:
                    _escrever_texto(celula, valor)

    planilha.freeze_panes = 'C2'  # o cabeçalho e as colunas SKU e Título ficam fixos ao rolar
    planilha.auto_filter.ref = f'A1:{get_column_letter(len(colunas))}{max(1, len(pagina["linhas"]) + 1)}'


# ---------------------------------------------------------------------------
# A ABA "LEIA-ME"
# ---------------------------------------------------------------------------
def _escrever_aba_leia_me(planilha, pagina, colunas, agora):
    planilha.title = NOME_DA_ABA_LEIA_ME
    planilha.column_dimensions['A'].width = 52
    planilha.column_dimensions['B'].width = 110
    linha_atual = [1]

    def escrever(rotulo='', valor='', negrito_rotulo=True):
        celula_rotulo = planilha.cell(row=linha_atual[0], column=1, value=rotulo)
        celula_rotulo.font = FONTE_ROTULO if negrito_rotulo else FONTE_CORPO
        celula_rotulo.alignment = Alignment(vertical='top', wrap_text=True)
        celula_valor = planilha.cell(row=linha_atual[0], column=2)
        _escrever_texto(celula_valor, valor)
        celula_valor.font = FONTE_CORPO
        celula_valor.alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)  # horizontal='left' também para números (ex.: total de SKUs)
        linha_atual[0] += 1

    def secao(texto):
        linha_atual[0] += 1
        for coluna in (1, 2):
            planilha.cell(row=linha_atual[0], column=coluna).fill = PREENCHIMENTO_SECAO
        celula = planilha.cell(row=linha_atual[0], column=1, value=texto)
        celula.font = FONTE_SECAO
        linha_atual[0] += 1

    titulo = planilha.cell(row=1, column=1, value='Comparativo de estoque — Full (debug)')
    titulo.font = FONTE_TITULO
    linha_atual[0] = 3

    filtro_ativo = next((filtro['rotulo'] for grupo in pagina['grupos_filtros'] for filtro in grupo['filtros'] if filtro['chave'] == pagina['filtro']), 'Todos')
    chave_da_ordem, decrescente = colunas_mod.separar_ordem(pagina['ordem'])
    ordem = colunas_mod.registro_de_ordenacao(pagina['colunas']).get(chave_da_ordem, {}).get('rotulo', chave_da_ordem)
    escrever('Empresa', pagina['empresa_nome'])
    escrever('Gerado em', timezone.localtime(agora).strftime('%d/%m/%Y %H:%M'))
    escrever('SKUs neste arquivo', pagina['paginacao']['total'])
    escrever('Busca', pagina['texto'] or '(nenhuma)')
    escrever('Filtro', filtro_ativo)
    escrever('Ordem', f"{ordem} ({'decrescente' if decrescente else 'crescente'})")
    escrever('', 'O arquivo tem a lista INTEIRA que a tela mostrava (todas as páginas), com a mesma busca, o mesmo filtro e a mesma ordem.', negrito_rotulo=False)

    secao('As 3 fontes, como estavam ao gerar este arquivo')
    for chave, nome in (('arquivo', 'Arquivo do ERP'), ('sysemp', 'API Sysemp'), ('ml', 'API Mercado Livre')):
        cartao = pagina['fontes'][chave]
        escrever(nome, cartao.get('mensagem') or '', negrito_rotulo=True)
        for item in cartao['itens']:
            escrever(f"    {item['rotulo']}", ('ATENÇÃO: ' if item['alerta'] else '') + str(item['valor']), negrito_rotulo=False)

    secao('Como ler')
    for rotulo, texto in (
        ('Célula vazia', 'Não sei (a tela mostra "—"). Não é zero.'),
        ('Zero em cinza itálico', 'Zero implícito: a API Sysemp só manda SKU com estoque maior que zero; com a coleta completa, o SKU que não veio vale 0. A coluna "API Sysemp | Locais com zero '
                                  'implícito" diz em quais locais isso aconteceu.'),
        ('Colunas "Conta | Δ ..."', 'Diferença = esquerda − direita, sem tolerância: zero = igual, qualquer outro número = diferente (pintado de amarelo). Positivo = o lado esquerdo tem mais.'),
        ('Quem é Full e quem é Flex', 'Vem do mapa de locais da empresa (Samvale: Full = local 4 da Sysemp; Flex = locais 2 + 10 e a coluna "Estoque" do arquivo). Na Magazine ainda não há mapa: '
                                       'só aparecem as colunas cruas do arquivo e o Mercado Livre.'),
        ('Mercado Livre', 'Só o que já estava gravado no banco (as consultas feitas na tela "Estoque no Full"). Este arquivo não consulta nada.'),
    ):
        escrever(rotulo, texto)

    secao('O que cada coluna da aba Comparativo significa')
    for coluna in colunas:
        escrever(coluna['titulo'], coluna['explicacao'])

    secao('O que cada marca de achado significa (coluna "Achados | Quais")')
    for _, rotulo, dica in _marcas_da_empresa(pagina):
        escrever(rotulo, dica)


# Função Objetivo: O .xlsx inteiro (em bytes, pronto para a resposta HTTP) a partir do que montar_comparativo(..., todas_as_linhas=True) devolveu.
def gerar_excel_comparativo(pagina, agora):
    colunas = _montar_colunas_excel(pagina, pagina['linhas'])
    planilha = Workbook()
    _escrever_aba_dados(planilha.active, pagina, colunas)
    _escrever_aba_leia_me(planilha.create_sheet(), pagina, colunas, agora)
    buffer = io.BytesIO()
    planilha.save(buffer)
    return buffer.getvalue()
